from __future__ import annotations
import os
import re
import csv
import time
import datetime
from typing import Dict, List, Optional, Tuple
from PyQt5 import QtCore, QtGui, QtWidgets
from .shared import *
from .shared import _get_base_dir, _profiles_dir, _tr, _json_mod


# Taxonomic columns that describe the expected classification of every query.
# They must be present in each BLAST table, repeated on every hit row.
QUERY_TAX_COLUMNS = ("Query_Order", "Query_Family", "Query_Genus", "Query_organism")
# Taxonomy of each BLAST hit, written by the BLAST panel only when "Fetch
# organism + taxonomy" is on. Without it every hit would score as a mismatch.
SUBJECT_TAX_COLUMNS = ("Subject_Order", "Subject_Family", "Subject_Genus",
                       "Subject_organism")

# Column names accepted as the identifier column of a query-taxonomy reference
# file. The first one present wins; if none is, the first column is used.
_REF_ID_HEADERS = ("query_name", "sample", "sample_id", "sampleid", "id",
                   "identifier", "code", "voucher", "specimen")

_REF_EXT = (".csv", ".xlsx", ".tsv", ".txt")


def _is_reference(path: str) -> bool:
    return path.lower().endswith(_REF_EXT)


def read_tax_reference(path: str):
    """Read a query-taxonomy reference table.

    The identifier column is the first one carrying a known ID header (Sample,
    ID, Query_name, ...), else the first column of the file. The FOUR columns
    that follow it are taken, in that order, as Order, Family, Genus and
    Organism, whatever their own headers say — or whether there are headers
    at all: unless the first row names a known ID header, it is also read as
    an entry, so a file without a header row keeps its first sample (a real
    header row becomes a harmless entry no sample is ever called).

    Returns (id_column_name, [4 taxonomy column names],
             {identifier: (order, family, genus, organism)}).
    """
    headers, rows = read_blast_rows(path)
    if not headers:
        raise ValueError(f"Empty or unreadable reference file: {os.path.basename(path)}")
    id_i = 0
    named_id = False
    for i, name in enumerate(headers):
        if str(name).strip().lower() in _REF_ID_HEADERS:
            id_i = i
            named_id = True
            break
    if not named_id:
        rows = [list(headers)] + list(rows)
    if len(headers) < id_i + 5:
        raise ValueError(
            f"{os.path.basename(path)}: the identifier column "
            f"'{headers[id_i]}' must be followed by 4 columns "
            f"(Order, Family, Genus, Organism); {len(headers) - id_i - 1} found."
        )
    tax_cols = [str(h).strip() for h in headers[id_i + 1:id_i + 5]]
    table = {}
    for row in rows:
        if id_i >= len(row):
            continue
        key = "" if row[id_i] is None else str(row[id_i]).strip()
        if not key or key in table:
            continue   # first occurrence wins; duplicates are ignored
        values = []
        for j in range(id_i + 1, id_i + 5):
            v = row[j] if j < len(row) else ""
            values.append("" if v is None else str(v).strip())
        table[key] = tuple(values)
    if not table:
        raise ValueError(
            f"{os.path.basename(path)}: no identifier found in column "
            f"'{headers[id_i]}'."
        )
    # Without a recognised ID header the first row may be data: name the
    # column by position rather than by what may be a sample code.
    id_name = (str(headers[id_i]).strip() if named_id
               else f"column {id_i + 1} (first row: '{str(headers[id_i]).strip()}')")
    return id_name, tax_cols, table


# ── Writing the reference taxonomy into a BLAST table ───────────────────────
# Module-level so both Best Sequence and the BLAST panel can apply a reference
# file to a table without duplicating the matching/writing logic.

# Secondary-variant records (secondary_variants.fa / unique_secondary_variants
# .fasta) are named "{sample}_var{i}": the suffix is removed so a variant
# competes with — and takes the taxonomy of — its host sample.
_VARIANT_SUFFIX = re.compile(r"_var\d+$")


def is_variant_header(header) -> bool:
    return bool(_VARIANT_SUFFIX.search(str(header).split(";", 1)[0].strip()))


def sample_id_of(query_name, strip_suffix: str = "") -> str:
    """Sample ID of a BLAST Query_name: the text before the first ';',
    without a secondary-variant '_var{i}' suffix, and with the configured
    suffix removed if present."""
    sample = _VARIANT_SUFFIX.sub("", str(query_name).split(";", 1)[0].strip())
    if strip_suffix and sample.endswith(strip_suffix):
        sample = sample[:-len(strip_suffix)]
    return sample


_ALL_FA = "_all.fa"


def sample_of_header(header, strip_suffix: str = "") -> str:
    """Sample a FASTA record is grouped under in Best Sequence: its sample ID
    (see sample_id_of) without a trailing '_all.fa', so a consensus
    ('{sample}_all.fa') and its variants ('{sample}_var{i}') share one sample
    even when the suffix field was cleared."""
    sample = sample_id_of(header, strip_suffix)
    if sample.endswith(_ALL_FA) and len(sample) > len(_ALL_FA):
        sample = sample[:-len(_ALL_FA)]
    return sample


# ── Cached file reads for the live checks of the panels ─────────────────────
# Keyed by (path, kind) and invalidated by the file's modification time, so a
# check that runs on every keystroke never re-reads an unchanged file.
_FILE_CACHE: dict = {}


def _cached(path: str, kind: str, loader):
    """loader(path), cached; an error is cached too and raised again."""
    try:
        stamp = os.path.getmtime(path)
    except OSError:
        return []
    hit = _FILE_CACHE.get((path, kind))
    if hit is None or hit[0] != stamp:
        try:
            hit = (stamp, loader(path))
        except Exception as exc:
            hit = (stamp, exc)
        _FILE_CACHE[(path, kind)] = hit
    if isinstance(hit[1], Exception):
        raise hit[1]
    return hit[1]


def fasta_headers(path: str) -> List[str]:
    """Headers (without '>') of a FASTA file, [] if unreadable."""
    def load(p):
        opener = __import__("gzip").open if p.lower().endswith(".gz") else open
        with opener(p, "rt", encoding="utf-8", errors="replace") as fh:
            return [ln[1:].strip() for ln in fh if ln.startswith(">")]
    try:
        return _cached(path, "fasta", load)
    except Exception:
        return []


def repeated_headers(headers, spaces_as_underscore: bool = False) -> Dict[str, int]:
    """{header: copies} for every header that occurs more than once.

    *spaces_as_underscore* compares headers as the BLAST panel submits them
    (spaces become '_'), so two that differ only in spaces are one."""
    counts: Dict[str, int] = {}
    for h in headers:
        key = h.replace(" ", "_") if spaces_as_underscore else h
        counts[key] = counts.get(key, 0) + 1
    return {h: n for h, n in counts.items() if n > 1}


def repeated_note(repeats: Dict[str, int], limit: int = 3) -> str:
    """'3 repeated header(s), e.g. a ×2, b ×2, c ×3' (the examples are capped)."""
    if not repeats:
        return ""
    shown = ", ".join(f"{h[:40]}{'…' if len(h) > 40 else ''} ×{n}"
                      for h, n in list(repeats.items())[:limit])
    more = f" and {len(repeats) - limit} more" if len(repeats) > limit else ""
    return f"{len(repeats)} repeated header(s), e.g. {shown}{more}"


def fasta_repeats(path: str) -> Dict[str, int]:
    """repeated_headers() of one FASTA file, cached while the file is unchanged."""
    return _cached(path, "repeats", lambda p: repeated_headers(fasta_headers(p)))


def table_column(path: str, column: str) -> List[str]:
    """Non-empty values of one column of a table (.xlsx/.tsv/.csv), [] if absent."""
    def load(p):
        headers, rows = read_blast_rows(p)
        if column not in headers:
            return []
        i = headers.index(column)
        return [str(r[i]).strip() for r in rows
                if i < len(r) and r[i] is not None and str(r[i]).strip()]
    try:
        return _cached(path, "col:" + column, load)
    except Exception:
        return []


def read_tax_reference_cached(path: str):
    """read_tax_reference for the live panel checks (same errors)."""
    return _cached(path, "taxref", read_tax_reference)


_ID_SEPARATORS = re.compile(r"[-_.\s]+")


def reference_match_check(samples, ref: dict) -> Tuple[str, str, int, int]:
    """How many of *samples* (sample IDs) the reference table knows.

    Returns (message, colour, found, total); message is "" without samples.
    Looks the IDs up as the run does (exact, then case-insensitive). When
    some are missing it also tries ignoring '-', '_', '.' and spaces, to
    point out the usual cause: the same codes written with different
    separators in the data and in the reference (DNS-1343 vs DNS_1343)."""
    samples = [s for s in dict.fromkeys(samples) if s]
    if not samples or not ref:
        return "", "", 0, 0
    fmt_warn = reference_format_warning(ref)
    if fmt_warn:
        return fmt_warn, RED, 0, len(samples)
    ref_lower = {k.lower() for k in ref}
    missing = [s for s in samples if s not in ref and s.lower() not in ref_lower]
    found = len(samples) - len(missing)
    empty = reference_empty_ids(samples, ref)
    empty_txt = (f" · {len(empty)} with empty taxonomy (e.g. {', '.join(empty[:3])})"
                 if empty else "")
    if not missing:
        return (f"Reference check: all {found} sample ID(s) found in the reference"
                f"{empty_txt}.", AMBER if empty else GREEN, found, len(samples))

    def norm(x):
        return _ID_SEPARATORS.sub("", x.lower())
    ref_norm = {}
    for k in ref:
        ref_norm.setdefault(norm(k), k)
    near = [(s, ref_norm[norm(s)]) for s in missing if norm(s) in ref_norm]
    msg = (f"Reference check: {found} of {len(samples)} sample ID(s) found "
           f"in the reference")
    if near:
        s, r = near[0]
        msg += (f" · {len(near)} more would match if '-', '_', '.' and spaces were "
                f"ignored (e.g. sample '{s}' vs reference '{r}'): use the same "
                f"separators in both")
    else:
        msg += " · not found, e.g. " + ", ".join(missing[:3])
    msg += empty_txt
    return msg, (RED if found == 0 else AMBER), found, len(samples)


def reference_format_warning(ref: dict) -> str:
    """Warning when the reference does not look like ID + Order, Family,
    Genus, Organism: taxon names carry no digits, so Order/Family/Genus
    columns that are mostly numbers or codes mean the columns are not where
    they should be (e.g. a lab sheet with plate / well columns after the ID).
    "" when the layout looks right."""
    values = [v for tax in ref.values() for v in tax[:3] if v]
    if not values:
        return ("Reference check: the Order / Family / Genus columns are empty — "
                "the 4 columns after the identifier must hold Order, Family, Genus "
                "and Organism.")
    n_digit = sum(1 for v in values if any(ch.isdigit() for ch in v))
    if n_digit / len(values) > 0.5:
        example = next(iter(ref.values()))
        return ("Reference check: the columns after the identifier look like numbers "
                f"or codes, not taxa (e.g. {' / '.join(x for x in example if x)}) — "
                "the identifier must be followed by Order, Family, Genus and Organism.")
    return ""


def reference_empty_ids(samples, ref: dict) -> List[str]:
    """Samples found in the reference but with all four taxonomy cells empty:
    they get no expected taxonomy, so none of their hits can be judged."""
    ref_lower = {k.lower(): v for k, v in ref.items()}
    out = []
    for s in dict.fromkeys(samples):
        tax = lookup_tax(s, ref, ref_lower) if s else None
        if tax is not None and not any(tax):
            out.append(s)
    return out


def reference_report_lines(samples, ref: dict, unknown=None, limit: int = 50) -> List[str]:
    """Run-log lines naming the samples a run cannot judge taxonomically:
    absent from the reference, present with empty taxonomy, plus a layout
    warning. [] when everything is in order."""
    lines = []
    warn = reference_format_warning(ref)
    if warn:
        lines += ["", "  " + warn]
    if unknown is None:
        ref_lower = {k.lower() for k in ref}
        unknown = [s for s in dict.fromkeys(samples)
                   if s and s not in ref and s.lower() not in ref_lower]
    for title, ids in (("Samples missing from the reference (no expected taxonomy):",
                        sorted(unknown)),
                       ("Samples in the reference with empty taxonomy:",
                        reference_empty_ids(samples, ref))):
        if ids:
            lines += ["", f"  {title} {len(ids)}"]
            lines += [f"    {x}" for x in ids[:limit]]
            if len(ids) > limit:
                lines.append(f"    … {len(ids) - limit} more")
    return lines


def lookup_tax(sample: str, ref: dict, ref_lower: dict):
    """Reference row of *sample*, matched exactly then case-insensitively."""
    tax = ref.get(sample)
    if tax is None:
        tax = ref_lower.get(sample.lower())
    return tax


def apply_reference_tax(path: str, ref: dict, ref_lower: dict, strip_suffix: str = "",
                        ref_path: str = ""):
    """Write the four Query_* columns into a BLAST table, in place.

    Every row is keyed by the sample ID of its Query_name (the text before
    the first ';', with the configured suffix removed), which is the
    identifier the reference file is expected to use. Columns already in the
    table are overwritten; the missing ones are appended at its right end,
    so the original layout and formatting are left alone.

    An .xlsx from the BLAST panel also gets its Tax_level_match column and
    its Summary and Best hit sheets brought up to date (*ref_path* is the
    reference file named in the Summary).

    Returns (rows seen, rows filled, sample IDs absent from the reference).
    """
    if path.lower().endswith(".xlsx"):
        return _apply_reference_xlsx(path, ref, ref_lower, strip_suffix, ref_path)
    return _apply_reference_text(path, ref, ref_lower, strip_suffix)


# Hit (Subject_*) taxonomy columns of a BLAST panel table, in the order of
# the "order"/"family"/"genus"/"organism" ranks of concordance_level.
_SUBJECT_RANK_COLUMNS = ("Subject_Order", "Subject_Family", "Subject_Genus",
                         "Subject_organism")


def _apply_reference_xlsx(path, ref, ref_lower, strip_suffix: str = "",
                          ref_path: str = ""):
    import openpyxl
    from copy import copy
    wb = openpyxl.load_workbook(path)
    sheet = blast_sheet(wb)
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    headers = [(str(c.value).strip() if c.value is not None else "")
               for c in header_row]
    if "Query_name" not in headers:
        wb.close()
        raise ValueError(f"{os.path.basename(path)} has no 'Query_name' column.")
    q_i = headers.index("Query_name")

    # Trailing blank headers are not part of the table: append after the
    # last named column, so no empty column is left in between.
    next_col = max((i for i, h in enumerate(headers, 1) if h), default=0)
    style_src = sheet.cell(row=1, column=q_i + 1)
    col_idx = {}
    for name in QUERY_TAX_COLUMNS:
        if name in headers:
            col_idx[name] = headers.index(name) + 1
        else:
            next_col += 1
            col_idx[name] = next_col
            cell = sheet.cell(row=1, column=next_col, value=name)
            cell._style = copy(style_src._style)

    n_rows = n_filled = 0
    unknown = set()
    for row in sheet.iter_rows(min_row=2):
        value = row[q_i].value if q_i < len(row) else None
        if value is None or str(value).strip() == "":
            continue
        n_rows += 1
        sample = sample_id_of(value, strip_suffix)
        tax = lookup_tax(sample, ref, ref_lower)
        if tax is None:
            unknown.add(sample)
            tax = ("", "", "", "")
        else:
            n_filled += 1
        for name, val in zip(QUERY_TAX_COLUMNS, tax):
            sheet.cell(row=row[0].row, column=col_idx[name], value=val)

    # A BLAST panel table with hit taxonomy: keep its Tax_level_match (the
    # deepest rank where hit and reference agree) in step with the new
    # Query_* values, as the BLAST panel's own reference step does.
    if all(c in headers for c in _SUBJECT_RANK_COLUMNS):
        if "Tax_level_match" in headers:
            m_col = headers.index("Tax_level_match") + 1
        else:
            m_col = next_col + 1
            cell = sheet.cell(row=1, column=m_col, value="Tax_level_match")
            cell._style = copy(style_src._style)
        s_idx = [headers.index(c) + 1 for c in _SUBJECT_RANK_COLUMNS]
        q_idx = [col_idx[c] for c in QUERY_TAX_COLUMNS]
        ranks = ("order", "family", "genus", "organism")
        for r in range(2, sheet.max_row + 1):
            if not str(sheet.cell(row=r, column=q_i + 1).value or "").strip():
                continue
            hit = {k: display_taxon(sheet.cell(row=r, column=c).value)
                   for k, c in zip(ranks, s_idx)}
            qtax = {k: display_taxon(sheet.cell(row=r, column=c).value)
                    for k, c in zip(ranks, q_idx)}
            sheet.cell(row=r, column=m_col, value=concordance_level(hit, qtax))

    # Its Summary and Best hit sheets are built from the hit table: rebuild
    # them, or they would keep the previous reference's identifications.
    from .blast_panel import refresh_summary_sheets
    refresh_summary_sheets(wb, path, {"tax_reference": ref_path,
                                      "strip_suffix": strip_suffix}
                           if ref_path else None)

    tmp = path + ".tmp"
    wb.save(tmp)
    wb.close()
    os.replace(tmp, path)
    return n_rows, n_filled, sorted(unknown)


def _apply_reference_text(path, ref, ref_lower, strip_suffix: str = ""):
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        first = fh.readline()
        fh.seek(0)
        sep = "\t" if "\t" in first else ","
        rows = [r for r in csv.reader(fh, delimiter=sep)]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError(f"Empty BLAST table: {os.path.basename(path)}")
    headers = [c.strip() for c in rows[0]]
    if "Query_name" not in headers:
        raise ValueError(f"{os.path.basename(path)} has no 'Query_name' column.")
    q_i = headers.index("Query_name")
    col_idx = {}
    for name in QUERY_TAX_COLUMNS:
        if name in headers:
            col_idx[name] = headers.index(name)
        else:
            col_idx[name] = len(headers)
            headers.append(name)
    width = len(headers)

    out = [headers]
    n_rows = n_filled = 0
    unknown = set()
    for row in rows[1:]:
        row = list(row) + [""] * (width - len(row))
        value = row[q_i] if q_i < len(row) else ""
        if str(value).strip():
            n_rows += 1
            sample = sample_id_of(value, strip_suffix)
            tax = lookup_tax(sample, ref, ref_lower)
            if tax is None:
                unknown.add(sample)
                tax = ("", "", "", "")
            else:
                n_filled += 1
            for name, val in zip(QUERY_TAX_COLUMNS, tax):
                row[col_idx[name]] = val
        out.append(row)

    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as fh:
        csv.writer(fh, delimiter=sep, lineterminator="\n").writerows(out)
    os.replace(tmp, path)
    return n_rows, n_filled, sorted(unknown)


# Taxonomic rank reached by the best concordant hit, deepest first.
TAX_LEVELS = ("organism", "genus", "family", "order", "none")


def display_taxon(value) -> str:
    """Readable form of a taxon name, kept for the report.

    Empty taxonomic cells may be written as 0 / '-' / '' in the tables.
    NCBI writes organism names with the genus capitalised and underscores
    instead of spaces (Palicourea_purpurea), so the underscores go and the
    original capitalisation is preserved. A name that arrives entirely in
    lower case is capitalised, so the genus reads correctly
    (epidendrum fimbriatum -> Epidendrum fimbriatum).
    """
    if value is None:
        return ""
    text = str(value).strip()
    if text in ("0", "0.0", "-", "", "N/A", "NA", "nan", "None"):
        return ""
    text = text.replace("_", " ")
    if text == text.lower():
        text = text[:1].upper() + text[1:]
    return text


# Open-nomenclature qualifiers: dropped before comparing species names, so
# "Palicourea cf. guianensis" is compared as "palicourea guianensis".
_SPECIES_QUALIFIERS = frozenset({"cf.", "cf", "aff.", "aff", "nr.", "nr", "near"})
# Epithets that do not name a species: "sp.", "spp.", "sp1", "sp.2", ...
_NOT_AN_EPITHET = re.compile(r"^spp?\.?\d*$")


def species_key(name: str) -> str:
    """'genus epithet' (lower case) of a species name, or "" when the name does
    not identify a species: a bare genus, "Genus sp.", "Genus sp. ABC123",
    "Genus spp.". Qualifiers (cf./aff./nr.) are skipped; authors and anything
    after the epithet are ignored."""
    words = [w for w in str(name or "").lower().split()
             if w not in _SPECIES_QUALIFIERS]
    if len(words) < 2 or _NOT_AN_EPITHET.match(words[1]):
        return ""
    return f"{words[0]} {words[1]}"


def concordance_level(hit: dict, qtax: dict) -> str:
    """Deepest rank shared by the expected query taxonomy and the subject.

    Both `hit` and `qtax` are dicts with "order"/"family"/"genus"/"organism"
    keys; both sides are compared in normalised (lower-case) form. Module-level
    so both Best Sequence and the BLAST panel can compute it without
    duplicating the comparison rules.
    """
    def same(rank):
        a = qtax.get(rank, "").lower()
        b = hit.get(rank, "").lower()
        return bool(a) and a == b

    # Species level only when BOTH sides name an actual species: two different
    # "Genus sp." / "Genus cf. x" records, or a bare genus, are not a species
    # match (they fall through to the genus comparison below).
    q_sp = species_key(qtax.get("organism", ""))
    if q_sp and q_sp == species_key(hit.get("organism", "")):
        return "organism"
    if same("genus"):
        return "genus"
    if same("family"):
        return "family"
    if same("order"):
        return "order"
    return "none"


# Fallback for a query that appears in no BLAST table at all: its expected
# taxonomy is simply unknown, since the tables are the only source for it.
_EMPTY_TAX = {"order": "", "family": "", "genus": "", "organism": ""}

# One-line meaning of each flag, printed next to its count in the run log.
_FLAG_HELP = {
    "no_blast_hit":          "(no hit at all in the BLAST table)",
    "hits_below_min_aln":    "(hits exist but all below the minimum alignment)",
    "tax_mismatch":          "(good hits, none matching the expected taxonomy)",
    "no_query_taxonomy":     "(no expected taxonomy for the sample: Query_* empty)",
    "low_taxonomic_support": "(best hit matches only at order level)",
    "near_tie":              "(runner-up within 1 point of the winner)",
    "missing_in":            "(sample not recovered in every run)",
    "ambs":                  "(selected sequence still carries ambiguities)",
}

# ── Scoring weights ──────────────────────────────────────────────────────────
# These are design constants, not user settings, so they are deliberately kept
# out of the panel: their values only make sense relative to the 100-point gap
# between the taxonomic ranks of TAX_BONUS, and raising them breaks the property
# the whole utility rests on — that taxonomic concordance outranks raw score.
#
# BITSCORE_WEIGHT matches that 100-point gap exactly, so the bit score can order
# candidates that reached the same rank but never promote one that reached a
# lower rank.
BITSCORE_WEIGHT = 100.0
# The penalties are tie-breakers. Within a sample the normalised bit scores of
# the candidates typically sit 1-5 points apart, so 2 points per ambiguity is the
# same order of magnitude: enough to separate two otherwise equivalent
# candidates, never enough to override taxonomy (that would take 50 ambiguities).
# The choice is heuristic, but the result is flat over 1-5: on a 517-sample test
# set, any value in that range picked the same sequences (0 was worse — 18
# samples fell back to length/reads alone; 20+ started costing identifications).
AMB_PENALTY = 2.0
GAP_PENALTY = 2.0

_FASTA_EXT = (".fa", ".fas", ".fasta", ".fna")
_BLAST_EXT = (".xlsx", ".tsv", ".csv", ".txt")


def _stem(path: str) -> str:
    """File name without directory and without its (possibly double) extension."""
    name = os.path.basename(path)
    lower = name.lower()
    if lower.endswith(".gz"):
        name = name[:-3]
        lower = name.lower()
    root, _ext = os.path.splitext(name)
    return root


def _is_fasta(path: str) -> bool:
    lower = path.lower()
    if lower.endswith(".gz"):
        lower = lower[:-3]
    return lower.endswith(_FASTA_EXT)


def _is_blast(path: str) -> bool:
    return path.lower().endswith(_BLAST_EXT)


# Files a run's folder holds besides its FASTA + BLAST table. A dropped folder
# is read one level deep, and these must not be taken for inputs: their FASTA
# or table-like extensions would otherwise pair with the wrong file.
_NOT_INPUT_PREFIXES = (
    "missing_seqs_", "nohit_seqs_", "blast_run_log_", "blastfile_run_log_",
    "bestseq-", "bestseq_run_log_", "secondary_variants", "~$",
)


def scan_folder(folder: str) -> Tuple[List[str], List[str]]:
    """FASTA files and BLAST tables directly inside *folder* (subfolders are
    not read), as (inputs, ignored names). Files a run writes beside its
    results (see _NOT_INPUT_PREFIXES) are ignored, and where a table exists
    as .xlsx and also as .tsv/.csv/.txt of the same name (the BLAST panel
    writes both) the .xlsx is the one kept."""
    try:
        names = sorted(os.listdir(folder), key=str.lower)
    except OSError:
        return [], []
    inputs, ignored = [], []
    for name in names:
        path = os.path.join(folder, name)
        if not os.path.isfile(path) or not (_is_fasta(path) or _is_blast(path)):
            continue
        if name.lower().startswith(_NOT_INPUT_PREFIXES):
            ignored.append(name)
        else:
            inputs.append(path)
    xlsx_stems = {_stem(f) for f in inputs if f.lower().endswith(".xlsx")}
    kept = []
    for f in inputs:
        if _is_blast(f) and not f.lower().endswith(".xlsx") and _stem(f) in xlsx_stems:
            ignored.append(os.path.basename(f))
        else:
            kept.append(f)
    return kept, ignored


# Name of the hit table's sheet in the BLAST panel's .xlsx, which also carries
# a Summary and a Best hit sheet ahead of it. Other tables use their first sheet.
BLAST_RESULTS_SHEET = "BLAST Results"


def blast_sheet(wb):
    """The sheet of a workbook that holds the BLAST hit table."""
    if BLAST_RESULTS_SHEET in wb.sheetnames:
        return wb[BLAST_RESULTS_SHEET]
    return wb[wb.sheetnames[0]]


def read_blast_header(path: str) -> List[str]:
    """Return the column names of a BLAST table (.xlsx / .tsv / .csv)."""
    if path.lower().endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = blast_sheet(wb)
            row = next(sheet.iter_rows(values_only=True), ())
            wb.close()
            return [str(c).strip() if c is not None else "" for c in row]
        except Exception:
            return []
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            line = fh.readline().rstrip("\n").rstrip("\r")
        sep = "\t" if "\t" in line else ","
        return [c.strip() for c in line.split(sep)]
    except Exception:
        return []


def read_blast_rows(path: str) -> Tuple[List[str], List[tuple]]:
    """Return (headers, data rows) of a BLAST table (.xlsx / .tsv / .csv)."""
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = blast_sheet(wb)
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        headers = [str(c).strip() if c is not None else "" for c in header]
        data = [r for r in rows if r is not None and any(v is not None for v in r)]
        wb.close()
        return headers, data
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        sample = fh.readline()
        fh.seek(0)
        sep = "\t" if "\t" in sample else ","
        reader = csv.reader(fh, delimiter=sep)
        rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return [c.strip() for c in rows[0]], [tuple(r) for r in rows[1:]]


# ═══════════════════════════════════════════════════════════════════════════
# PAIR DROP ZONE  (one FASTA + its BLAST table per comparison)
# ═══════════════════════════════════════════════════════════════════════════

class _PairDropZone(QtWidgets.QFrame):
    """Drop zone that pairs each FASTA with its BLAST table by file name."""

    pairsChanged = QtCore.pyqtSignal(list)      # list of dicts (complete pairs only)

    _ROW_H   = 84
    _CHROME  = 200
    _EMPTY_H = 200

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("drop_zone")
        self.setAcceptDrops(True)
        self.setFixedHeight(self._EMPTY_H)

        self._fastas: List[str] = []
        self._blasts: List[str] = []
        # Files that came from a dropped folder (paired by name only: see
        # _match), and what was skipped from those folders, shown in the zone.
        self._from_folder: set = set()
        self._notices: List[str] = []
        self._seq_cache: Dict[str, int] = {}
        self._col_cache: Dict[str, List[str]] = {}
        # True when a taxonomy reference file is supplied: the Query_* columns
        # are then written into the tables at run time, so missing ones are a
        # note, not an error.
        self._ref_mode = False

        layout = QtWidgets.QVBoxLayout(self)
        layout.setContentsMargins(10, 6, 10, 6)
        layout.setSpacing(8)

        self._lbl_src_empty  = "Drag/Add FASTA + BLAST files here"
        self._lbl_src_filled = "Runs"
        self._lbl = make_label(self._lbl_src_empty, size=18, color=TEXT_SEC)
        self._lbl.setAlignment(QtCore.Qt.AlignCenter)
        self._lbl.setToolTip(
            "Drop every FASTA (.fa/.fas/.fasta) together with its BLAST table\n"
            "(.xlsx/.tsv/.csv), or the folders that hold them (only the files\n"
            "directly inside each folder are read). Files are paired by name, so\n"
            "each pair must share the same base name — e.g. run1.fa + run1.xlsx.\n"
            "A single FASTA and a single table left without a match are paired\n"
            "anyway, except when they come from a folder.\n"
            "One pair classifies that run; two or more compare them."
        )

        self._drag_icon_lbl = QtWidgets.QLabel()
        self._drag_icon_lbl.setAlignment(QtCore.Qt.AlignCenter)
        self._drag_icon_lbl.hide()

        self._rows_container = QtWidgets.QWidget()
        self._rows_layout = QtWidgets.QVBoxLayout(self._rows_container)
        self._rows_layout.setContentsMargins(0, 0, 0, 0)
        self._rows_layout.setSpacing(8)
        self._rows_container.hide()

        btn_layout = QtWidgets.QHBoxLayout()
        btn_layout.setSpacing(8)
        btn_layout.setAlignment(QtCore.Qt.AlignCenter)
        self._browse_btn = QtWidgets.QPushButton("Add")
        self._browse_btn.setObjectName("secondary_btn")
        self._browse_btn.setFixedWidth(130)
        self._browse_btn.clicked.connect(self._browse_files)
        self._clear_btn = QtWidgets.QPushButton("Clear")
        self._clear_btn.setObjectName("danger_btn")
        self._clear_btn.setFixedWidth(130)
        self._clear_btn.clicked.connect(self.clear)
        self._clear_btn.hide()
        btn_layout.addWidget(self._browse_btn)
        btn_layout.addWidget(self._clear_btn)

        _exp = QtWidgets.QSizePolicy
        self._top_spacer = QtWidgets.QSpacerItem(0, 0, _exp.Minimum, _exp.Expanding)
        self._bot_spacer = QtWidgets.QSpacerItem(0, 0, _exp.Minimum, _exp.Expanding)
        layout.addSpacerItem(self._top_spacer)
        layout.addWidget(self._drag_icon_lbl)
        layout.addWidget(self._lbl)
        layout.addWidget(self._rows_container)
        layout.addLayout(btn_layout)
        layout.addSpacerItem(self._bot_spacer)

    # ── i18n ──────────────────────────────────────────────────────────────

    def retranslateUi(self):
        ctx = "BestSeqPairDropZone"
        self._browse_btn.setText(_tr(ctx, "Add"))
        self._clear_btn.setText(_tr(ctx, "Clear"))
        self._update_display()

    def changeEvent(self, event):
        if event.type() == QtCore.QEvent.LanguageChange:
            self.retranslateUi()
        super().changeEvent(event)

    # ── File inspection helpers ───────────────────────────────────────────

    def _count_seqs(self, path: str) -> int:
        if path in self._seq_cache:
            return self._seq_cache[path]
        try:
            opener = __import__("gzip").open if path.endswith(".gz") else open
            with opener(path, "rt", encoding="utf-8", errors="replace") as fh:
                n = sum(1 for ln in fh if ln.startswith(">"))
        except Exception:
            n = 0
        self._seq_cache[path] = n
        return n

    def _columns(self, path: str) -> List[str]:
        if path not in self._col_cache:
            self._col_cache[path] = read_blast_header(path)
        return self._col_cache[path]

    def _missing_tax_columns(self, path: str) -> List[str]:
        cols = set(self._columns(path))
        return [c for c in QUERY_TAX_COLUMNS if c not in cols]

    def _missing_subject_columns(self, path: str) -> List[str]:
        cols = set(self._columns(path))
        return [c for c in SUBJECT_TAX_COLUMNS if c not in cols]

    def matched_in_table(self, fasta: str, blast: str) -> Tuple[int, int]:
        """(FASTA headers found as a Query_name of the table, headers). The
        run looks hits up by the full header (spaces as '_', as the BLAST
        panel submits them), so an unmatched header gets no hits at all."""
        headers = fasta_headers(fasta)
        queries = set(table_column(blast, "Query_name"))
        n = sum(1 for h in headers
                if h in queries or h.replace(" ", "_") in queries)
        return n, len(headers)

    def set_reference_mode(self, enabled: bool):
        """Accept tables without the Query_* columns (a reference file fills them)."""
        if enabled != self._ref_mode:
            self._ref_mode = enabled
            self._update_display()

    def refresh_columns(self):
        """Re-read the table headers, which change once a run writes the
        Query_* columns into them."""
        self._col_cache = {}
        self._update_display()

    # ── Pairing ───────────────────────────────────────────────────────────

    def _match(self) -> Dict[str, Optional[str]]:
        """{fasta: its BLAST table or None}, matched by base file name.  When
        exactly one FASTA and one table are left unmatched they are paired
        anyway: the 'BLAST web results' tab writes blastfile-<ts>.tsv, whose
        name never matches the FASTA that was searched."""
        blast_by_stem = {_stem(b): b for b in self._blasts}
        match = {fa: blast_by_stem.get(_stem(fa)) for fa in self._fastas}
        used = {_stem(b) for b in match.values() if b}
        free_fa = [fa for fa, bl in match.items() if bl is None]
        free_bl = [s for s in blast_by_stem if s not in used]
        # Not for files from a folder: next to many others, the only two left
        # over are as likely to be unrelated as a pair.
        if (len(free_fa) == 1 and len(free_bl) == 1
                and free_fa[0] not in self._from_folder
                and blast_by_stem[free_bl[0]] not in self._from_folder):
            match[free_fa[0]] = blast_by_stem[free_bl[0]]
        return match

    def pairs(self) -> List[dict]:
        """Complete FASTA+BLAST pairs (see _match)."""
        return [{"fasta": fa, "blast": bl} for fa, bl in self._match().items() if bl]

    def _orphans(self) -> Tuple[List[str], List[str]]:
        match = self._match()
        used = {_stem(b) for b in match.values() if b}
        return ([fa for fa, bl in match.items() if bl is None],
                [b for b in self._blasts if _stem(b) not in used])

    def is_valid(self) -> Tuple[bool, str]:
        """Ready to run? Returns (ok, message shown next to the drop zone)."""
        pairs = self.pairs()
        if not pairs:
            return False, "Add at least one FASTA with its BLAST table."
        no_hit_tax = [os.path.basename(p["blast"]) for p in pairs
                      if self._missing_subject_columns(p["blast"])]
        if no_hit_tax:
            return False, ("No hit taxonomy (Subject_* columns) in: "
                           + ", ".join(no_hit_tax[:3]) + ("…" if len(no_hit_tax) > 3 else "")
                           + " — re-run BLAST with 'Fetch organism + taxonomy' on.")
        unmatched = [os.path.basename(p["fasta"]) for p in pairs
                     if self.matched_in_table(p["fasta"], p["blast"])[0] == 0]
        if unmatched:
            return False, ("No sequence of " + ", ".join(unmatched[:3])
                           + ("…" if len(unmatched) > 3 else "")
                           + " is in its BLAST table (Query_name): check that the "
                             "table belongs to that FASTA.")
        bad = [os.path.basename(p["blast"]) for p in pairs
               if self._missing_tax_columns(p["blast"])]
        if bad and not self._ref_mode:
            return False, ("Missing query taxonomy columns in: " + ", ".join(bad[:3])
                           + ("…" if len(bad) > 3 else ""))
        note = ("Single run: sequences will be classified by taxonomic match "
                "(no selection between runs)." if len(pairs) == 1 else "")
        orphan_fa, orphan_bl = self._orphans()
        if orphan_fa or orphan_bl:
            unpaired = (f"{len(orphan_fa) + len(orphan_bl)} unpaired file(s) "
                        f"will be ignored.")
            return True, (note + "  " + unpaired) if note else unpaired
        return True, note

    # ── Rows ──────────────────────────────────────────────────────────────

    def _make_row(self, fasta: str, blast: Optional[str], index: int) -> QtWidgets.QWidget:
        row = QtWidgets.QWidget()
        row.setObjectName("file_row")
        row.setStyleSheet(f"""
            QWidget#file_row {{
                background-color: {GRAY_BG};
                border-radius: 7px;
                border: 1px solid {GRAY_LINE};
            }}
            QWidget#file_row:hover {{
                background-color: {BLUE_LIGHT};
                border-color: #B8D4F0;
            }}
        """)
        row.setToolTip(fasta if blast is None else f"{fasta}\n{blast}")

        hl = QtWidgets.QHBoxLayout(row)
        hl.setContentsMargins(12, 8, 10, 8)
        hl.setSpacing(10)

        icon = make_label("🧬" if blast else "⚠", size=15)
        icon.setFixedWidth(24)

        center = QtWidgets.QVBoxLayout()
        center.setSpacing(1)
        name_lbl = make_label(os.path.basename(fasta), size=15, color=TEXT_PRI)
        name_lbl.setWordWrap(False)
        repeats = fasta_repeats(fasta)
        if repeats:
            # The run keeps the FIRST record of a repeated header (see _read_fasta).
            name_lbl.setTextFormat(QtCore.Qt.RichText)
            name_lbl.setText(
                f"{__import__("html").escape(os.path.basename(fasta))}  <span style='color:#B45309;"
                f" font-size:13px;'>⚠ {len(repeats)} repeated header(s) — "
                f"only the first record of each is used</span>")
            row.setToolTip(row.toolTip() + "\n\n" + repeated_note(repeats, 8))
        center.addWidget(name_lbl)

        n_seqs = self._count_seqs(fasta)
        if blast is None:
            detail = make_label(
                f"{n_seqs:,} sequences  ·  no BLAST table with this name — ignored",
                size=14, color=RED)
        else:
            n_in, n_all = self.matched_in_table(fasta, blast)
            if n_in == 0:
                in_table = "none in the BLAST table"
            elif n_in < n_all:
                in_table = f"{n_in:,} in the BLAST table"
            else:
                in_table = "all in the BLAST table"
            n_seqs = f"{n_seqs:,} seqs ({in_table})"
            missing = self._missing_tax_columns(blast)
            if n_in == 0:
                detail = make_label(
                    f"{n_seqs}  ·  📊 {os.path.basename(blast)}  ·  "
                    f"its Query_name values do not match these headers",
                    size=14, color=RED)
            elif missing and self._ref_mode:
                detail = make_label(
                    f"{n_seqs}  ·  📊 {os.path.basename(blast)}  ·  "
                    f"{', '.join(missing)} → will be written from the reference",
                    size=14, color=TEXT_HINT)
            elif missing:
                detail = make_label(
                    f"{n_seqs}  ·  📊 {os.path.basename(blast)}  ·  "
                    f"missing: {', '.join(missing)}",
                    size=14, color=RED)
            else:
                detail = make_label(
                    f"{n_seqs}  ·  📊 {os.path.basename(blast)}  ·  "
                    f"query taxonomy ✓",
                    size=14, color=TEXT_HINT)
        detail.setWordWrap(False)
        center.addWidget(detail)

        remove_btn = QtWidgets.QPushButton("✕")
        remove_btn.setFixedSize(26, 26)
        remove_btn.setToolTip("Remove this run")
        remove_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: transparent; color: {TEXT_HINT};
                border: none; border-radius: 5px;
                font-size: 13px; font-weight: bold;
            }}
            QPushButton:hover {{ background-color: {RED_LT}; color: {RED}; }}
            QPushButton:pressed {{ background-color: {RED}; color: white; }}
        """)
        remove_btn.clicked.connect(lambda checked, i=index: self._remove(i))

        hl.addWidget(icon)
        hl.addLayout(center, 1)
        hl.addWidget(remove_btn)
        return row

    def _remove(self, index: int):
        entries = self._entries()
        if not (0 <= index < len(entries)):
            return
        fasta, blast = entries[index]
        if fasta in self._fastas:
            self._fastas.remove(fasta)
        if blast and blast in self._blasts:
            self._blasts.remove(blast)
        self._update_display()
        self.pairsChanged.emit(self.pairs())

    def _entries(self) -> List[Tuple[str, Optional[str]]]:
        """Rows to display: every FASTA with its BLAST table (or None)."""
        return list(self._match().items())

    def _adjust_height(self):
        n = len(self._entries()) + len(self._orphans()[1]) + len(self._notices)
        if n == 0:
            self.setFixedHeight(self._EMPTY_H)
        else:
            self.setFixedHeight(self._CHROME + n * self._ROW_H - (n - 1) * 8)

    def _update_display(self):
        entries = self._entries()
        orphan_blasts = self._orphans()[1]

        while self._rows_layout.count() > 0:
            item = self._rows_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

        if not entries and not orphan_blasts and not self._notices:
            self._rows_container.hide()
            self._clear_btn.hide()
            self._lbl.setText(_tr("BestSeqPairDropZone", self._lbl_src_empty))
            self.setProperty("filled", "false")
        else:
            for i, (fa, bl) in enumerate(entries):
                self._rows_layout.addWidget(self._make_row(fa, bl, i))
            for bl in orphan_blasts:
                lbl = make_label(
                    f"⚠  {os.path.basename(bl)} — no FASTA with this name, ignored",
                    size=14, color=RED)
                lbl.setContentsMargins(12, 4, 10, 4)
                self._rows_layout.addWidget(lbl)
            for note in self._notices:
                lbl = make_label(note, size=14, color=RED)
                lbl.setWordWrap(True)
                lbl.setContentsMargins(12, 4, 10, 4)
                self._rows_layout.addWidget(lbl)
            self._rows_container.show()
            self._clear_btn.show()
            n_pairs = len(self.pairs())
            self._lbl.setText(
                f"{_tr('BestSeqPairDropZone', self._lbl_src_filled)} "
                f"({n_pairs} paired · {len(self._fastas)} FASTA · {len(self._blasts)} BLAST)"
            )
            self.setProperty("filled", "true")

        _sp = QtWidgets.QSizePolicy
        mode = (_sp.Expanding if not entries and not orphan_blasts and not self._notices
                else _sp.Fixed)
        self._top_spacer.changeSize(0, 0, _sp.Minimum, mode)
        self._bot_spacer.changeSize(0, 0, _sp.Minimum, mode)
        self.layout().invalidate()

        self._adjust_height()
        refresh_style(self)
        refresh_style(self._rows_container)

    # ── Input ─────────────────────────────────────────────────────────────

    def _browse_files(self):
        files, _ = QtWidgets.QFileDialog.getOpenFileNames(
            self, _tr("BestSeqPairDropZone", "Select FASTA and BLAST files"), "",
            "FASTA + BLAST (*.fa *.fas *.fasta *.fna *.xlsx *.tsv *.csv);;All (*)"
        )
        if files:
            self._add_files(files)

    def _add_folders(self, folders):
        """Add the FASTA files and BLAST tables directly inside each folder.

        Pairs are matched by base name alone, so a base name that would come
        from two places (two folders, or a folder and a file already loaded)
        is not loaded at all: whichever pair it paired with could be the wrong
        one, and one would silently replace the other. The zone says so."""
        found, notes = [], []
        for folder in folders:
            files, ignored = scan_folder(folder)
            name = os.path.basename(os.path.normpath(folder))
            if not files:
                notes.append(f"⚠  {name}: no FASTA or BLAST table directly inside it.")
            if ignored:
                notes.append(f"{name}: {len(ignored)} file(s) ignored (run outputs or "
                             f"duplicate tables): {', '.join(ignored[:3])}"
                             + ("…" if len(ignored) > 3 else ""))
            found += files

        have = [f for f in self._fastas + self._blasts if f not in found]
        where: Dict[Tuple[bool, str], set] = {}
        for f in found + have:
            where.setdefault((_is_fasta(f), _stem(f)), set()).add(f)
        clash = {k for k, v in where.items() if len(v) > 1 and any(f in found for f in v)}
        if clash:
            skipped = sorted({os.path.basename(f) for f in found
                              if (_is_fasta(f), _stem(f)) in clash})
            notes.append("⚠  Not loaded, the same name comes from more than one place: "
                         + ", ".join(skipped[:4]) + ("…" if len(skipped) > 4 else "")
                         + ". Click Clear if one is already loaded, then drop those folders one at a time.")
            found = [f for f in found if (_is_fasta(f), _stem(f)) not in clash]

        before = len(self._fastas) + len(self._blasts)
        self._add_files(found, from_folder=True)
        if notes or len(self._fastas) + len(self._blasts) == before:
            self._notices.extend(n for n in notes if n not in self._notices)
            self._update_display()
            self.pairsChanged.emit(self.pairs())

    def _add_files(self, paths, from_folder: bool = False):
        added = 0
        for p in paths:
            if from_folder:
                self._from_folder.add(p)
            if _is_fasta(p):
                if p not in self._fastas:
                    self._fastas.append(p)
                    added += 1
            elif _is_blast(p):
                if p not in self._blasts:
                    self._blasts.append(p)
                    added += 1
        if added:
            self._fastas.sort(key=lambda x: os.path.basename(x).lower())
            self._blasts.sort(key=lambda x: os.path.basename(x).lower())
            self._update_display()
            self.pairsChanged.emit(self.pairs())

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()
            self.setProperty("dragging", "true")
            urls = e.mimeData().urls()
            if urls:
                path = urls[0].toLocalFile()
                if path:
                    icon = QtWidgets.QFileIconProvider().icon(QtCore.QFileInfo(path))
                    self._drag_icon_lbl.setPixmap(icon.pixmap(48, 48))
                    self._drag_icon_lbl.show()
            refresh_style(self)

    def dragLeaveEvent(self, e):
        self.setProperty("dragging", "false")
        self._drag_icon_lbl.hide()
        refresh_style(self)

    def dropEvent(self, e):
        self.setProperty("dragging", "false")
        self._drag_icon_lbl.hide()
        refresh_style(self)
        paths = [u.toLocalFile() for u in e.mimeData().urls()]
        folders = [p for p in paths if p and os.path.isdir(p)]
        files = [p for p in paths if p and not os.path.isdir(p) and (_is_fasta(p) or _is_blast(p))]
        if files:
            self._add_files(files)
        if folders:
            self._add_folders(folders)

    def clear(self):
        self._fastas = []
        self._blasts = []
        self._from_folder = set()
        self._notices = []
        self._seq_cache = {}
        self._col_cache = {}
        self._update_display()
        self.pairsChanged.emit([])


# ═══════════════════════════════════════════════════════════════════════════
# QUERY TAXONOMY REFERENCE DROP ZONE  (one .csv / .xlsx / .tsv file)
# ═══════════════════════════════════════════════════════════════════════════

class _ElidingLabel(QtWidgets.QLabel):
    """Single-line label that shrinks its text to the width it is given.

    Used instead of word wrap so the row keeps one constant height: a long
    file path is elided in the middle rather than pushing the text out of
    the frame.
    """

    def __init__(self, text="", parent=None):
        super().__init__(text, parent)
        self._full = text
        self.setWordWrap(False)
        self.setSizePolicy(QtWidgets.QSizePolicy.Ignored,
                           QtWidgets.QSizePolicy.Preferred)
        self.setMinimumWidth(60)

    def setFullText(self, text: str):
        self._full = text or ""
        self._apply()

    def fullText(self) -> str:
        return self._full

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply()

    def _apply(self):
        fm = self.fontMetrics()
        w = max(self.width(), self.minimumWidth())
        text = self._full
        if fm.width(text) > w:
            text = fm.elidedText(text, QtCore.Qt.ElideMiddle, w)
        if text != self.text():
            super().setText(text)


class _RefDropZone(QtWidgets.QFrame):
    """Compact drop zone for the single query-taxonomy reference table.

    The reference file is optional, so this stays one slim row: the hint (or
    the chosen file name) on the left, the buttons on the right. The height
    follows the contents instead of being fixed, which is what used to clip
    the text, and the label is elided rather than wrapped so a long path can
    never push the row out of shape.
    """

    fileChanged = QtCore.pyqtSignal(str)   # path ('' when cleared)

    _HINT = "Drag the reference file here  (.csv, .xlsx, .tsv)"

    # Empty: white, with the same dashed line as the drop zones above it, so it
    # stands out from the grey panel and the spot to drop the file is not lost.
    # Dragging a file over it greys the fill and turns the line blue, as in
    # those zones; one that cannot be used, red; a loaded file is green.
    _QSS = f"""
    QFrame#ref_drop_zone {{
        background-color: {WHITE};
        border: 2px dashed #E6E6E3;
        border-radius: 8px;
        padding: 4px 10px;
    }}
    QFrame#ref_drop_zone[dragging="true"] {{
        background-color: {DROP_DRAG_BG};
        border: {DROP_DRAG_BORDER};
    }}
    QFrame#ref_drop_zone[dragging="invalid"] {{
        background-color: {RED_LT};
        border: 2px dashed {RED};
    }}
    QFrame#ref_drop_zone[filled="true"] {{
        background-color: {GREEN_LT};
        border: 1px solid {GREEN_MID};
    }}
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("ref_drop_zone")
        self.setStyleSheet(self._QSS)
        self.setAcceptDrops(True)
        self.setSizePolicy(QtWidgets.QSizePolicy.Preferred,
                           QtWidgets.QSizePolicy.Fixed)
        self.path = ""

        layout = QtWidgets.QHBoxLayout(self)
        layout.setContentsMargins(10, 6, 10, 6)
        layout.setSpacing(10)

        self._lbl = _ElidingLabel(self._HINT)
        self._lbl.setStyleSheet(f"font-size:15px; color:{TEXT_SEC};")
        # The style sheet size is invisible to QFontMetrics, so set it on the
        # font as well: that is what the elision measures against.
        _f = self._lbl.font()
        _f.setPixelSize(15)
        self._lbl.setFont(_f)
        self._lbl.setAlignment(QtCore.Qt.AlignLeft | QtCore.Qt.AlignVCenter)
        layout.addWidget(self._lbl, 1)

        self._browse_btn = QtWidgets.QPushButton("Add")
        self._browse_btn.setObjectName("secondary_btn")
        self._browse_btn.setMinimumWidth(96)
        self._browse_btn.clicked.connect(self._browse)
        self._clear_btn = QtWidgets.QPushButton("Clear")
        self._clear_btn.setObjectName("danger_btn")
        self._clear_btn.setMinimumWidth(96)
        self._clear_btn.clicked.connect(self.clear)
        self._clear_btn.hide()
        layout.addWidget(self._browse_btn, 0)
        layout.addWidget(self._clear_btn, 0)

        self._tip = (
            "One table (.csv / .xlsx / .tsv) holding the expected classification\n"
            "of every sample. The identifier must be the same sample ID as in the\n"
            "FASTA headers, and the 4 columns after it are read as Order, Family,\n"
            "Genus and Organism, in that order."
        )
        self._render()

    # ── Display ───────────────────────────────────────────────────────────

    def _render(self):
        if self.path:
            self._lbl.setFullText(f"📑  {os.path.basename(self.path)}")
            self._lbl.setToolTip(self.path)
            self._clear_btn.show()
            self.setProperty("filled", "true")
        else:
            self._lbl.setFullText(_tr("BestSeqRefDropZone", self._HINT))
            self._lbl.setToolTip(self._tip)
            self._clear_btn.hide()
            self.setProperty("filled", "false")
        refresh_style(self)

    def set_path(self, path: str):
        self.path = path or ""
        self._render()
        self.fileChanged.emit(self.path)

    def clear(self):
        self.set_path("")

    def retranslateUi(self):
        ctx = "BestSeqRefDropZone"
        self._browse_btn.setText(_tr(ctx, "Add"))
        self._clear_btn.setText(_tr(ctx, "Clear"))
        self._render()

    def changeEvent(self, event):
        if event.type() == QtCore.QEvent.LanguageChange:
            self.retranslateUi()
        super().changeEvent(event)

    # ── Input ─────────────────────────────────────────────────────────────

    def _browse(self):
        f, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, _tr("BestSeqRefDropZone", "Select the taxonomy reference file"),
            "", "Reference table (*.csv *.xlsx *.tsv *.txt);;All (*)"
        )
        if f:
            self.set_path(f)

    def _dropped_paths(self, mime) -> List[str]:
        paths = [u.toLocalFile() for u in mime.urls()]
        return [p for p in paths if p and _is_reference(p)]

    def dragEnterEvent(self, e):
        if not e.mimeData().hasUrls():
            return
        e.acceptProposedAction()
        # Say up front whether the file can be used, the way the pair zone does.
        ok = bool(self._dropped_paths(e.mimeData()))
        self.setProperty("dragging", "true" if ok else "invalid")
        refresh_style(self)

    def dragLeaveEvent(self, e):
        self.setProperty("dragging", "false")
        refresh_style(self)

    def dropEvent(self, e):
        self.setProperty("dragging", "false")
        refresh_style(self)
        paths = self._dropped_paths(e.mimeData())
        if paths:
            self.set_path(paths[0])


# ═══════════════════════════════════════════════════════════════════════════
# BEST SEQUENCE PANEL
# ═══════════════════════════════════════════════════════════════════════════

class BestSeqPanel(QtWidgets.QWidget):
    bestSeqRequested = QtCore.pyqtSignal(list, dict)   # pairs, config dict
    stopRequested    = QtCore.pyqtSignal()             # user clicked Stop

    _SLOT_KEYS = ("info", "files", "identical", "select", "progress", "result", "review")

    def __init__(self, parent=None):
        super().__init__(parent)

        outer_layout = QtWidgets.QVBoxLayout(self)
        outer_layout.setContentsMargins(0, 0, 0, 0)
        outer_layout.setSpacing(0)

        scroll = QtWidgets.QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QtWidgets.QFrame.NoFrame)

        self._inner = QtWidgets.QWidget()
        self._layout = QtWidgets.QVBoxLayout(self._inner)
        self._layout.setContentsMargins(20, 20, 20, 8)
        self._layout.setSpacing(14)
        scroll.setWidget(self._inner)
        outer_layout.addWidget(scroll, 0)

        # ── Title + description ──
        self._lbl_title = make_label("Best Sequence Selector", size=19, bold=True)
        self._lbl_desc = make_label(
            "Drag-and-drop each FASTA together with its BLAST table (.xlsx, .tsv, .csv); "
            "files are paired by base name (run1.fa + run1.xlsx).\n"
            "Two or more runs — picks the best consensus sequence per sample: those "
            "identical in every run are kept as they are, the rest are resolved with the "
            "BLAST hits.\n"
            "A single run — classifies its sequences by taxonomic match, splitting them "
            "into identified and no-hit FASTA files.",
            color=TEXT_SEC
        )
        self._lbl_desc.setWordWrap(True)
        self._lbl_summary_line = make_label(
            "Pair each FASTA with its BLAST table (same base name) to pick the best "
            "sequence per sample, or to classify a single run by taxonomic match.",
            color=TEXT_SEC)
        self._lbl_summary_line.setWordWrap(True)
        self._layout.addWidget(self._lbl_title)
        self._layout.addWidget(self._lbl_summary_line)
        self._layout.addWidget(make_collapsible(self._lbl_desc))

        self._lbl_req = QtWidgets.QLabel(
            "<table cellspacing='0' cellpadding='0'><tr>"
            "<td valign='top'>⚠&nbsp;&nbsp;</td>"
            "<td>Every BLAST table must carry the expected classification of the sample in "
            "<b>each hit row</b>: <b>Query_Order</b>, <b>Query_Family</b>, <b>Query_Genus</b> "
            "and <b>Query_organism</b>.<br>"
            "Empty ranks may be left blank or as 0. Files missing these columns are rejected, "
            "unless a <b>query taxonomy reference file</b> is supplied below: the four columns "
            "are then written into every BLAST table before the run.</td>"
            "</tr></table>"
        )
        self._lbl_req.setWordWrap(True)
        self._lbl_req.setStyleSheet("color:#B45309; font-size:16px;")
        self._lbl_req.hide()   # shown only while a table lacks the Query_* columns

        # ── Settings group ──
        self._settings_box = QtWidgets.QGroupBox("Selection Settings")
        self._settings_box.setStyleSheet(group_box_style())
        sg = QtWidgets.QFormLayout(self._settings_box)
        sg.setLabelAlignment(QtCore.Qt.AlignRight)
        sg.setSpacing(10)
        sg.setContentsMargins(16, 16, 16, 16)

        self._minaln_spin = QtWidgets.QSpinBox()
        self._minaln_spin.setRange(0, 100000)
        self._minaln_spin.setValue(100)
        self._minaln_spin.setFixedWidth(120)
        self._minaln_spin.setToolTip(
            "Hits with a shorter alignment are treated as spurious and discarded,\n"
            "so they cannot drive the taxonomic decision."
        )
        self._lbl_minaln = QtWidgets.QLabel("Minimum alignment length (bp):")
        sg.addRow(self._lbl_minaln, self._minaln_spin)

        self._suffix_edit = QtWidgets.QLineEdit("_all.fa")
        self._suffix_edit.setFixedWidth(300)
        self._suffix_edit.setPlaceholderText("suffix stripped from the sample ID (optional)")
        self._suffix_edit.setToolTip(
            "The sample ID is the text before the first ';' in the header.\n"
            "This suffix is removed from it, e.g. DNS-1343_all.fa;758;807 → DNS-1343."
        )
        self._lbl_suffix = QtWidgets.QLabel("Strip suffix from sample ID:")
        self._lbl_sample_preview = QtWidgets.QLabel("")
        self._lbl_sample_preview.setTextFormat(QtCore.Qt.RichText)
        self._lbl_sample_preview.setStyleSheet(f"color:{TEXT_HINT}; font-size:14px;")
        suffix_row = QtWidgets.QWidget()
        sr = QtWidgets.QHBoxLayout(suffix_row)
        sr.setContentsMargins(0, 0, 0, 0)
        sr.setSpacing(12)
        sr.addWidget(self._suffix_edit)
        sr.addWidget(self._lbl_sample_preview, 1)
        sg.addRow(self._lbl_suffix, suffix_row)
        self._suffix_edit.textChanged.connect(self._on_suffix_changed)

        # ── Optional query-taxonomy reference file ──
        # Without it every BLAST table must already carry the Query_* columns.
        # With it they are written into the tables from a single list, keyed by
        # the sample ID left after the suffix above is stripped.
        self._ref_check = QtWidgets.QCheckBox(
            "I have a reference file with the query taxonomy")
        self._ref_check.setToolTip(
            "Adds Query_Order / Query_Family / Query_Genus / Query_organism to every\n"
            "BLAST table from a single reference list, instead of preparing each\n"
            "table by hand. The tables are rewritten in place."
        )
        self._ref_check.toggled.connect(self._on_ref_toggled)
        sg.addRow(self._ref_check)

        self._ref_zone = _RefDropZone()
        self._ref_zone.fileChanged.connect(self._on_ref_file)
        self._ref_zone.hide()
        sg.addRow(self._ref_zone)

        self._lbl_ref_info = QtWidgets.QLabel("")
        self._lbl_ref_info.setWordWrap(True)
        self._lbl_ref_info.setStyleSheet(f"color:{TEXT_HINT}; font-size:14px;")
        self._lbl_ref_info.hide()
        sg.addRow(self._lbl_ref_info)

        # ── Scoring explanation ──
        self._lbl_rule = make_label(
            "Score (when comparing two or more runs) = taxonomic rank of the "
            "best concordant hit "
            "(species 400 · genus 300 · family 200 · order 100)  +  its bit score "
            f"×{BITSCORE_WEIGHT:g} normalised within the sample  −  {AMB_PENALTY:g} per "
            f"ambiguity  −  {GAP_PENALTY:g} per estimated gap. "
            "Ties are broken by longer sequence, then more reads. "
            "The weights are fixed: taxonomic concordance always outranks raw score.",
            size=15, color=TEXT_HINT)
        self._lbl_rule.setWordWrap(True)
        rule_box = make_collapsible(self._lbl_rule, "How the score works")

        # ── Drop zone (first: the settings below preview its files) ──
        self._drop = _PairDropZone()
        self._drop.pairsChanged.connect(self._on_pairs)
        self._layout.addWidget(self._drop)

        self._lbl_status = make_label("", size=15, color=RED)
        self._lbl_status.setWordWrap(True)
        self._layout.addWidget(self._lbl_status)
        self._layout.addWidget(self._lbl_req)
        self._layout.addWidget(self._settings_box)
        self._layout.addWidget(rule_box)
        self._layout.addStretch()

        # ── Live progress display ──
        self._log_slots = {k: "" for k in self._SLOT_KEYS}
        self._log = QtWidgets.QPlainTextEdit()
        self._log.setReadOnly(True)
        self._log.setFont(QtGui.QFont("Consolas", 9))
        self._log.setMinimumHeight(200)
        self._log.setSizePolicy(
            QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Fixed
        )
        self._log.setStyleSheet(
            f"QPlainTextEdit {{ background:{GRAY_BG}; border:1px solid {GRAY_LINE}; "
            f"border-radius:6px; padding:6px; color:{TEXT_PRI}; margin:0 20px 8px 20px; "
            f"font-family:'Consolas','Courier New',monospace; }}"
        )
        self._log.hide()
        outer_layout.addWidget(self._log, 0)

        # ── Elapsed-time timer ──
        self._info_base   = ""
        self._start_time  = 0.0
        self._elapsed_timer = QtCore.QTimer(self)
        self._elapsed_timer.setInterval(1000)
        self._elapsed_timer.timeout.connect(self._tick_elapsed)

        # ── Footer ──
        footer = QtWidgets.QWidget()
        footer.setObjectName("bestseq_footer")
        footer.setStyleSheet(f"""
            QWidget#bestseq_footer {{
                background: {GRAY_CARD};
                border-top: 1px solid {GRAY_LINE};
            }}
        """)
        fl = QtWidgets.QHBoxLayout(footer)
        fl.setContentsMargins(20, 10, 20, 10)

        self._clear_btn = QtWidgets.QPushButton("Clear")
        self._clear_btn.setObjectName("danger_btn")
        self._clear_btn.setFixedHeight(44)
        self._clear_btn.setFixedWidth(140)
        self._clear_btn.clicked.connect(self._reset)
        fl.addWidget(self._clear_btn)

        self._open_folder_btn = QtWidgets.QPushButton("Open folder")
        self._open_folder_btn.setObjectName("secondary_btn")
        self._open_folder_btn.setFixedHeight(44)
        self._open_folder_btn.hide()
        self._open_folder_btn.clicked.connect(self._open_output_folder)
        fl.addWidget(self._open_folder_btn)

        self._open_results_btn = QtWidgets.QPushButton("Open results")
        self._open_results_btn.setObjectName("secondary_btn")
        self._open_results_btn.setFixedHeight(44)
        self._open_results_btn.hide()
        self._open_results_btn.clicked.connect(self._open_results_file)
        fl.addWidget(self._open_results_btn)

        self._stop_btn = QtWidgets.QPushButton("Stop")
        self._stop_btn.setObjectName("danger_btn")
        self._stop_btn.setFixedHeight(44)
        self._stop_btn.setFixedWidth(120)
        self._stop_btn.hide()
        self._stop_btn.clicked.connect(self.stopRequested)
        fl.addWidget(self._stop_btn)

        fl.addStretch()

        self._run_btn = QtWidgets.QPushButton("Select best sequences  →")
        self._run_btn.setObjectName("primary_btn")
        self._run_btn.setFixedHeight(44)
        self._run_btn.setFixedWidth(300)
        self._run_btn.setEnabled(False)
        self._run_btn.clicked.connect(self._emit_run)
        self._set_run_style(False)
        fl.addWidget(self._run_btn)
        outer_layout.addWidget(footer)

        self._last_outdir = ""
        self._last_report = ""
        self._ref_ok = False          # reference file present and readable
        self._ref_found = None        # sample IDs found in it (None: not checked)
        self._describe_reference()

        self.installEventFilter(self)

    # ── Layout helpers ────────────────────────────────────────────────────

    def eventFilter(self, obj, event):
        """Detect when the panel is resized to adjust the log height."""
        if obj == self and event.type() == QtCore.QEvent.Resize:
            self._adjust_log_height()
        return super().eventFilter(obj, event)

    def _adjust_log_height(self):
        """Set the log height to 30% of the panel."""
        if self._log.isVisible():
            self._log.setFixedHeight(max(int(self.height() * 0.3), 200))

    def showEvent(self, event):
        super().showEvent(event)
        self._adjust_log_height()

    def _set_run_style(self, enabled: bool):
        if enabled:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color: {BLUE}; color: white; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
                f"QPushButton:hover {{ background-color: #0C4A82; }}"
            )
        else:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color: {GRAY_LINE}; color: {TEXT_HINT}; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
            )

    # ── i18n ──────────────────────────────────────────────────────────────

    def retranslateUi(self):
        ctx = "BestSeqPanel"
        self._lbl_title.setText(_tr(ctx, "Best Sequence Selector"))
        self._lbl_desc.setText(_tr(ctx,
            "Drag-and-drop each FASTA together with its BLAST table (.xlsx, .tsv, .csv); "
            "files are paired by base name. With two or more runs it picks the best "
            "consensus sequence per sample; with a single run it classifies its sequences "
            "by taxonomic match."))
        self._settings_box.setTitle(_tr(ctx, "Selection Settings"))
        self._lbl_minaln.setText(_tr(ctx, "Minimum alignment length (bp):"))
        self._lbl_suffix.setText(_tr(ctx, "Strip suffix from sample ID:"))
        self._ref_check.setText(_tr(ctx, "I have a reference file with the query taxonomy"))
        self._ref_zone.retranslateUi()
        self._clear_btn.setText(_tr(ctx, "Clear"))
        self._open_folder_btn.setText(_tr(ctx, "Open folder"))
        self._open_results_btn.setText(_tr(ctx, "Open results"))
        self._run_btn.setText(_tr(ctx, self._run_text()))
        self._drop.retranslateUi()

    def changeEvent(self, event):
        if event.type() == QtCore.QEvent.LanguageChange:
            self.retranslateUi()
        super().changeEvent(event)

    # ── Slots ─────────────────────────────────────────────────────────────

    def _on_ref_toggled(self, checked: bool):
        self._ref_zone.setVisible(checked)
        self._lbl_ref_info.setVisible(checked)
        self._describe_reference()
        self._drop.set_reference_mode(checked and self._ref_ok)
        self._on_pairs(None)

    def _on_ref_file(self, path: str):
        self._describe_reference()
        self._drop.set_reference_mode(self._ref_check.isChecked() and self._ref_ok)
        self._on_pairs(None)

    def _run_text(self) -> str:
        """One run is classified; two or more are compared."""
        return ("Classify sequences  →" if len(self._drop.pairs()) == 1
                else "Select best sequences  →")

    def _on_suffix_changed(self, *_):
        self._update_sample_preview()
        if self._ref_check.isChecked():
            self._describe_reference()

    def _update_sample_preview(self):
        """'header → sample' for the first loaded FASTA, so the grouping is
        visible before running."""
        pairs = self._drop.pairs()
        headers = fasta_headers(pairs[0]["fasta"]) if pairs else []
        if not headers:
            self._lbl_sample_preview.setText("")
            return
        h = headers[0]
        shown = h if len(h) <= 40 else h[:40] + "…"
        sample = sample_of_header(h, self._suffix_edit.text().strip())
        self._lbl_sample_preview.setText(
            f"{shown}  →  sample <b style='color:{BLUE}'>{sample or '—'}</b>")

    def _reference_samples(self) -> List[str]:
        """Sample IDs the reference file must know: those of the FASTA
        headers, derived as the run does for the BLAST Query_name."""
        suffix = self._suffix_edit.text().strip()
        return [sample_id_of(h, suffix)
                for p in self._drop.pairs() for h in fasta_headers(p["fasta"])]

    def _describe_reference(self):
        """Validate the reference file and describe how it will be read."""
        self._ref_found = None
        self._ref_ok = False
        path = self._ref_zone.path
        if not path:
            self._lbl_ref_info.setText(
                "The identifier must be the same sample ID as in the FASTA headers "
                "(the text before the first ';', with the suffix above removed). "
                "The <b>4 columns following it</b> are read, in this order, as "
                "<b>Order, Family, Genus and Organism</b>."
            )
            self._lbl_ref_info.setStyleSheet(f"color:{TEXT_HINT}; font-size:14px;")
            return
        try:
            id_col, tax_cols, table = read_tax_reference_cached(path)
        except Exception as exc:
            self._lbl_ref_info.setText(f"⚠  {exc}")
            self._lbl_ref_info.setStyleSheet(f"color:{RED}; font-size:14px;")
            return
        self._ref_ok = True
        check, colour, found, total = reference_match_check(
            self._reference_samples(), table)
        if total:
            self._ref_found = found
        self._lbl_ref_info.setText(
            f"{len(table):,} entries  ·  identifier: <b>{id_col}</b>  ·  "
            f"<b>{' · '.join(tax_cols)}</b> → Query_Order · Query_Family · "
            f"Query_Genus · Query_organism.<br>"
            f"These columns will be written into every BLAST table dropped above "
            f"(the files are modified in place)."
            + (f"<br><span style='color:{colour}'><b>{check}</b></span>" if check else "")
        )
        self._lbl_ref_info.setStyleSheet(f"color:{TEXT_HINT}; font-size:14px;")

    def _on_pairs(self, pairs=None):
        self._update_sample_preview()
        if self._ref_check.isChecked():
            self._describe_reference()
        self._run_btn.setText(_tr("BestSeqPanel", self._run_text()))
        ok, msg = self._drop.is_valid()
        # The reference file is what makes a table without Query_* columns
        # acceptable, so its own problems are reported first.
        if self._ref_check.isChecked() and self._drop.pairs():
            if not self._ref_zone.path:
                ok, msg = False, ("Add the query taxonomy reference file, "
                                  "or uncheck that option.")
            elif not self._ref_ok:
                ok, msg = False, ("The query taxonomy reference file cannot be used "
                                  "(see the message in the settings).")
        self._lbl_status.setText(msg)
        self._lbl_status.setStyleSheet(
            f"color:{TEXT_HINT};" if ok else f"color:{RED};"
        )
        # The requirement note only matters while a table lacks the columns
        # and no usable reference file will write them.
        need = any(self._drop._missing_tax_columns(p["blast"])
                   for p in self._drop.pairs())
        self._lbl_req.setVisible(
            need and not (self._ref_check.isChecked() and self._ref_ok))
        self._run_btn.setEnabled(ok)
        self._set_run_style(ok)

    def _emit_run(self):
        if (self._ref_check.isChecked() and self._ref_found == 0
                and QtWidgets.QMessageBox.question(
                    self, "Reference file",
                    "No sample ID of the FASTA files was found in the reference "
                    "file, so no query taxonomy would be written (see the check "
                    "in the settings).\n\nRun anyway?",
                    QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
                    QtWidgets.QMessageBox.No) != QtWidgets.QMessageBox.Yes):
            return
        # The scoring weights are design constants (see the module header),
        # not user settings; they travel in cfg so the worker and the run log
        # keep reporting the values actually applied.
        cfg = {
            "min_alignment":   self._minaln_spin.value(),
            "bitscore_weight": BITSCORE_WEIGHT,
            "amb_penalty":     AMB_PENALTY,
            "gap_penalty":     GAP_PENALTY,
            "strip_suffix":    self._suffix_edit.text().strip(),
            "tax_reference":   (self._ref_zone.path
                                if self._ref_check.isChecked() else ""),
        }
        self.bestSeqRequested.emit(self._drop.pairs(), cfg)

    def _open_output_folder(self):
        if self._last_outdir and os.path.isdir(self._last_outdir):
            QtGui.QDesktopServices.openUrl(QtCore.QUrl.fromLocalFile(self._last_outdir))

    def _open_results_file(self):
        if self._last_report and os.path.isfile(self._last_report):
            QtGui.QDesktopServices.openUrl(QtCore.QUrl.fromLocalFile(self._last_report))

    def _reset(self):
        self._drop.clear()
        for k in self._SLOT_KEYS:
            self._log_slots[k] = ""
        self._info_base = ""
        self._elapsed_timer.stop()
        self._log.clear()
        self._log.hide()
        self._lbl_status.setText("")
        self._lbl_req.hide()
        self._lbl_sample_preview.setText("")
        self._open_folder_btn.hide()
        self._open_results_btn.hide()
        self._stop_btn.hide()
        self._run_btn.show()
        self._run_btn.setEnabled(False)
        self._set_run_style(False)
        self._last_outdir = ""
        self._last_report = ""
        self._clear_btn.setEnabled(True)

    # ── Public API (called by MainWindow) ──────────────────────────────────

    def load_run(self, paths: List[str]) -> bool:
        """Start a fresh selection with *paths* (a FASTA and its BLAST table,
        e.g. handed over by the BLAST panel): whatever an earlier selection left
        in the panel (files, log, result buttons) is cleared first, so the
        files come in as a single run and are classified, not compared with
        the previous ones. Returns False, leaving the panel as it is, while a
        run is in progress."""
        # isHidden(), not isVisible(): the panel is behind BLAST when this is
        # called, so isVisible() is False even while a run is in progress.
        if not self._stop_btn.isHidden():
            return False
        self._reset()
        self._drop._add_files(paths)
        return True

    def _rebuild_log(self):
        sep = "─" * 56
        lines = [
            self._log_slots.get("info",      ""),
            sep,
            self._log_slots.get("files",     ""),
            self._log_slots.get("identical", ""),
            self._log_slots.get("select",    ""),
            self._log_slots.get("progress",  ""),
            sep,
            self._log_slots.get("result",    ""),
            self._log_slots.get("review",    ""),
        ]
        self._log.setPlainText("\n".join(lines))

    def update_status(self, key: str, text: str):
        self._log.show()
        self._adjust_log_height()
        if key == "info":
            self._info_base = text
        self._log_slots[key] = text
        self._rebuild_log()

    def _tick_elapsed(self):
        elapsed = int(time.monotonic() - self._start_time)
        h, rem  = divmod(elapsed, 3600)
        m, s    = divmod(rem, 60)
        t_str   = f"{h}h {m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"
        self._log_slots["info"] = f"{self._info_base}  │  Time: {t_str}"
        self._rebuild_log()

    def set_running(self, running: bool):
        self._run_btn.setVisible(not running)
        self._stop_btn.setVisible(running)
        self._clear_btn.setEnabled(not running)
        if running:
            self._start_time = time.monotonic()
            self._elapsed_timer.start()
            self._log.show()
            self._adjust_log_height()
        else:
            self._elapsed_timer.stop()

    def set_progress(self, current: int, total: int):
        if total > 0:
            pct = int(current * 100 / total)
            bar_len = 28
            filled = int(bar_len * current / total)
            bar = "█" * filled + " " * (bar_len - filled)
            self.update_status("progress", f"Progress    │ [{bar}] {pct}%")

    def on_finished(self, outdir: str):
        self.set_running(False)
        # A run with a reference file rewrites the tables, so their headers
        # (and the "missing column" notes) are stale.
        if self._ref_check.isChecked():
            self._drop.refresh_columns()
        self._last_outdir = outdir
        if outdir and os.path.isdir(outdir):
            self._open_folder_btn.show()
            for ext in (".xlsx", ".tsv"):
                matches = sorted(
                    (os.path.join(outdir, f) for f in os.listdir(outdir)
                     if f.endswith(ext) and f.startswith("bestseq-")),
                    key=os.path.getmtime, reverse=True
                )
                if matches:
                    self._last_report = matches[0]
                    self._open_results_btn.show()
                    break
        ok, _msg = self._drop.is_valid()
        self._run_btn.setEnabled(ok)
        self._set_run_style(ok)

    def on_error(self, msg: str):
        self.set_running(False)
        self.update_status("result", f"ERROR       │ {error_summary(msg)}")
        show_error_dialog(self, "Best Sequence error", msg)
        ok, _msg = self._drop.is_valid()
        self._run_btn.setEnabled(ok)
        self._set_run_style(ok)


# ═══════════════════════════════════════════════════════════════════════════
# BEST SEQUENCE WORKER
# ═══════════════════════════════════════════════════════════════════════════

class _BestSeqWorker(QtCore.QThread):
    """
    Select the best consensus sequence per sample among one or more runs.

    With a single run there is nothing to choose between: every sequence is kept
    and the module works as a classifier, splitting the run by taxonomic match.

    With two or more runs, sequences that are identical in every run where the
    sample appears are taken as they are. The rest are scored with the BLAST
    hits of each candidate:

      score = taxonomic bonus of the best concordant hit
            + bitscore_weight * (bit score of that hit, normalised in the sample)
            - amb_penalty * ambs
            - gap_penalty * estgaps

    The taxonomic bonus compares the expected classification of the query
    (Query_Order / Query_Family / Query_Genus / Query_organism, which must be
    present on every hit row) with the classification of the subject.

    A secondary variant only replaces its sample's barcode when the top hit of
    the variant reaches a deeper taxonomic level than the top hit of the
    barcode. Otherwise both identify the same taxon and the bit-score gap is
    noise (a variant 1-2 bases away from the barcode), so the barcode is kept.

    The selected sequences are written to '_all', split by why they are or are
    not identified: '_identified' (best hit concordant at some rank),
    '_no_blast_hit' (BLAST found nothing similar: no hit, or only hits shorter
    than the minimum alignment) and '_tax_mismatch' (good hits that contradict
    the expected taxonomy). A sequence with hits but no expected taxonomy to
    compare them with (sample missing from the reference, or empty there) also
    goes to '_tax_mismatch', flagged 'no_query_taxonomy' in the report. The
    split is decided by taxonomy alone — being identical in every run makes a
    consensus reproducible, not identified.
    """

    statusUpdated   = QtCore.pyqtSignal(str, str)   # (slot_key, text)
    progressUpdated = QtCore.pyqtSignal(int, int)   # current, total
    taskFinished    = QtCore.pyqtSignal(str)        # output directory
    taskError       = QtCore.pyqtSignal(str)

    TAX_BONUS = {"organism": 400, "genus": 300, "family": 200, "order": 100, "none": 0}

    # Secondary-variant header fields -> the metrics used for scoring.
    _VARIANT_KEYS = {"len": "length", "coverage": "reads",
                     "ns": "ambs", "fixed_indels": "estgaps"}

    # Report layout: 4 colour zones (identity · sequence metrics · BLAST · decision)
    _COLUMNS = [
        # zone 1 - which sequence was kept
        "Sample", "Decision", "N_files", "Selected_file", "Header",
        # zone 2 - sequence metrics
        "Length", "Reads", "Ambs", "Estgaps",
        # zone 3 - BLAST evidence: expected taxonomy and the best hit side by
        # side, so a mismatch shows at which rank it breaks
        "N_hits", "N_hits_raw", "Tax_level",
        "Query_Order", "Query_Family", "Query_Genus", "Query_organism",
        "Hit_Order", "Hit_Family", "Hit_Genus", "Hit_organism",
        "Best_hit_acc", "P_identity", "Alignment_length", "Bit_score",
        # zone 4 - how the decision was made
        "Score", "Runner_up_file", "Runner_up_score", "Flag",
    ]
    _ZONE_2 = 5      # first column of zone 2
    _ZONE_3 = 9      # first column of zone 3
    _ZONE_4 = 24     # first column of zone 4
    _NUMERIC_NAMES = frozenset({
        "N_files", "Length", "Reads", "Ambs", "Estgaps", "N_hits", "N_hits_raw",
        "P_identity", "Alignment_length", "Bit_score", "Score", "Runner_up_score",
    })

    _PROGRESS_UNITS = 1000   # loading = first half, selection = second half

    def __init__(self, pairs: List[dict], cfg: dict, parent=None):
        super().__init__(parent)
        self.pairs = list(pairs)
        self.cfg   = dict(cfg)
        self._stop = False
        self._last_pct = -1

    def _emit_progress(self, fraction: float):
        """progressUpdated on a fixed scale, only when the integer % changes:
        every emit rebuilds the whole log widget in the GUI thread."""
        units = int(self._PROGRESS_UNITS * max(0.0, min(1.0, fraction)))
        pct = units * 100 // self._PROGRESS_UNITS
        if pct != self._last_pct:
            self._last_pct = pct
            self.progressUpdated.emit(units, self._PROGRESS_UNITS)

    def stop(self):
        self._stop = True

    # ── Parsing helpers ───────────────────────────────────────────────────

    @staticmethod
    def _read_fasta(path):
        """Return (ordered {header: sequence} dict, duplicated headers).

        Headers are kept verbatim and are the key the BLAST table is matched
        on, so a repeated header cannot be told apart: the FIRST record wins
        and the repeats are returned so the caller can report them instead of
        one silently overwriting the other."""
        seqs = {}
        dups: List[str] = []
        header = None
        chunks = []

        def _keep():
            if header in seqs:
                dups.append(header)
            else:
                seqs[header] = "".join(chunks)

        opener = __import__("gzip").open if path.endswith(".gz") else open
        with opener(path, "rt", encoding="utf-8", errors="replace") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                if line.startswith(">"):
                    if header is not None:
                        _keep()
                    header = line[1:]
                    chunks = []
                else:
                    chunks.append(line)
        if header is not None:
            _keep()
        return seqs, dups

    def _parse_header(self, header: str) -> dict:
        """DNS-1343_all.fa;758;807;ambs=0;estgaps=0 -> sample id + metrics.
        Secondary variants (DNS-1343_var1;type=corrected;len=..;coverage=..;
        fixed_indels=..;Ns=..) map to their host sample and the same metrics."""
        parts = header.split(";")
        # A consensus is '{sample}_all.fa' and its variants '{sample}_var{i}':
        # grouped under one sample even when the suffix field was cleared.
        sample = sample_of_header(header, self.cfg.get("strip_suffix", ""))
        # ambs stays None when the header has no 'ambs='/'Ns=' field; the
        # caller then counts the ambiguities in the sequence itself.
        info = {"sample": sample, "length": None, "reads": None,
                "ambs": None, "estgaps": 0}
        for i, part in enumerate(parts[1:], start=1):
            if "=" in part:
                key, value = part.split("=", 1)
                key = key.strip().lower()
                key = self._VARIANT_KEYS.get(key, key)
                if key in info:
                    try:
                        info[key] = int(float(value))
                    except ValueError:
                        pass
            else:
                try:
                    value = int(float(part))
                except ValueError:
                    continue
                if i == 1:
                    info["length"] = value
                elif i == 2:
                    info["reads"] = value
        return info

    @staticmethod
    def _display(value) -> str:
        return display_taxon(value)

    @staticmethod
    def _clean(value) -> str:
        """Normalised form used for comparisons only (never shown)."""
        return display_taxon(value).lower()

    @staticmethod
    def _to_float(value) -> float:
        # Tables edited in Excel can hold numbers as text ("99.5%", "99,5");
        # reading those as 0 would silently drop the hit (alignment filter).
        v = parse_number(value)
        return 0.0 if v is None else v

    # ── Query taxonomy taken from a reference file ────────────────────────

    def _sample_of(self, query_name) -> str:
        """Sample ID of a BLAST Query_name, parsed like a FASTA header."""
        return sample_id_of(query_name, self.cfg.get("strip_suffix", ""))

    def _apply_reference_tax(self, path: str, ref: dict, ref_lower: dict):
        """Write the four Query_* columns into a BLAST table, in place."""
        return apply_reference_tax(path, ref, ref_lower, self.cfg.get("strip_suffix", ""),
                                   self.cfg.get("tax_reference", ""))

    def _sync_blast_tsv(self, path: str, ref: dict, ref_lower: dict) -> str:
        """After the reference went into a BLAST .xlsx, write it into the .tsv
        of the same name too (the BLAST panel's pair), so the two agree and an
        .xlsx rebuilt from the .tsv later keeps it. Each file is edited in
        place: rebuilding the .xlsx from the .tsv instead would drop whatever
        was changed in it by hand. Returns the run-log line ("" when there is
        no such .tsv); a failure is a warning, since this run reads the .xlsx."""
        if not path.lower().endswith(".xlsx"):
            return ""
        tsv_path = os.path.splitext(path)[0] + ".tsv"
        if not os.path.isfile(tsv_path):
            return ""
        from .blast_panel import apply_reference_and_tax_match
        tsv_name = os.path.basename(tsv_path)
        try:
            n_rows, n_filled, _unknown, _n_match = apply_reference_and_tax_match(
                tsv_path, ref, ref_lower, self.cfg.get("strip_suffix", ""))
        except PermissionError:
            msg = f"{tsv_name} not updated: open in another program"
        except Exception as exc:
            msg = f"{tsv_name} not updated: {exc}"
        else:
            return f"{tsv_name} updated too: {n_filled}/{n_rows} rows filled"
        self.statusUpdated.emit("files", f"Warning     │ {msg}")
        return "⚠ " + msg

    def _load_blast(self, path: str) -> Tuple[Dict[str, List[dict]],
                                              Dict[str, dict],
                                              Dict[str, int]]:
        """Parse one BLAST table.

        Returns (hits, query_tax, raw_counts):
          hits        {query_name: [hit, ...]} ordered by Hit_rank, keeping only
                      hits at or above the minimum alignment length;
          query_tax   {query_name: expected taxonomy} read from every row,
                      including the rows dropped by that filter, so the expected
                      classification is still reported for a query whose hits
                      were all too short;
          raw_counts  {query_name: rows present in the table} before filtering.
        """
        headers, rows = read_blast_rows(path)
        if not headers:
            raise ValueError(f"Empty or unreadable BLAST table: {os.path.basename(path)}")
        idx = {name: i for i, name in enumerate(headers)}
        missing = [c for c in QUERY_TAX_COLUMNS if c not in idx]
        if missing:
            raise ValueError(
                f"{os.path.basename(path)} is missing the query taxonomy column(s): "
                f"{', '.join(missing)}"
            )
        if "Query_name" not in idx:
            raise ValueError(
                f"{os.path.basename(path)} has no 'Query_name' column."
            )
        missing = [c for c in SUBJECT_TAX_COLUMNS if c not in idx]
        if missing:
            raise ValueError(
                f"{os.path.basename(path)} has no hit taxonomy ({', '.join(missing)}): "
                f"re-run BLAST with 'Fetch organism + taxonomy' on."
            )

        def get(row, name, default=None):
            i = idx.get(name)
            if i is None or i >= len(row):
                return default
            return row[i]

        min_aln = float(self.cfg.get("min_alignment", 100))
        table: Dict[str, List[dict]] = {}
        query_tax: Dict[str, dict] = {}
        raw_counts: Dict[str, int] = {}
        for row in rows:
            query = get(row, "Query_name")
            if query is None or str(query).strip() == "":
                continue
            query = str(query).strip()
            raw_counts[query] = raw_counts.get(query, 0) + 1
            # The expected taxonomy is a property of the query, repeated on
            # every hit row, so it is read before the alignment-length filter.
            if query not in query_tax:
                query_tax[query] = {
                    "order":    self._display(get(row, "Query_Order")),
                    "family":   self._display(get(row, "Query_Family")),
                    "genus":    self._display(get(row, "Query_Genus")),
                    "organism": self._display(get(row, "Query_organism")),
                }
            alen = self._to_float(get(row, "Alignment_length"))
            if alen < min_aln:
                continue   # spurious short hit: carries no taxonomic information
            hit = {
                "rank":     self._to_float(get(row, "Hit_rank", 99)) or 99,
                "acc":      str(get(row, "Subject_accession.ver", "") or ""),
                "pident":   self._to_float(get(row, "P_identity")),
                "alen":     alen,
                "bit":      self._to_float(get(row, "Bit_score")),
                "order":    self._display(get(row, "Subject_Order")),
                "family":   self._display(get(row, "Subject_Family")),
                "genus":    self._display(get(row, "Subject_Genus")),
                "organism": self._display(get(row, "Subject_organism")),
            }
            table.setdefault(query, []).append(hit)
        for hits in table.values():
            hits.sort(key=lambda h: h["rank"])
        return table, query_tax, raw_counts

    # ── Scoring ───────────────────────────────────────────────────────────

    def _evaluate(self, hits: List[dict],
                  qtax: dict) -> Tuple[Optional[dict], str]:
        """Best taxonomically concordant hit of one candidate sequence."""
        best = None
        best_level = "none"
        for hit in hits:
            level = concordance_level(hit, qtax)
            deeper = self.TAX_BONUS[level] > self.TAX_BONUS[best_level]
            same_level_better_bit = (level == best_level and best is not None
                                     and hit["bit"] > best["bit"])
            if best is None or deeper or same_level_better_bit:
                best = hit
                best_level = level
        return best, best_level

    # ── XLSX export (same look as the BLAST results workbook) ─────────────

    def _tsv_to_xlsx(self, tsv_path: str) -> str:
        """Convert *tsv_path* to a formatted xlsx. Returns xlsx path or '' on failure."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError as e:
            self.statusUpdated.emit("result", f"XLSX skip  │ openpyxl not available: {e}")
            return ""
        try:
            with open(tsv_path, "r", encoding="utf-8") as fh:
                lines = [l for l in fh.read().splitlines() if l.strip()]
        except Exception as e:
            self.statusUpdated.emit("result", f"XLSX skip  │ could not read TSV: {e}")
            return ""
        if not lines:
            return ""

        headers = lines[0].split("\t")
        n_cols  = len(headers)

        # openpyxl 3.x requires 8-char ARGB hex strings (alpha + RGB).
        # Zone 1 → sample identity, Zone 2 → sequence metrics,
        # Zone 3 → BLAST evidence,  Zone 4 → decision.
        H1 = PatternFill(patternType="solid", fgColor="FF1A365D")   # header navy
        H2 = PatternFill(patternType="solid", fgColor="FF0D5E6E")   # header teal
        H3 = PatternFill(patternType="solid", fgColor="FF7C3200")   # header burnt-orange
        H4 = PatternFill(patternType="solid", fgColor="FF4C1D6B")   # header purple
        D1 = PatternFill(patternType="solid", fgColor="FFE8F1FB")   # data light-blue
        D2 = PatternFill(patternType="solid", fgColor="FFE8F5F6")   # data light-teal
        D3 = PatternFill(patternType="solid", fgColor="FFFEF3E8")   # data light-orange
        D4 = PatternFill(patternType="solid", fgColor="FFF3EAFB")   # data light-purple

        white_bold  = Font(color="FFFFFFFF", bold=True, size=10)
        normal_font = Font(size=10)
        bold_font   = Font(size=10, bold=True)          # rows resolved with BLAST
        flag_font   = Font(size=10, color="FFB45309")   # rows carrying a warning
        hdr_align   = Alignment(horizontal="center", vertical="center")
        dat_align   = Alignment(vertical="center", wrap_text=False)
        thin        = Side(style="thin", color="FFCCCCCC")
        border      = Border(left=thin, right=thin, top=thin, bottom=thin)

        _numeric_idx = frozenset(
            i for i, h in enumerate(headers) if h in self._NUMERIC_NAMES)

        def _num(v):
            if v == "":
                return v
            try:
                return int(v)
            except ValueError:
                pass
            try:
                return float(v)
            except ValueError:
                return v          # leave genuinely non-numeric text as-is

        def _zone(ci):
            if ci < self._ZONE_2:
                return 1
            if ci < self._ZONE_3:
                return 2
            if ci < self._ZONE_4:
                return 3
            return 4

        _H = {1: H1, 2: H2, 3: H3, 4: H4}
        _D = {1: D1, 2: D2, 3: D3, 4: D4}

        wb = Workbook()
        ws = wb.active
        ws.title = "Best sequences"

        ws.append(headers)
        for ci in range(n_cols):
            cell = ws.cell(row=1, column=ci + 1)
            cell.fill      = _H[_zone(ci)]
            cell.font      = white_bold
            cell.alignment = hdr_align
            cell.border    = border
        ws.row_dimensions[1].height = 22

        try:
            _dec_idx = headers.index("Decision")
        except ValueError:
            _dec_idx = 1
        try:
            _flag_idx = headers.index("Flag")
        except ValueError:
            _flag_idx = -1

        # One row per sample: alternate the tint every other row for legibility,
        # and bold the samples that actually needed a BLAST-based decision.
        for rn, line in enumerate(lines[1:], start=2):
            raw_vals = line.split("\t")
            while len(raw_vals) < n_cols:
                raw_vals.append("")
            raw_vals = raw_vals[:n_cols]

            decided = (raw_vals[_dec_idx] == "resolved_by_score"
                       if _dec_idx < n_cols else False)
            tinted = (rn % 2 == 0)

            vals = [_num(v) if ci in _numeric_idx else v
                    for ci, v in enumerate(raw_vals)]
            ws.append(vals)
            for ci in range(n_cols):
                cell = ws.cell(row=rn, column=ci + 1)
                if tinted:
                    cell.fill = _D[_zone(ci)]
                if ci == _flag_idx and raw_vals[ci]:
                    cell.font = flag_font
                else:
                    cell.font = bold_font if decided else normal_font
                cell.alignment = dat_align
                cell.border    = border

        ws.freeze_panes    = "A2"
        ws.auto_filter.ref = ws.dimensions

        for ci, col_cells in enumerate(ws.columns):
            width = max((len(str(c.value or "")) for c in col_cells), default=8)
            ws.column_dimensions[get_column_letter(ci + 1)].width = min(width + 2, 55)

        xlsx_path = tsv_path.rsplit(".", 1)[0] + ".xlsx"
        try:
            wb.save(xlsx_path)
        except Exception as e:
            self.statusUpdated.emit("result", f"XLSX error │ {e}")
            return ""
        return xlsx_path

    # ── Main run ─────────────────────────────────────────────────────────

    def run(self):
        try:
            self._run_selection()
        except Exception as e:
            import traceback
            self.taskError.emit(f"{e}\n{traceback.format_exc()}")

    def _run_selection(self):
        cfg        = self.cfg
        run_start  = datetime.datetime.now()
        mydate     = run_start.strftime("%Y%m%d-%H%M%S")
        output_dir = cfg["outdir"]
        os.makedirs(output_dir, exist_ok=True)

        amb_penalty = float(cfg.get("amb_penalty", AMB_PENALTY))
        gap_penalty = float(cfg.get("gap_penalty", GAP_PENALTY))
        bit_weight  = float(cfg.get("bitscore_weight", BITSCORE_WEIGHT))

        n_files = len(self.pairs)
        self.statusUpdated.emit(
            "info",
            f"Runs: {n_files}  │  Min alignment: {cfg.get('min_alignment', 100)} bp"
            f"  │  Bit weight: {bit_weight:g}"
        )

        # ── Query taxonomy written into the BLAST tables from a reference ──
        # Without a reference every table must already carry the Query_* columns
        # (_load_blast rejects it otherwise). With one, they are written now, in
        # place, so a single list of samples serves every run.
        ref_path  = cfg.get("tax_reference", "")
        ref_lines: List[str] = []
        if ref_path:
            id_col, tax_cols, ref_table = read_tax_reference(ref_path)
            ref_lower = {k.lower(): v for k, v in ref_table.items()}
            self.statusUpdated.emit(
                "files",
                f"Reference   │ {len(ref_table)} entries · identifier '{id_col}'"
            )
            ref_lines += [
                "",
                "  Query taxonomy reference:",
                f"    File        : {os.path.abspath(ref_path)}",
                f"    Identifier  : {id_col}",
                f"    Columns     : {' · '.join(tax_cols)} → "
                f"{' · '.join(QUERY_TAX_COLUMNS)}",
                f"    Entries     : {len(ref_table)}",
            ]
            all_unknown = set()
            for i, pair in enumerate(self.pairs):
                if self._stop:
                    break
                name = os.path.basename(pair["blast"])
                self.statusUpdated.emit(
                    "files",
                    f"Reference   │ [{i + 1}/{n_files}] writing query taxonomy into {name}…"
                )
                try:
                    n_rows, n_filled, unknown = self._apply_reference_tax(
                        pair["blast"], ref_table, ref_lower)
                except PermissionError:
                    self.taskError.emit(
                        f"Could not write the query taxonomy into {name}: the file is "
                        f"open in another program. Close it and run again."
                    )
                    return
                all_unknown.update(unknown)
                ref_lines.append(
                    f"    {name}: {n_filled}/{n_rows} rows filled"
                    + (f" · {len(unknown)} sample(s) not in the reference"
                       if unknown else "")
                )
                tsv_line = self._sync_blast_tsv(pair["blast"], ref_table, ref_lower)
                if tsv_line:
                    ref_lines.append(f"      {tsv_line}")
            _fmt = reference_format_warning(ref_table)
            if _fmt:
                self.statusUpdated.emit("files", "Warning     │ " + _fmt)
            if all_unknown:
                self.statusUpdated.emit(
                    "files",
                    f"Reference   │ Query taxonomy written · {len(all_unknown)} "
                    f"sample(s) not found in the reference"
                )
            if all_unknown or _fmt:
                # Empty-taxonomy samples are listed once the run's samples
                # are known (after loading, below).
                ref_lines += reference_report_lines([], ref_table,
                                                    unknown=all_unknown)
            else:
                self.statusUpdated.emit(
                    "files",
                    f"Reference   │ Query taxonomy written into {n_files} table(s)  ✓"
                )
            if self._stop:
                self.statusUpdated.emit("result", "Stopped     │ selection cancelled")
                self.taskFinished.emit(output_dir)
                return

        # ── Load every FASTA + BLAST pair ──
        candidates: Dict[str, List[dict]] = {}
        dup_lines: List[str] = []   # run-log lines for FASTAs with repeated headers
        nomatch_lines: List[str] = []   # FASTAs whose headers are missing from the table
        file_labels: List[str] = []
        total_seqs = 0
        for i, pair in enumerate(self.pairs):
            if self._stop:
                break
            label = _stem(pair["fasta"])
            file_labels.append(label)
            self.statusUpdated.emit(
                "files", f"Loading     │ [{i + 1}/{n_files}] {os.path.basename(pair['fasta'])}"
            )
            seqs, dups = self._read_fasta(pair["fasta"])
            if dups:
                dup_lines.append(
                    f"    {os.path.basename(pair['fasta'])}: {len(dups)} repeated "
                    f"header(s), first record kept — e.g. {dups[0]}")
                self.statusUpdated.emit(
                    "files",
                    f"Warning     │ {os.path.basename(pair['fasta'])}: {len(dups)} "
                    f"repeated header(s) ignored (first record kept)")
            blast, query_tax, raw_counts = self._load_blast(pair["blast"])
            # Hits are looked up by the full header: say so when headers are
            # absent from the table, or those sequences silently get no hits.
            n_in = sum(1 for h in seqs
                       if h in raw_counts or h.replace(" ", "_") in raw_counts)
            if n_in < len(seqs):
                nomatch_lines.append(
                    f"    {os.path.basename(pair['fasta'])}: {len(seqs) - n_in} of "
                    f"{len(seqs)} header(s) not in {os.path.basename(pair['blast'])}")
                if n_in == 0:
                    self.statusUpdated.emit(
                        "files",
                        f"Warning     │ {os.path.basename(pair['fasta'])}: no header found "
                        f"in its BLAST table (Query_name) — every sequence gets no hit")
            total_seqs += len(seqs)
            for header, seq in seqs.items():
                info = self._parse_header(header)
                if info["ambs"] is None:
                    info["ambs"] = sum(1 for c in seq.upper() if c not in "ACGT-")
                if info["length"] is None:
                    info["length"] = len(seq)
                # The BLAST panel submits headers with spaces replaced by '_',
                # so a header with spaces is looked up in that form too.
                qkey = header
                if qkey not in raw_counts and " " in header:
                    qkey = header.replace(" ", "_")
                hits = blast.get(qkey, [])
                qtax = query_tax.get(qkey, _EMPTY_TAX)
                best_hit, level = self._evaluate(hits, qtax)
                # Level of the TOP hit (highest bit score), used to decide
                # whether a secondary variant really beats the barcode.
                top_level = (concordance_level(max(hits, key=lambda h: h["bit"]), qtax)
                             if hits else "none")
                candidates.setdefault(info["sample"], []).append({
                    "file": label, "header": header, "seq": seq, "info": info,
                    "n_hits": len(hits), "n_hits_raw": raw_counts.get(qkey, 0),
                    "qtax": qtax, "best_hit": best_hit, "level": level,
                    "top_level": top_level,
                })
            self._emit_progress(0.5 * (i + 1) / n_files)

        if self._stop:
            self.statusUpdated.emit("result", "Stopped     │ selection cancelled")
            self.taskFinished.emit(output_dir)
            return
        if not candidates:
            self.taskError.emit("No FASTA sequences found in the provided files.")
            return

        self.statusUpdated.emit(
            "files",
            f"Loaded      │ {n_files} runs · {total_seqs} sequences · "
            f"{len(candidates)} samples"
        )
        if ref_path:
            _empty = reference_empty_ids(list(candidates), ref_table)
            if _empty:
                ref_lines += ["", f"  Samples in the reference with empty taxonomy: {len(_empty)}"]
                ref_lines += [f"    {x}" for x in _empty[:50]]
                self.statusUpdated.emit(
                    "files", f"Warning     │ {len(_empty)} sample(s) with empty taxonomy "
                             f"in the reference (e.g. {', '.join(_empty[:3])})")

        # ── Output files ──
        tsv_path = os.path.join(output_dir, f"bestseq-{mydate}.tsv")
        fa_all   = os.path.join(output_dir, f"bestseq-{mydate}_all.fasta")
        fa_id    = os.path.join(output_dir, f"bestseq-{mydate}_identified.fasta")
        fa_nohit = os.path.join(output_dir, f"bestseq-{mydate}_no_blast_hit.fasta")
        fa_mism  = os.path.join(output_dir, f"bestseq-{mydate}_tax_mismatch.fasta")

        n_samples   = len(candidates)
        n_identical = 0
        n_selected  = 0
        n_single    = 0
        n_one_run   = 0
        n_written   = {"all": 0, "identified": 0, "no_blast_hit": 0,
                       "below_min_aln": 0, "tax_mismatch": 0, "no_query_tax": 0}
        n_variant_overruled = 0   # variants that won on score with no taxonomic gain
        level_counts = {lvl: 0 for lvl in TAX_LEVELS}
        flag_counts: Dict[str, int] = {}

        try:
            fh_tsv   = open(tsv_path, "w", encoding="utf-8")
            fh_all   = open(fa_all,   "w", encoding="utf-8")
            fh_id    = open(fa_id,    "w", encoding="utf-8")
            fh_nohit = open(fa_nohit, "w", encoding="utf-8")
            fh_mism  = open(fa_mism,  "w", encoding="utf-8")
        except PermissionError as e:
            self.taskError.emit(f"Could not write output files (locked/permission denied):\n{e}")
            return

        try:
            fh_tsv.write("\t".join(self._COLUMNS) + "\n")

            for done, sample in enumerate(sorted(candidates), start=1):
                if self._stop:
                    break
                cands = candidates[sample]
                identical = len({c["seq"].upper() for c in cands}) == 1

                max_bit = max((c["best_hit"]["bit"] if c["best_hit"] else 0.0)
                              for c in cands) or 1.0
                for cand in cands:
                    bit = cand["best_hit"]["bit"] if cand["best_hit"] else 0.0
                    cand["score"] = (self.TAX_BONUS[cand["level"]]
                                     + bit_weight * bit / max_bit
                                     - amb_penalty * (cand["info"]["ambs"] or 0)
                                     - gap_penalty * (cand["info"]["estgaps"] or 0))

                # deterministic ordering: score, then longer, then more reads, then file
                cands.sort(key=lambda c: (-c["score"],
                                          -(c["info"]["length"] or 0),
                                          -(c["info"]["reads"] or 0),
                                          c["file"]))
                # A secondary variant only beats the barcode with a real
                # taxonomic gain: its top hit must reach a deeper level than
                # the top hit of the best barcode candidate. Otherwise both
                # point to the same taxon and the score gap is noise.
                if is_variant_header(cands[0]["header"]):
                    bc = next((c for c in cands if not is_variant_header(c["header"])),
                              None)
                    if bc is not None and (self.TAX_BONUS[cands[0]["top_level"]]
                                           <= self.TAX_BONUS[bc["top_level"]]):
                        cands.remove(bc)
                        cands.insert(0, bc)
                        n_variant_overruled += 1
                best      = cands[0]
                runner_up = cands[1] if len(cands) > 1 else None

                # A sample can have several candidates per file (its barcode
                # plus its secondary variants), so presence counts files.
                n_cand_files = len({c["file"] for c in cands})
                in_all_files = n_cand_files == n_files
                if len(cands) == 1 and n_files == 1:
                    # Nothing to compare: the run is being classified by its
                    # taxonomic match, not chosen against other runs.
                    decision = "single_run"
                    n_single += 1
                elif identical and n_files > 1 and n_cand_files == 1:
                    # Recovered in one run only: nothing to be identical to.
                    decision = "only_one_run"
                    n_one_run += 1
                elif identical:
                    decision = ("identical_in_all_runs" if in_all_files
                                else "identical_in_available_runs")
                    n_identical += 1
                else:
                    decision = "resolved_by_score"
                    n_selected += 1
                level_counts[best["level"]] = level_counts.get(best["level"], 0) + 1

                # Flags say WHY the taxonomy is (or is not) supported, and are
                # deliberately distinct: no hit at all, hits that exist but were
                # all too short, hits that contradict the expected taxonomy
                # (a contamination / mislabelling candidate), and a hit that only
                # reaches order level (weak but consistent).
                flags = []
                if best["n_hits"] == 0:
                    flags.append("hits_below_min_aln" if best["n_hits_raw"]
                                 else "no_blast_hit")
                elif not any(best["qtax"].values()):
                    # Nothing to compare the hits with (e.g. the sample is
                    # missing from the reference file): not a mismatch.
                    flags.append("no_query_taxonomy")
                elif best["level"] == "none":
                    flags.append("tax_mismatch")
                elif best["level"] == "order":
                    flags.append("low_taxonomic_support")
                if (not identical and runner_up is not None
                        and abs(best["score"] - runner_up["score"]) < 1.0):
                    flags.append("near_tie")
                if not in_all_files and n_files > 1:
                    flags.append("missing_in_%d_run(s)" % (n_files - n_cand_files))
                if is_variant_header(best["header"]):
                    # The secondary variant beat the barcode (e.g. the dominant
                    # haplotype was a contaminant): worth a manual look.
                    flags.append("secondary_variant_selected")
                if best["info"]["ambs"]:
                    flags.append("ambs=%d" % best["info"]["ambs"])
                for f in flags:
                    key = f.split("=")[0]
                    flag_counts[key] = flag_counts.get(key, 0) + 1

                # The expected taxonomy comes from the BLAST table itself, so it
                # is reported even when every hit of this query was filtered out.
                hit  = best["best_hit"]
                qtax = best["qtax"]

                fh_tsv.write("\t".join(str(x) for x in [
                    sample, decision, len(cands), best["file"], best["header"],
                    best["info"]["length"] if best["info"]["length"] is not None else "",
                    best["info"]["reads"] if best["info"]["reads"] is not None else "",
                    best["info"]["ambs"], best["info"]["estgaps"],
                    best["n_hits"], best["n_hits_raw"], best["level"],
                    qtax["order"], qtax["family"], qtax["genus"], qtax["organism"],
                    hit["order"] if hit else "",
                    hit["family"] if hit else "",
                    hit["genus"] if hit else "",
                    hit["organism"] if hit else "",
                    hit["acc"] if hit else "",
                    round(hit["pident"], 3) if hit else "",
                    int(hit["alen"]) if hit else "",
                    round(hit["bit"], 1) if hit else "",
                    round(best["score"], 2),
                    runner_up["file"] if runner_up and not identical else "",
                    round(runner_up["score"], 2) if runner_up and not identical else "",
                    ";".join(flags),
                ]) + "\n")

                # Original header kept verbatim, sequence on a single line.
                # The two subsets split strictly on taxonomic concordance: a
                # sequence identical in every run is a reproducible consensus,
                # not a verified identification, so it only reaches
                # '_identified' if its hits actually match the expected taxonomy.
                # Not identified is split by WHY: BLAST found nothing similar
                # (no hit / only short hits) vs. similar sequences of another
                # taxon (contamination / mislabelling candidate).
                record = ">%s\n%s\n" % (best["header"], best["seq"])
                fh_all.write(record)
                n_written["all"] += 1
                if best["level"] != "none":
                    fh_id.write(record)
                    n_written["identified"] += 1
                elif best["n_hits"] == 0:
                    fh_nohit.write(record)
                    n_written["no_blast_hit"] += 1
                    if best["n_hits_raw"]:
                        n_written["below_min_aln"] += 1
                else:
                    # Includes hits with no expected taxonomy to compare with
                    # (flag no_query_taxonomy): not identified, for review.
                    fh_mism.write(record)
                    n_written["tax_mismatch"] += 1
                    if not any(best["qtax"].values()):
                        n_written["no_query_tax"] += 1

                self._emit_progress(0.5 + 0.5 * done / n_samples)
                if done % 25 == 0 or done == n_samples:
                    self.statusUpdated.emit(
                        "select",
                        f"Classifying │ {done}/{n_samples} sequences"
                        if n_files == 1 else
                        f"Selecting   │ {done}/{n_samples} samples · "
                        f"{n_identical} identical · {n_selected} decided by BLAST"
                    )
        finally:
            for fh in (fh_tsv, fh_all, fh_id, fh_nohit, fh_mism):
                try:
                    fh.close()
                except Exception:
                    pass

        # A FASTA that received no sequence is not left behind as an empty file.
        fa_counts = {fa_all: "all", fa_id: "identified",
                     fa_nohit: "no_blast_hit", fa_mism: "tax_mismatch"}
        for fa_path, key in fa_counts.items():
            if n_written[key] == 0:
                try:
                    os.remove(fa_path)
                except OSError:
                    pass

        def _fa_line(fa_path, key):
            n = n_written[key]
            return (f"{os.path.basename(fa_path)}  ({n} seqs)" if n
                    else "none (0 seqs, file not written)")

        if n_files == 1:
            self.statusUpdated.emit(
                "select",
                f"Classified  │ {n_single}/{n_samples} sequences by taxonomic match"
            )
        else:
            self.statusUpdated.emit(
                "identical",
                f"Identical   │ {n_identical} sample(s) identical across runs"
                + (f" · {n_one_run} found in one run only" if n_one_run else "")
            )
            self.statusUpdated.emit(
                "select",
                f"Selected    │ {n_identical + n_selected + n_one_run}/{n_samples} "
                f"samples · {n_selected} decided by BLAST"
            )

        xlsx_path = self._tsv_to_xlsx(tsv_path)

        elapsed = datetime.datetime.now() - run_start
        elapsed_str = str(elapsed).split(".")[0]
        status_str  = "Stopped" if self._stop else "Completed"
        kept = "sequences" if n_files == 1 else "best sequences"
        result_msg = (
            f"{status_str}   │ {n_written['all']} {kept} · "
            f"{n_written['identified']} identified · "
            f"{n_written['no_blast_hit']} no BLAST hit · "
            f"{n_written['tax_mismatch']} taxonomic mismatch"
            + (f" ({n_written['no_query_tax']} without expected taxonomy)"
               if n_written["no_query_tax"] else "")
        )
        self.statusUpdated.emit("result", result_msg)

        # What deserves a manual look, most important first.
        review = []
        for key in ("tax_mismatch", "no_query_taxonomy", "no_blast_hit",
                    "hits_below_min_aln", "secondary_variant_selected",
                    "near_tie", "low_taxonomic_support", "missing_in"):
            n = sum(v for k, v in flag_counts.items() if k.startswith(key))
            if n:
                review.append(f"{n} {key}")
        self.statusUpdated.emit(
            "review",
            "Review      │ " + (" · ".join(review) if review else "nothing flagged")
            + ("  (see the Flag column)" if review else ""))

        # ── Run log ──
        log_lines = [
            "Best Sequence Selection Log",
            "=" * 60,
            f"Date/Time  : {run_start.strftime('%Y-%m-%d %H:%M:%S')}",
            f"Status     : {status_str}",
            f"Total time : {elapsed_str}",
            "",
            "Input runs (FASTA + BLAST table):",
        ]
        for pair in self.pairs:
            log_lines.append(f"  {os.path.abspath(pair['fasta'])}")
            log_lines.append(f"    BLAST: {os.path.abspath(pair['blast'])}")
        log_lines += [
            "",
            "Parameters:",
            f"  Minimum alignment length : {cfg.get('min_alignment', 100)} bp",
            f"  Bit-score weight         : {bit_weight:g}",
            f"  Penalty per ambiguity    : {amb_penalty:g}",
            f"  Penalty per estimated gap: {gap_penalty:g}",
            f"  Sample ID suffix removed : {cfg.get('strip_suffix', '') or '(none)'}",
            f"  Query taxonomy reference : "
            f"{os.path.basename(ref_path) if ref_path else '(none - read from the tables)'}",
            *ref_lines,
            *(["", "  Repeated FASTA headers (only the first record of each was used):"]
              + dup_lines if dup_lines else []),
            *(["", "  FASTA headers missing from the BLAST table (those got no hits):"]
              + nomatch_lines if nomatch_lines else []),
            "",
            "Scoring:",
            "  score = taxonomic bonus of the best concordant hit",
            "          (species 400 / genus 300 / family 200 / order 100 / none 0)",
            f"        + {bit_weight:g} x (bit score of that hit, normalised within the sample)",
            f"        - {amb_penalty:g} x ambs  -  {gap_penalty:g} x estgaps",
            "  Ties are broken by longer sequence, then more reads, then file name.",
            "  A secondary variant replaces the barcode only if its top hit reaches a",
            "  deeper taxonomic level than the barcode's top hit.",
            "",
            "Results:",
            f"  Samples                  : {n_samples}",
            (f"  Sequences classified     : {n_single}" if n_files == 1 else
             f"  Identical across runs    : {n_identical}"),
            *([] if n_files == 1 else
              [f"  Found in one run only    : {n_one_run}",
               f"  Resolved by score        : {n_selected}"]),
            "",
            f"  Variants overruled       : {n_variant_overruled}"
            "  (won on score, same taxonomic level as the barcode)",
            "",
            "  Classification of the selected sequences:",
            f"    identified          : {n_written['identified']}",
            f"    no BLAST hit        : {n_written['no_blast_hit']}"
            f"  (no hit at all: {n_written['no_blast_hit'] - n_written['below_min_aln']}"
            f" · only hits below the minimum alignment: {n_written['below_min_aln']})",
            f"    taxonomic mismatch  : {n_written['tax_mismatch']}"
            + (f"  (of which {n_written['no_query_tax']} without expected taxonomy:"
               f" flag no_query_taxonomy)" if n_written["no_query_tax"] else ""),
            "",
            "  Taxonomic level reached by the selected sequence:",
        ]
        for lvl in TAX_LEVELS:
            log_lines.append(f"    {lvl:<10}: {level_counts.get(lvl, 0)}")
        log_lines += ["", "  Flags:"]
        if flag_counts:
            for key in sorted(flag_counts):
                # missing_in_N_run(s) carries the run count in the key itself
                help_key = "missing_in" if key.startswith("missing_in") else key
                log_lines.append(
                    f"    {key:<22}: {flag_counts[key]:<4} {_FLAG_HELP.get(help_key, '')}")
        else:
            log_lines.append("    (none)")
        log_lines += [
            "",
            "Output files:",
            f"  Folder                   : {output_dir}",
            f"  Report TSV               : {os.path.basename(tsv_path)}",
            f"  Report XLSX              : {os.path.basename(xlsx_path) if xlsx_path else 'N/A'}",
            f"  All best sequences       : {_fa_line(fa_all, 'all')}",
            f"  Taxonomically identified : {_fa_line(fa_id, 'identified')}",
            f"  No BLAST hit             : {_fa_line(fa_nohit, 'no_blast_hit')}",
            f"  Taxonomic mismatch       : {_fa_line(fa_mism, 'tax_mismatch')}",
            "",
            "NOTE: the subsets split strictly on taxonomic concordance and together add",
            "      up to the '_all' file. '_identified' holds the sequences whose best hit",
            "      matched the expected taxonomy at some rank. '_no_blast_hit' holds those",
            "      for which BLAST found nothing similar in the database (no hit, or only",
            "      short local hits below the minimum alignment): an unknown or poorly",
            "      represented taxon, or a non-target / low-quality sequence.",
            "      '_tax_mismatch' holds those with good hits of another taxon, and those",
            "      with hits but no expected taxonomy to compare with (flag",
            "      no_query_taxonomy: complete the reference). A sequence",
            "      identical in every run is a reproducible consensus, not a verified",
            "      identification: without a taxonomic hit it is split like any other.",
            "      Use the 'Decision' column of the report to tell those apart from the",
            "      ones that also disagreed between runs.",
            "",
            "      Review 'tax_mismatch' first: those samples have good BLAST hits that",
            "      contradict the expected taxonomy, which is what a contamination, an",
            "      index hop or a mislabelled voucher looks like. Compare the Query_* and",
            "      Hit_* columns of the report to see at which rank the match breaks.",
            "",
        ]
        log_path = os.path.join(output_dir, f"bestseq_run_log_{mydate}.txt")
        try:
            with open(log_path, "w", encoding="utf-8") as lf:
                lf.write("\n".join(log_lines))
        except Exception as exc:
            self.statusUpdated.emit("result", f"{result_msg}  │  Log error: {exc}")

        self.taskFinished.emit(output_dir)

from __future__ import annotations
import os
import re
import datetime
import xlsxwriter
from collections import Counter
from typing import Optional
from PyQt5 import QtCore, QtGui, QtWidgets
from .shared import *
from .shared import _get_base_dir, _tr


# Excel caps a worksheet at 1,048,576 rows (one is the header).
_XLSX_MAX_DATA_ROWS = 1048575

# "Split by header field" asks for confirmation above this many files.
_SPLIT_FILE_WARN = 100


def header_field_value(header: str, sep: str, field) -> Optional[str]:
    """Value of one header field, or None when the header does not have it.

    `field` is a 1-based position (int) or the key of a "key=value" field
    (str), found wherever it sits in the header. A "key=value" field yields
    its value part. With a separator set, a header that does not contain it
    has no fields at all. Shared by Filter, Split and their previews."""
    if sep and sep not in header:
        return None
    parts = header.split(sep) if sep else [header]
    if isinstance(field, str):
        for part in parts:
            key, eq, val = part.partition("=")
            if eq and key.strip() == field:
                return val.strip()
        return None
    if not 1 <= field <= len(parts):
        return None
    raw = parts[field - 1].strip()
    return raw.split("=", 1)[1].strip() if "=" in raw else raw


def header_keys(headers, sep: str) -> list:
    """Keys of the "key=value" fields found in the headers, in order."""
    keys = {}
    for h in headers:
        if sep and sep not in h:
            continue
        for part in (h.split(sep) if sep else [h]):
            key, eq, _v = part.partition("=")
            if eq and key.strip():
                keys[key.strip()] = None
    return list(keys)


def field_label(field) -> str:
    return f"field {field}" if isinstance(field, int) else f'key "{field}"'


def extract_id(header: str, rule: Optional[dict] = None) -> str:
    """Sequence ID of a FASTA header (without '>') under an ID rule:

      id_mode "before"           - the text before the first `id_sep`
                                   (";", "|", " " or any custom string);
                                   what the panel always sends
      id_mode "token" / no sep   - the header up to the first whitespace
      id_suffix                  - removed from the end of the ID if present
                                   (e.g. "_all.fa" for ONTbarcoder consensus)

    Shared by "Append info" and the exact-ID mode of "Extract by pattern", and
    by their live previews, so what the preview shows is what the run does."""
    rule = rule or {}
    header = header.strip()
    if rule.get("id_mode") == "before" and rule.get("id_sep"):
        sep = rule["id_sep"]
        sid = (header.split()[0] if header.split() else "") if sep == " " \
            else header.split(sep, 1)[0].strip()
    else:
        sid = header.split()[0] if header.split() else ""
    suffix = rule.get("id_suffix", "")
    if suffix and sid.endswith(suffix) and len(sid) > len(suffix):
        sid = sid[:-len(suffix)]
    return sid


class _FastaToolsWorker(QtCore.QThread):
    progress = QtCore.pyqtSignal(str)
    log_line = QtCore.pyqtSignal(str)
    # Not named "finished": that would shadow QThread's own finished() signal.
    done     = QtCore.pyqtSignal(list)
    error    = QtCore.pyqtSignal(str)

    def __init__(self, files: list, operation: str, params: dict, out_dir: str, parent=None):
        super().__init__(parent)
        self._files     = files
        self._operation = operation
        self._params    = params
        self._out_dir   = out_dir
        self._stop      = False

    def stop(self):
        self._stop = True

    @staticmethod
    def _read_fasta(path):
        import gzip
        opener = gzip.open if path.lower().endswith(".gz") else open
        records = []
        header = None
        seq_parts = []
        with opener(path, "rt", encoding="utf-8", errors="replace") as fh:
            for line in fh:
                line = line.rstrip("\r\n")
                if line.startswith(">"):
                    if header is not None:
                        records.append((header, "".join(seq_parts)))
                    header = line[1:]
                    seq_parts = []
                elif header is not None:
                    seq_parts.append(line)
            if header is not None:
                records.append((header, "".join(seq_parts)))
        return records

    @staticmethod
    def _write_fasta(records, path) -> bool:
        """Write *records* to *path*. With no records nothing is written (an
        empty FASTA is no use to anyone) and False is returned."""
        if not records:
            return False
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            for hdr, seq in records:
                fh.write(f">{hdr}\n{seq}\n")
        return True

    @staticmethod
    def _stem(path):
        name = os.path.basename(path)
        for ext in (".fasta.gz", ".fas.gz", ".fa.gz", ".fna.gz",
                    ".fasta", ".fas", ".fa", ".fna"):
            if name.lower().endswith(ext):
                return name[: -len(ext)]
        return os.path.splitext(name)[0]

    def _warn_repeats(self, name, records):
        """Log repeated headers of *name* (they make several operations, and the
        tools that read the result, treat different sequences as one)."""
        from .best_seq_panel import repeated_headers, repeated_note
        repeats = repeated_headers([h for h, _s in records])
        if repeats:
            self.log_line.emit(
                f"⚠ {name}: {repeated_note(repeats)}. Different sequences "
                f"with the same header are not told apart by the tools that "
                f"read this file.\n")

    def _iter_inputs(self):
        """Yield (display_name, stem, records) per input to process.
        In merge mode, concatenates all files and yields a single entry.
        When two files share the same stem, prefixes with the parent folder name."""
        if self._params.get("merge") and len(self._files) > 1:
            self.progress.emit("Merging files…")
            all_records = []
            for path in self._files:
                if self._stop:
                    return
                all_records.extend(self._read_fasta(path))
            self._warn_repeats(f"Merged {len(self._files)} files", all_records)
            yield f"Merged {len(self._files)} files", "merged", all_records
        else:
            stems = [self._stem(p) for p in self._files]
            dup_stems = {s for s in stems if stems.count(s) > 1}
            used_stems: set = set()
            for path, stem in zip(self._files, stems):
                if self._stop:
                    return
                name = os.path.basename(path)
                self.progress.emit(f"Processing {name}…")
                if stem in dup_stems:
                    parent = os.path.basename(os.path.dirname(os.path.abspath(path)))
                    stem = f"{parent}_{stem}"
                # Guarantee a unique output stem even when the parent folder also
                # matches, so a later file never silently overwrites an earlier one.
                base_stem = stem
                counter = 2
                while stem in used_stems:
                    stem = f"{base_stem}_{counter}"
                    counter += 1
                used_stems.add(stem)
                records = self._read_fasta(path)
                self._warn_repeats(name, records)
                yield name, stem, records

    @staticmethod
    def _cell_to_str(v) -> str:
        """Convert an Excel cell value to a string for ID matching.
        openpyxl returns integer-valued cells as float (e.g. 1234.0); turn those
        into "1234" so they match a "1234" FASTA id instead of becoming "1234_0"
        after the dot is normalized to an underscore."""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    @staticmethod
    def _normalize_field(s: str) -> str:
        import unicodedata, re as _re
        s = str(s).strip()
        s = unicodedata.normalize("NFD", s)
        s = "".join(c for c in s if unicodedata.category(c) != "Mn")
        s = _re.sub(r"[ .,]+", "_", s)
        s = _re.sub(r"_+", "_", s)
        return s.strip("_")

    def _run_unique(self):
        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            seen = set()
            unique = []
            for hdr, seq in records:
                key = seq.upper()
                if key not in seen:
                    seen.add(key)
                    unique.append((hdr, seq))
            removed = len(records) - len(unique)
            out_path = os.path.join(self._out_dir, f"{stem}_unique.fasta")
            wrote = self._write_fasta(unique, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Total sequences in :  {len(records)}\n"
                f"  Unique sequences   :  {len(unique)}\n"
                f"  Duplicates removed :  {removed}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    def _run_identical(self):
        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            counts = Counter(seq.upper() for _, seq in records)
            identical = [(hdr, seq) for hdr, seq in records if counts[seq.upper()] > 1]
            dup_groups = sum(1 for c in counts.values() if c > 1)

            # FASTA output
            out_path = os.path.join(self._out_dir, f"{stem}_identical.fasta")
            wrote = self._write_fasta(identical, out_path)

            # Excel: group IDs by sequence
            seq_to_ids: dict = {}
            for hdr, seq in records:
                key = seq.upper()
                if counts[key] > 1:
                    seq_to_ids.setdefault(key, []).append(hdr)

            xlsx_path = os.path.join(self._out_dir, f"{stem}_duplicate_groups.xlsx")
            wb = xlsxwriter.Workbook(xlsx_path)
            hf = wb.add_format({"bold": True, "bg_color": "#4F81BD",
                                "font_color": "#FFFFFF", "border": 1})
            cf = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
            nf = wb.add_format({"border": 1, "align": "center", "valign": "top"})

            ws_i = wb.add_worksheet("IDs")
            ws_i.activate()
            ws_g = wb.add_worksheet("Groups")

            for col, h in enumerate(["Group", "ID"]):
                ws_i.write(0, col, h, hf)
            ws_i.set_column(0, 0, 8)
            ws_i.set_column(1, 1, 30)

            for col, h in enumerate(["Group", "Count", "Sequence"]):
                ws_g.write(0, col, h, hf)
            ws_g.set_column(0, 1, 8)
            ws_g.set_column(2, 2, 50)

            if len(identical) > _XLSX_MAX_DATA_ROWS:
                self.log_line.emit(
                    f"  ⚠ {len(identical)} IDs exceed Excel's row limit: the 'IDs' sheet "
                    f"keeps the first {_XLSX_MAX_DATA_ROWS}; the FASTA has them all.")
            id_row = 1
            for g_idx, (seq_key, ids) in enumerate(seq_to_ids.items(), start=1):
                ws_g.write(g_idx, 0, g_idx, nf)
                ws_g.write(g_idx, 1, len(ids), nf)
                # Excel caps cell text at 32767 chars; truncate longer sequences
                # so xlsxwriter does not raise/drop the cell silently.
                if len(seq_key) > 32767:
                    ws_g.write(g_idx, 2, seq_key[:32764] + "...", cf)
                else:
                    ws_g.write(g_idx, 2, seq_key, cf)
                for id_val in ids:
                    ws_i.write(id_row, 0, g_idx, nf)
                    ws_i.write(id_row, 1, id_val, cf)
                    id_row += 1
            wb.close()
            outputs.append(xlsx_path)

            self.log_line.emit(
                f"{display_name}\n"
                f"  Total sequences in       :  {len(records)}\n"
                f"  Identical sequences out  :  {len(identical)}\n"
                f"  Distinct duplicated seqs :  {dup_groups}\n"
                f"  Excel groups file        :  {os.path.basename(xlsx_path)}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    def _run_grep(self):
        import re
        pattern_str  = self._params.get("pattern", "")
        pattern_file = self._params.get("pattern_file", "")
        use_regex    = self._params.get("use_regex", False)
        exact_rule   = self._params.get("exact_id")   # ID rule dict, or None
        ignore_case  = self._params.get("ignore_case", False)

        patterns = []
        if pattern_file:
            if not os.path.isfile(pattern_file):
                raise ValueError(f"Pattern file not found: {pattern_file}")
            with open(pattern_file, encoding="utf-8", errors="replace") as f:
                patterns = [ln.strip() for ln in f if ln.strip()]
            if not patterns:
                raise ValueError(f"Pattern file is empty: {os.path.basename(pattern_file)}")
        elif pattern_str:
            patterns = [pattern_str]

        if not patterns:
            raise ValueError("No pattern specified for grep operation.")

        mode_str = ("exact ID" if exact_rule is not None else
                    "regex" if use_regex else "plain text")
        if ignore_case:
            mode_str += ", ignore case"
        pat_summary = patterns[0] if len(patterns) == 1 else f"{len(patterns)} patterns from file"

        if exact_rule is not None:
            # Whole-ID match: "DNS-1" selects DNS-1 only, never DNS-10/DNS-11.
            wanted = {p.lower() for p in patterns} if ignore_case else set(patterns)
            def matches(hdr):
                sid = extract_id(hdr, exact_rule)
                return (sid.lower() if ignore_case else sid) in wanted
        elif use_regex:
            flags = re.IGNORECASE if ignore_case else 0
            compiled = [re.compile(p, flags) for p in patterns]
            def matches(hdr):
                return any(rx.search(hdr) for rx in compiled)
        elif ignore_case:
            lowered = [p.lower() for p in patterns]
            def matches(hdr):
                h = hdr.lower()
                return any(p in h for p in lowered)
        else:
            def matches(hdr):
                return any(p in hdr for p in patterns)

        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            matched = [(hdr, seq) for hdr, seq in records if matches(hdr)]
            unmatched = len(records) - len(matched)
            out_path = os.path.join(self._out_dir, f"{stem}_grep.fasta")
            wrote = self._write_fasta(matched, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Pattern ({mode_str})  :  {pat_summary}\n"
                f"  Total sequences in  :  {len(records)}\n"
                f"  Matched             :  {len(matched)}\n"
                f"  Not matched         :  {unmatched}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    def _run_append(self):
        import openpyxl
        excel_path = self._params.get("excel_file", "")
        separator  = self._params.get("separator", "|")
        id_rule    = self._params.get("id_rule") or {}

        if not excel_path or not os.path.isfile(excel_path):
            raise ValueError("No valid Excel file specified.")

        lookup = self.read_excel_lookup(excel_path)

        excel_name = os.path.basename(excel_path)
        n_excel_ids = len(lookup)

        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            updated_list = []
            n_updated = 0
            n_not_found = 0
            not_found_ex = []
            for hdr, seq in records:
                raw_id = extract_id(hdr, id_rule)
                seq_id = self._normalize_field(raw_id) if raw_id else ""
                fields = lookup.get(seq_id, []) if seq_id else []
                if fields:
                    new_hdr = hdr + separator + separator.join(fields)
                    n_updated += 1
                else:
                    new_hdr = hdr
                    n_not_found += 1
                    if len(not_found_ex) < 3:
                        not_found_ex.append(raw_id or "(empty)")
                updated_list.append((new_hdr, seq))
            out_path = os.path.join(self._out_dir, f"{stem}_append.fasta")
            wrote = self._write_fasta(updated_list, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Excel file          :  {excel_name}  ({n_excel_ids} IDs)\n"
                f"  Total sequences in  :  {len(records)}\n"
                f"  ID rule             :  {describe_id_rule(id_rule)}\n"
                f"  IDs updated         :  {n_updated}\n"
                f"  IDs not in Excel    :  {n_not_found}"
                + (f"  (e.g. {', '.join(not_found_ex)})" if not_found_ex else "")
                + "\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    @classmethod
    def read_excel_lookup(cls, excel_path: str) -> dict:
        """{normalized ID: [normalized fields]} from column 1 / columns 2+ of
        the first sheet. Read-only mode: only cell values are needed, and it
        streams large sheets instead of loading them whole."""
        import openpyxl
        if not excel_path.lower().endswith(".xlsx"):
            raise ValueError("Only .xlsx files are supported (save the sheet as "
                             "Excel Workbook .xlsx).")
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        try:
            ws = wb.active
            lookup: dict = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    continue
                row_id = cls._normalize_field(cls._cell_to_str(row[0]))
                lookup[row_id] = [
                    nv for v in row[1:]
                    if v is not None and cls._cell_to_str(v).strip()
                    for nv in (cls._normalize_field(cls._cell_to_str(v)),)
                    if any(c.isalnum() for c in nv)
                ]
            return lookup
        finally:
            wb.close()

    def _run_reformat(self):
        mode      = self._params.get("mode", "linearize")
        wrap_cols = int(self._params.get("wrap_cols", 80))

        if mode == "linearize":
            suffix = "_linearized"
            mode_desc = "multi-line → single line"
        else:
            def _wrap(seq, w=wrap_cols):
                return "\n".join(seq[i: i + w] for i in range(0, len(seq), w)) if seq else ""
            suffix = "_wrapped"
            mode_desc = f"single line → multi-line  ({wrap_cols} nt/row)"

        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            reformatted = records if mode == "linearize" else [(hdr, _wrap(seq)) for hdr, seq in records]
            out_path = os.path.join(self._out_dir, f"{stem}{suffix}.fasta")
            wrote = self._write_fasta(reformatted, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Mode                :  {mode_desc}\n"
                f"  Total sequences     :  {len(records)}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    # NCBI translation table for each genetic code offered in the panel.
    _GCODE_TABLE = {"Standard": 1, "Vertebrate mitochondrial": 2,
                    "Invertebrate mitochondrial": 5, "Plant plastid": 11}

    def _run_orf_trim(self):
        from .orf_trim_fasta import trim_records
        gname = self._params.get("genetic_code", "Invertebrate mitochondrial")
        table = self._GCODE_TABLE[gname]
        min_cov = float(self._params.get("min_coverage", 0.95))
        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            out, counts, warns = trim_records(records, table, min_cov)
            out_path = os.path.join(self._out_dir, f"{stem}_orf.fasta")
            wrote = self._write_fasta(out, out_path)
            lens = Counter(len(s) for _h, s in out).most_common(5)
            n_empty = len(records) - len(out)
            msg = (
                f"{display_name}\n"
                f"  Genetic code        :  {gname} (table {table})\n"
                f"  Total sequences     :  {len(records)}\n"
                f"  Trimmed to ORF      :  {counts['trimmed']}\n"
                f"  Unchanged (clean)   :  {counts['unchanged']}\n"
                f"  Left intact (ORF < {min_cov:.0%}):  {counts['short_orf']}\n"
                + (f"  Empty, dropped      :  {n_empty}\n" if n_empty else "")
                + f"  Most common lengths :  "
                + ", ".join(f"{ln} bp ×{n}" for ln, n in lens) + "\n"
            )
            if warns:
                shown = warns[:20]
                msg += "  Short ORF (possible NUMT / pseudogene / wrong code):\n"
                msg += "".join(f"    {w}\n" for w in shown)
                if len(warns) > len(shown):
                    msg += f"    … and {len(warns) - len(shown)} more\n"
            self.log_line.emit(msg)
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    def _run_sort(self):
        separator = self._params.get("separator", "|")
        levels    = self._params.get("levels", [{"field": 1, "order": "asc"}])

        sep_desc = f'"{separator}"' if separator else "none (whole header)"
        levels_desc = "  |  ".join(
            f"Field {lv['field']} {'↑' if lv['order'] == 'asc' else '↓'}"
            for lv in levels
        )

        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            working = list(records)
            for level in reversed(levels):
                idx     = level["field"] - 1
                reverse = level["order"] == "desc"

                def _key(rec, i=idx, sep=separator):
                    parts = rec[0].split(sep) if sep else [rec[0]]
                    val   = parts[i].strip() if i < len(parts) else ""
                    # Natural sort: split into (text, int) chunks so that
                    # "Sample_10" > "Sample_2" instead of "Sample_10" < "Sample_2".
                    return tuple(
                        int(c) if c.isdigit() else c.lower()
                        for c in re.split(r"(\d+)", val)
                    )

                working = sorted(working, key=_key, reverse=reverse)

            out_path = os.path.join(self._out_dir, f"{stem}_sorted.fasta")
            wrote = self._write_fasta(working, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Total sequences  :  {len(records)}\n"
                f"  Separator        :  {sep_desc}\n"
                f"  Sort levels      :  {levels_desc}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    def _run_filter_fields(self):
        separator = self._params.get("separator", "|")
        criteria  = self._params.get("criteria", [])
        if not criteria:
            raise ValueError("No filter criteria specified.")

        # Parse each criterion into (field_index, op, value, numeric).
        # Operators: =, !=, >=, <=, >, <, contains, not contains
        parsed = []
        for c in criteria:
            field = c.get("field", 1)     # 1-based position (int) or key (str)
            op  = c.get("op", "=")
            raw = c.get("value", "").strip()
            # Text operators must always compare as strings, even when the value
            # happens to parse as a number (e.g. "contains 5"). Otherwise the
            # numeric branch in _passes() has no case for contains/not contains
            # and the criterion is silently ignored.
            if op in ("contains", "not contains"):
                val     = raw.lower()
                numeric = False
            else:
                try:
                    val     = float(raw)
                    numeric = True
                except ValueError:
                    val     = raw.lower()
                    numeric = False
            parsed.append((field, op, val, numeric))

        def _ok(header, field, op, val, numeric):
            raw_v = header_field_value(header, separator, field)
            if raw_v is None:
                return False
            if numeric:
                try:
                    fv = float(raw_v)
                except ValueError:
                    return False
                return {">=": fv >= val, "<=": fv <= val, ">": fv > val,
                        "<": fv < val, "=": fv == val, "!=": fv != val}[op]
            sv = raw_v.lower()
            return {"=": sv == val, "!=": sv != val, "contains": val in sv,
                    "not contains": val not in sv}[op]

        match_any = self._params.get("match", "all") == "any"
        combine   = any if match_any else all

        def _passes(header):
            return combine(_ok(header, f, op, v, num) for f, op, v, num in parsed)

        sep_desc  = f'"{separator}"' if separator else "none (whole header)"
        crit_desc = ("  OR  " if match_any else "  AND  ").join(
            f"{field_label(f)} {o} {v}" for f, o, v, _ in parsed
        )
        crit_label = "Criteria (OR)" if match_any else "Criteria (AND)"
        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            matched   = [(h, s) for h, s in records if _passes(h)]
            unmatched = len(records) - len(matched)
            out_path  = os.path.join(self._out_dir, f"{stem}_filtered.fasta")
            wrote = self._write_fasta(matched, out_path)
            self.log_line.emit(
                f"{display_name}\n"
                f"  Separator        :  {sep_desc}\n"
                f"  {crit_label:<17}:  {crit_desc}\n"
                f"  Total sequences  :  {len(records)}\n"
                f"  Passed           :  {len(matched)}\n"
                f"  Excluded         :  {unmatched}\n"
            )
            if wrote:
                outputs.append(out_path)
            else:
                self.log_line.emit(
                    f"  No sequences: {os.path.basename(out_path)} not written\n")
        return outputs

    @staticmethod
    def _safe_filename(value: str) -> str:
        """A header field value made safe as part of a file name on Windows,
        macOS and Linux."""
        import re as _re
        # Keep letters, digits, "." and "-"; anything else (";", "=", spaces,
        # reserved characters...) becomes "_".
        s = _re.sub(r'[^\w.\-]+', "_", value.strip()).strip("._")
        return s[:80] or "empty"

    def _run_split(self):
        """Split each input into several FASTA files, written to <stem>_split/.

        Modes: "count" (N sequences per file), "files" (N files of near-equal
        size), "bases" (at most N bases per file; a longer single sequence
        still gets a file of its own — a sequence is never cut) and "field"
        (one file per distinct value of a header field, in order of first
        appearance). Every sequence ends up in exactly one file.
        """
        mode  = self._params.get("mode", "count")
        value = int(self._params.get("value", 1))
        separator = self._params.get("separator", "|")
        field = self._params.get("field", 1)    # position (int) or key (str)

        def _field_value(header):
            return header_field_value(header, separator, field)

        mode_desc = {
            "count": f"{value} sequences per file",
            "files": f"into {value} files",
            "bases": f"at most {value:,} bases per file",
            "field": f'one file per value of {field_label(field)} '
                     f'(separator "{separator}")',
        }[mode]

        outputs = []
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            if not records:
                self.log_line.emit(f"{display_name}\n  No sequences — nothing to split.\n")
                continue

            # ── Group the records (order kept) ──
            named_groups = []          # (file name suffix, records)
            if mode == "field":
                by_value: dict = {}
                for rec in records:
                    v = _field_value(rec[0])
                    key = "no_field" if v is None else self._safe_filename(v)
                    by_value.setdefault(key, []).append(rec)
                # Different values can map to the same safe name, and Windows
                # file names ignore case: keep every file name distinct.
                used: set = set()
                for key, recs in by_value.items():
                    name, n = key, 2
                    while name.lower() in used:
                        name, n = f"{key}_{n}", n + 1
                    used.add(name.lower())
                    named_groups.append((name, recs))
            else:
                if mode == "count":
                    groups = [records[i:i + value] for i in range(0, len(records), value)]
                elif mode == "files":
                    k = min(value, len(records))
                    size, extra = divmod(len(records), k)
                    groups, pos = [], 0
                    for i in range(k):
                        n = size + (1 if i < extra else 0)
                        groups.append(records[pos:pos + n])
                        pos += n
                else:  # bases
                    groups, cur, cur_bases = [], [], 0
                    for rec in records:
                        n = len(rec[1])
                        if cur and cur_bases + n > value:
                            groups.append(cur)
                            cur, cur_bases = [], 0
                        cur.append(rec)
                        cur_bases += n
                    if cur:
                        groups.append(cur)
                width = max(2, len(str(len(groups))))
                named_groups = [(f"part{i:0{width}d}", g) for i, g in enumerate(groups, 1)]

            # ── Write ──
            split_dir = os.path.join(self._out_dir, f"{stem}_split")
            for suffix, recs in named_groups:
                if self._stop:
                    break
                path = os.path.join(split_dir, f"{stem}_{suffix}.fasta")
                self._write_fasta(recs, path)
                outputs.append(path)

            sizes = [len(r) for _s, r in named_groups]
            lines = [
                display_name,
                f"  Split            :  {mode_desc}",
                f"  Total sequences  :  {len(records)}",
                f"  Files written    :  {len(named_groups)}  →  {os.path.basename(split_dir)}/",
                f"  Sequences / file :  {min(sizes)}–{max(sizes)}" if len(set(sizes)) > 1
                else f"  Sequences / file :  {sizes[0]}",
            ]
            if mode == "bases":
                bases = [sum(len(s) for _h, s in r) for _n, r in named_groups]
                lines.append(f"  Bases / file     :  {min(bases):,}–{max(bases):,}")
                over = sum(1 for b in bases if b > value)
                if over:
                    lines.append(f"  Note             :  {over} file(s) hold a single sequence "
                                 f"longer than {value:,} bases (sequences are never cut)")
            if mode == "field":
                missing = sum(len(r) for n, r in named_groups if n == "no_field")
                if missing:
                    lines.append(f"  Without {field_label(field)}  :  {missing} sequence(s) "
                                 f"→ {stem}_no_field.fasta")
                shown = named_groups[:20]
                lines.append("  Values           :  " + ", ".join(
                    f"{n} ({len(r)})" for n, r in shown)
                    + (f", … +{len(named_groups) - 20} more" if len(named_groups) > 20 else ""))
            self.log_line.emit("\n".join(lines) + "\n")
        return outputs

    # Stop codons per genetic code (NCBI tables 1, 2, 5 and 11; DNA alphabet)
    _STOPS = {
        "Standard":                   {"TAA", "TAG", "TGA"},
        "Vertebrate mitochondrial":   {"TAA", "TAG", "AGA", "AGG"},
        "Invertebrate mitochondrial": {"TAA", "TAG"},
        "Plant plastid":              {"TAA", "TAG", "TGA"},
    }
    _IUPAC = {
        "A": "A", "C": "C", "G": "G", "T": "T", "R": "AG", "Y": "CT",
        "S": "CG", "W": "AT", "K": "GT", "M": "AC", "B": "CGT", "D": "AGT",
        "H": "ACT", "V": "ACG", "N": "ACGT",
    }
    _COMPLEMENT = str.maketrans("ACGTRYSWKMBDHVN", "TGCAYRSWMKVHDBN")

    @classmethod
    def _is_stop(cls, codon: str, stops: set) -> bool:
        """True if every possible reading of the codon (IUPAC expanded) is a stop."""
        if codon in stops:
            return True
        opts = [cls._IUPAC.get(ch) for ch in codon]
        if None in opts or all(len(o) == 1 for o in opts):
            return False
        return all(a + b + c in stops for a in opts[0] for b in opts[1] for c in opts[2])

    @classmethod
    def _scan_stops(cls, seq: str, stops: set):
        """Count stop codons in all 6 frames (every codon, terminal included).
        Returns (best_frame, n_stops, positions) for the frame with fewest stops;
        frames are labelled +1..+3 (forward) and -1..-3 (reverse complement).
        Positions are 1-based nucleotide starts on the forward (ungapped) sequence."""
        L = len(seq)
        rc = seq.translate(cls._COMPLEMENT)[::-1]
        best = None
        for strand, s_ in (("+", seq), ("-", rc)):
            for f in range(3):
                pos = [i for i in range(f, L - 2, 3) if cls._is_stop(s_[i:i + 3], stops)]
                if strand == "-":
                    pos = [L - i - 2 for i in pos][::-1]
                else:
                    pos = [i + 1 for i in pos]
                if best is None or len(pos) < len(best[1]):
                    best = (f"{strand}{f + 1}", pos)
        return best[0], len(best[1]), best[1]

    @classmethod
    def _seq_stats(cls, seq: str, stops=None) -> dict:
        """Per-sequence composition. Gaps ('-', '.') are excluded from the length.
        Ambiguous = IUPAC codes other than A/C/G/T/U (N included).
        GC% is computed over unambiguous bases (A/C/G/T/U).
        Stop codons: all 6 frames are scanned codon by codon (terminal codon
        included); the frame with fewest stops is reported."""
        u = seq.upper().replace("U", "T")
        gaps = u.count("-") + u.count(".")
        length = len(u) - gaps
        a, c, g, t = u.count("A"), u.count("C"), u.count("G"), u.count("T")
        n = u.count("N")
        ambig = n + sum(u.count(ch) for ch in "RYKMSWBDHV")
        acgt = a + c + g + t
        invalid = max(length - acgt - ambig, 0)

        ungapped = u.replace("-", "").replace(".", "")
        homo = max((len(m.group()) for m in re.finditer(r"A+|C+|G+|T+", ungapped)), default=0)

        best_frame, best_stops, stop_pos = "", 0, []
        if stops is not None and len(ungapped) >= 3:
            best_frame, best_stops, stop_pos = cls._scan_stops(ungapped, stops)
        return {
            "length": length, "gaps": gaps, "n": n, "ambig": ambig,
            "invalid": invalid, "gc": c + g, "acgt": acgt,
            "a": a, "c": c, "g": g, "t": t, "homopolymer": homo,
            "frame": best_frame, "stops": best_stops, "stop_pos": stop_pos,
            "ambig_pct": 100.0 * ambig / length if length else 0.0,
            "gc_pct": 100.0 * (c + g) / acgt if acgt else 0.0,
        }

    def _run_stats(self):
        import statistics
        outputs = []
        gcode = self._params.get("genetic_code")
        stops = self._STOPS.get(gcode)
        for display_name, stem, records in self._iter_inputs():
            if self._stop:
                break
            rows = [(hdr, self._seq_stats(seq, stops)) for hdr, seq in records]
            lengths = [s["length"] for _, s in rows]
            n_seq = len(rows)
            total = sum(lengths)
            total_ambig = sum(s["ambig"] for _, s in rows)
            total_acgt = sum(s["acgt"] for _, s in rows)
            total_gc = sum(s["gc"] for _, s in rows)

            n50 = l50 = 0
            acc = 0
            for i, ln in enumerate(sorted(lengths, reverse=True), 1):
                acc += ln
                if acc * 2 >= total:
                    n50, l50 = ln, i
                    break

            dup_seqs = n_seq - len({seq.upper() for _, seq in records})
            summary = [
                ("File(s)",                          display_name),
                ("Number of sequences",              n_seq),
                ("Total bases",                      total),
                ("Minimum length",                   min(lengths) if lengths else 0),
                ("Maximum length",                   max(lengths) if lengths else 0),
                ("Mean length",                      round(total / n_seq, 2) if n_seq else 0),
                ("Median length",                    statistics.median(lengths) if lengths else 0),
                ("Std. deviation of length",         round(statistics.pstdev(lengths), 2) if lengths else 0),
                ("N50",                              n50),
                ("L50",                              l50),
                ("Ambiguous bases (IUPAC, incl. N)", total_ambig),
                ("Ambiguous bases (%)",              round(100.0 * total_ambig / total, 4) if total else 0),
                ("  of which N",                     sum(s["n"] for _, s in rows)),
                ("Sequences with ambiguous bases",   sum(1 for _, s in rows if s["ambig"])),
                ("GC content (%)",                   round(100.0 * total_gc / total_acgt, 2) if total_acgt else 0),
                ("Gap characters ('-' '.')",         sum(s["gaps"] for _, s in rows)),
                ("Invalid characters",               sum(s["invalid"] for _, s in rows)),
                ("Empty sequences",                  sum(1 for ln in lengths if ln == 0)),
                ("Duplicated sequences (extra copies)", dup_seqs),
            ]
            for base in "acgt":
                summary.append((f"{base.upper()} (%)",
                                round(100.0 * sum(s[base] for _, s in rows) / total_acgt, 2)
                                if total_acgt else 0))
            homos = [s["homopolymer"] for _, s in rows]
            summary.append(("Longest homopolymer (max)", max(homos) if homos else 0))
            summary.append(("Longest homopolymer (mean)",
                            round(sum(homos) / n_seq, 2) if n_seq else 0))
            if stops is not None:
                summary.append(("Genetic code", gcode))
                summary.append(("Sequences with stop codons (best of 6 frames)",
                                sum(1 for _, s in rows if s["stops"])))

            xlsx_path = os.path.join(self._out_dir, f"{stem}_stats.xlsx")
            wb = xlsxwriter.Workbook(xlsx_path)
            hf = wb.add_format({"bold": True, "bg_color": "#4F81BD",
                                "font_color": "#FFFFFF", "border": 1})
            cf = wb.add_format({"border": 1, "valign": "top"})
            nf = wb.add_format({"border": 1, "align": "center", "valign": "top"})
            pf = wb.add_format({"border": 1, "align": "center", "valign": "top",
                                "num_format": "0.00"})

            ws_s = wb.add_worksheet("Summary")
            ws_s.activate()
            ws_s.write(0, 0, "Metric", hf)
            ws_s.write(0, 1, "Value", hf)
            ws_s.set_column(0, 0, 38)
            ws_s.set_column(1, 1, 24)
            for r, (k, v) in enumerate(summary, start=1):
                ws_s.write(r, 0, k, cf)
                ws_s.write(r, 1, v, nf)

            if n_seq > _XLSX_MAX_DATA_ROWS:
                self.log_line.emit(
                    f"  ⚠ {n_seq} sequences exceed Excel's row limit: the 'Sequences' "
                    f"sheet keeps the first {_XLSX_MAX_DATA_ROWS} (Summary covers all).")
            ws_q = wb.add_worksheet("Sequences")
            cols = ["ID", "Length", "Ambiguous", "Ambiguous (%)", "N",
                    "GC (%)", "A (%)", "C (%)", "G (%)", "T (%)",
                    "Longest homopolymer", "Gaps", "Invalid"]
            if stops is not None:
                cols += ["Best frame", "Stop codons", "Stop positions (nt)"]
            for c, h in enumerate(cols):
                ws_q.write(0, c, h, hf)
            ws_q.set_column(0, 0, 40)
            ws_q.set_column(1, len(cols) - 1, 14)
            for r, (hdr, s) in enumerate(rows, start=1):
                acgt = s["acgt"]
                comp = [100.0 * s[b] / acgt if acgt else 0.0 for b in "acgt"]
                vals = [(s["length"], nf), (s["ambig"], nf), (s["ambig_pct"], pf),
                        (s["n"], nf), (s["gc_pct"], pf)]
                vals += [(v, pf) for v in comp]
                vals += [(s["homopolymer"], nf), (s["gaps"], nf), (s["invalid"], nf)]
                if stops is not None:
                    vals += [(s["frame"], nf), (s["stops"], nf),
                             (", ".join(map(str, s["stop_pos"])), cf)]
                ws_q.write(r, 0, hdr, cf)
                for c, (v, fmt) in enumerate(vals, start=1):
                    ws_q.write(r, c, v, fmt)
            ws_q.freeze_panes(1, 1)
            ws_q.autofilter(0, 0, max(n_seq, 1), len(cols) - 1)

            # Length histogram with a "nice" bin width (~20 bins)
            ws_h = wb.add_worksheet("Length histogram")
            ws_h.write(0, 0, "Bin start", hf)
            ws_h.write(0, 1, "Bin end", hf)
            ws_h.write(0, 2, "Sequences", hf)
            ws_h.set_column(0, 2, 14)
            if lengths:
                lo, hi = min(lengths), max(lengths)
                raw = max((hi - lo) / 20.0, 1)
                width = next(w for m in (1, 10, 100, 1000, 10000, 100000, 10 ** 9)
                             for w in (m, 2 * m, 5 * m) if w >= raw)
                start = (lo // width) * width
                n_bins = (hi - start) // width + 1
                counts = Counter((ln - start) // width for ln in lengths)
                for b in range(n_bins):
                    ws_h.write(b + 1, 0, start + b * width, nf)
                    ws_h.write(b + 1, 1, start + (b + 1) * width - 1, nf)
                    ws_h.write(b + 1, 2, counts.get(b, 0), nf)
                chart = wb.add_chart({"type": "column"})
                chart.add_series({
                    "name":       "Sequences",
                    "categories": ["Length histogram", 1, 0, n_bins, 0],
                    "values":     ["Length histogram", 1, 2, n_bins, 2],
                    "gap":        10,
                })
                chart.set_title({"name": f"Length distribution (bin = {width} bp)"})
                chart.set_x_axis({"name": "Length (bin start, bp)"})
                chart.set_y_axis({"name": "Sequences"})
                chart.set_legend({"none": True})
                chart.set_size({"width": 760, "height": 400})
                ws_h.insert_chart(1, 4, chart)
            wb.close()
            outputs.append(xlsx_path)

            self.log_line.emit(
                f"{display_name}\n"
                f"  Sequences          :  {n_seq}\n"
                f"  Length min/max/mean:  {summary[3][1]} / {summary[4][1]} / {summary[5][1]}\n"
                f"  Total bases        :  {total}\n"
                f"  N50                :  {n50}\n"
                f"  Ambiguous bases    :  {total_ambig}  ({summary[11][1]}%)\n"
                f"  GC content         :  {summary[14][1]}%\n"
                f"  Longest homopolymer:  {max(homos) if homos else 0}\n"
                + (f"  With stop codons   :  {summary[-1][1]} sequence(s)\n"
                   if stops is not None else "")
                + f"  Stats file         :  {os.path.basename(xlsx_path)}\n"
            )
        return outputs

    def run(self):
        try:
            os.makedirs(self._out_dir, exist_ok=True)
            if self._operation == "stats":
                outputs = self._run_stats()
            elif self._operation == "unique":
                outputs = self._run_unique()
            elif self._operation == "identical":
                outputs = self._run_identical()
            elif self._operation == "grep":
                outputs = self._run_grep()
            elif self._operation == "append":
                outputs = self._run_append()
            elif self._operation == "reformat":
                outputs = self._run_reformat()
            elif self._operation == "orf_trim":
                outputs = self._run_orf_trim()
            elif self._operation == "sort":
                outputs = self._run_sort()
            elif self._operation == "filter_fields":
                outputs = self._run_filter_fields()
            elif self._operation == "split":
                outputs = self._run_split()
            else:
                raise ValueError(f"Unknown operation: {self._operation}")
            if not self._stop:
                self.done.emit(outputs)
        except Exception as exc:
            import traceback
            # ValueError = a message written for the user; anything else is
            # unexpected and carries the traceback for the details dialog.
            msg = str(exc) if isinstance(exc, ValueError) else \
                f"{type(exc).__name__}: {exc}\n{traceback.format_exc()}"
            self.error.emit(msg)


class _DragDropLineEdit(QtWidgets.QLineEdit):
    """QLineEdit that accepts a single file dragged onto it."""

    _DRAG_STYLE = (
        f"QLineEdit {{ border: {DROP_DRAG_BORDER}; background-color: {DROP_DRAG_BG};"
        f" border-radius: 4px; }}"
    )

    def __init__(self, accepted_extensions=None, parent=None):
        super().__init__(parent)
        self._accepted_ext = [e.lower() for e in (accepted_extensions or [])]
        self.setAcceptDrops(True)

    def _is_valid(self, url):
        if not url.isLocalFile():
            return False
        if self._accepted_ext:
            return any(url.toLocalFile().lower().endswith(e) for e in self._accepted_ext)
        return True

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and any(self._is_valid(u) for u in event.mimeData().urls()):
            event.acceptProposedAction()
            self.setStyleSheet(self._DRAG_STYLE)
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls() and any(self._is_valid(u) for u in event.mimeData().urls()):
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self.setStyleSheet("")
        super().dragLeaveEvent(event)

    def dropEvent(self, event):
        self.setStyleSheet("")
        for url in event.mimeData().urls():
            if self._is_valid(url):
                self.setText(url.toLocalFile())
                event.acceptProposedAction()
                return
        event.ignore()


def describe_id_rule(rule: Optional[dict]) -> str:
    rule = rule or {}
    if rule.get("id_mode") == "before" and rule.get("id_sep"):
        sep = "space" if rule["id_sep"] == " " else f'"{rule["id_sep"]}"'
        text = f"text before first {sep}"
    else:
        text = "header up to first space"
    if rule.get("id_suffix"):
        text += f', minus suffix "{rule["id_suffix"]}"'
    return text


class _IdRuleWidget(QtWidgets.QWidget):
    """How the sequence ID is taken from each FASTA header, plus a live
    preview line the panel fills in. Used by Append info and exact-ID grep."""
    changed = QtCore.pyqtSignal()

    _SEPS = [("|", '"|"  pipe'), (";", '";"  semicolon'), (",", '","  comma'),
             (" ", "␣  space"), ("	", "⇥  tab"), ("_", '"_"  underscore'),
             ("-", '"-"  hyphen'), ("", "Custom…")]

    def __init__(self, parent=None, default_sep: str = ";"):
        super().__init__(parent)
        lay = QtWidgets.QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(6)

        row = QtWidgets.QHBoxLayout()
        row.setSpacing(8)
        row.addWidget(make_label("ID = header text before first:", color=TEXT_SEC))
        self._sep = QtWidgets.QComboBox()
        for value, label in self._SEPS:
            self._sep.addItem(label, value)
        # Index restored by reset().
        self._default_sep_idx = max(0, self._sep.findData(default_sep))
        self._sep.setCurrentIndex(self._default_sep_idx)
        row.addWidget(self._sep)
        self._custom = QtWidgets.QLineEdit()
        self._custom.setPlaceholderText("separator")
        self._custom.setFixedWidth(90)
        row.addWidget(self._custom)
        row.addSpacing(12)
        row.addWidget(make_label("Remove suffix:", color=TEXT_SEC))
        self._suffix = QtWidgets.QLineEdit()
        self._suffix.setPlaceholderText("optional, e.g. _all.fa")
        self._suffix.setFixedWidth(170)
        row.addWidget(self._suffix)
        row.addStretch()
        lay.addLayout(row)

        self._preview = QtWidgets.QLabel("")
        self._preview.setTextFormat(QtCore.Qt.RichText)
        self._preview.setWordWrap(True)
        lay.addWidget(self._preview)

        self._sep.currentIndexChanged.connect(self._on_changed)
        for edit in (self._custom, self._suffix):
            edit.textChanged.connect(self._on_changed)
        self._sync()

    def _sync(self):
        self._custom.setVisible(self._sep.currentData() == "")

    def _on_changed(self, *_):
        self._sync()
        self.changed.emit()

    def rule(self) -> dict:
        sep = self._sep.currentData()
        if sep == "":
            sep = self._custom.text()
        # An empty custom separator falls back to the first space.
        return {"id_mode": "before", "id_sep": sep,
                "id_suffix": self._suffix.text().strip()}

    def reset(self):
        for w in (self._sep, self._custom, self._suffix):
            w.blockSignals(True)
        self._sep.setCurrentIndex(self._default_sep_idx)
        self._custom.clear()
        self._suffix.clear()
        for w in (self._sep, self._custom, self._suffix):
            w.blockSignals(False)
        self._sync()
        self._preview.setText("")

    def set_preview(self, html: str, ok: Optional[bool] = None):
        color = TEXT_SEC if ok is None else (GREEN if ok else RED)
        self._preview.setStyleSheet(f"color:{color}; font-size:14px;")
        self._preview.setText(html)


class _SeparatorPicker(QtWidgets.QWidget):
    """Field-separator choice shared by Filter, Sort and Split: a drop-down of
    common separators (optionally "None (whole header)"), whose last entry,
    Custom, enables a text box."""
    changed = QtCore.pyqtSignal()
    _CUSTOM = "__custom__"

    def __init__(self, allow_none: bool = False, stretch: bool = True, parent=None):
        super().__init__(parent)
        self.row = QtWidgets.QHBoxLayout(self)
        self.row.setContentsMargins(0, 0, 0, 0)
        self.row.setSpacing(12)
        self.row.addWidget(make_label("Field separator:", color=TEXT_SEC))
        self._combo = QtWidgets.QComboBox()
        if allow_none:
            self._combo.addItem("None (whole header)", "")
        for label, sep in (('"|"  pipe', "|"), ('";"  semicolon', ";"),
                           ('","  comma', ","), ("␣  space", " "),
                           ("⇥  tab", "\t"), ('"_"  underscore', "_"),
                           ('"-"  hyphen', "-")):
            self._combo.addItem(label, sep)
        self._combo.addItem("Custom…", self._CUSTOM)
        self._combo.setMinimumWidth(190)
        self.row.addWidget(self._combo)
        self._edit = QtWidgets.QLineEdit()
        self._edit.setFixedWidth(64)
        self._edit.setPlaceholderText("e.g. :")
        self._edit.setMaxLength(5)
        # Enabled (Custom chosen): white with a blue border; disabled: greyed out.
        self._edit.setStyleSheet(
            f"QLineEdit {{ background: white; color: {TEXT_PRI};"
            f" border: 1.5px solid {BLUE}; border-radius: 4px; padding: 2px 4px; }}"
            f"QLineEdit:disabled {{ background: {GRAY_BG}; color: {TEXT_HINT};"
            f" border: 1px solid {GRAY_LINE}; }}")
        self.row.addWidget(self._edit)
        if stretch:
            self.row.addStretch()
        self._combo.currentIndexChanged.connect(self._on_choice)
        self._edit.textChanged.connect(self.changed)
        self.reset()

    def _is_custom(self) -> bool:
        return self._combo.currentData() == self._CUSTOM

    def _on_choice(self, _i):
        self._edit.setEnabled(self._is_custom())
        self.changed.emit()

    def separator(self) -> str:
        """The chosen separator, the custom text (may be empty) or "" for None."""
        return self._edit.text() if self._is_custom() else self._combo.currentData()

    def custom_empty(self) -> bool:
        return self._is_custom() and not self._edit.text()

    def reset(self):
        """First entry: "None" where allowed (Sort), otherwise the pipe."""
        self._combo.setCurrentIndex(0)
        self._edit.clear()
        self._edit.setEnabled(False)


class _FieldSelector(QtWidgets.QWidget):
    """Picks a header field by position or, for "key=value" fields, by key
    name (found wherever it sits in the header)."""
    changed = QtCore.pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        row = QtWidgets.QHBoxLayout(self)
        row.setContentsMargins(0, 0, 0, 0)
        row.setSpacing(6)
        self._mode = QtWidgets.QComboBox()
        self._mode.addItems(["Position", "Key name"])
        self._mode.setFixedWidth(175)
        self._mode.setToolTip(
            "Position: the Nth field after splitting by the separator.\n"
            "Key name: a \"key=value\" field (e.g. ambs=0), found wherever it\n"
            "sits in the header — robust when headers differ in structure.")
        row.addWidget(self._mode)
        self._spin = QtWidgets.QSpinBox()
        self._spin.setRange(1, 99)
        self._spin.setFixedWidth(70)
        row.addWidget(self._spin)
        self._key = QtWidgets.QComboBox()
        self._key.setEditable(True)
        self._key.setFixedWidth(150)
        self._key.lineEdit().setPlaceholderText("key, e.g. ambs")
        self._key.hide()
        row.addWidget(self._key)
        self._mode.currentIndexChanged.connect(self._on_mode)
        self._spin.valueChanged.connect(self.changed)
        self._key.currentTextChanged.connect(self.changed)

    def _on_mode(self, _i):
        by_key = self.by_key()
        self._spin.setVisible(not by_key)
        self._key.setVisible(by_key)
        self.changed.emit()

    def by_key(self) -> bool:
        return self._mode.currentIndex() == 1

    def value(self):
        """1-based position (int) or key name (str)."""
        return self._key.currentText().strip() if self.by_key() else self._spin.value()

    def set_position(self, n: int):
        self._spin.setValue(n)

    def set_max(self, n: int, quiet: bool = False):
        if quiet:
            self._spin.blockSignals(True)
        self._spin.setMaximum(max(1, n))
        if quiet:
            self._spin.blockSignals(False)

    def set_keys(self, keys: list):
        """Offer the keys found in the headers; keep what is already typed."""
        cur = self._key.currentText()
        self._key.blockSignals(True)
        self._key.clear()
        self._key.addItems(keys)
        self._key.setCurrentText(cur or (keys[0] if keys else ""))
        self._key.blockSignals(False)

    def reset(self):
        self._mode.setCurrentIndex(0)
        self._spin.setValue(1)
        self._key.setCurrentText("")


def _first_header(path: str) -> str:
    """First header of a FASTA file (without '>'), "" if none. Raises OSError."""
    import gzip as _gz
    opener = _gz.open if path.lower().endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            if line.startswith(">"):
                return line[1:].rstrip()
    return ""


class FastaToolsPanel(QtWidgets.QWidget):
    _FASTA_GCODES = list(_FastaToolsWorker._STOPS)


    def __init__(self, parent=None):
        super().__init__(parent)
        self._sort_field_count: int = 99

        outer = QtWidgets.QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        scroll = QtWidgets.QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QtWidgets.QFrame.NoFrame)
        inner = QtWidgets.QWidget()
        self._layout = QtWidgets.QVBoxLayout(inner)
        self._layout.setContentsMargins(20, 20, 20, 20)
        self._layout.setSpacing(16)
        scroll.setWidget(inner)
        outer.addWidget(scroll, 1)

        self._layout.addWidget(make_label("FASTA Tools", size=19, bold=True))
        desc = make_label(
            "Drag one or multiple FASTA files (.fa, .fas, .fasta) and choose an operation to run on them.",
            color=TEXT_SEC,
        )
        desc.setWordWrap(True)
        self._layout.addWidget(desc)

        self._drop = MultiDropZone()
        self._drop.filesDropped.connect(self._on_files)
        self._layout.addWidget(self._drop)

        self._merge_chk = QtWidgets.QCheckBox("Merge files before processing")
        self._merge_chk.setToolTip(
            "Checked: all files are concatenated into a single dataset\n"
            "and the operation runs on the combined sequences.\n"
            "Output is saved as a single merged file.\n\n"
            "Unchecked: each file is processed individually\n"
            "and a separate output file is generated for each one."
        )
        self._layout.addWidget(self._merge_chk)

        ops_box = QtWidgets.QGroupBox("Operation")
        ops_box.setStyleSheet(group_box_style())
        ops_layout = QtWidgets.QVBoxLayout(ops_box)
        ops_layout.setContentsMargins(16, 16, 16, 16)
        ops_layout.setSpacing(2)

        self._radio_stats     = QtWidgets.QRadioButton("FASTA stats")
        self._radio_stats.setToolTip(
            "Summary statistics and per-sequence table, saved as <name>_stats.xlsx\n"
            "  • Sheet 'Summary': number of sequences, total bases, min/max/mean/median\n"
            "    length, N50/L50, ambiguous bases, GC content, gaps, duplicates\n"
            "  • Sheet 'Sequences': length, ambiguous count and %, N, GC %, gaps per sequence\n\n"
            "Ambiguous = IUPAC codes other than A/C/G/T/U (N included).\n"
            "GC % is computed over unambiguous bases. Gaps ('-', '.') are not counted in length.\n"
            "Also reports A/C/G/T composition, longest homopolymer, a length histogram\n"
            "sheet with chart and, optionally, stop codons (best of 6 frames)."
        )
        self._radio_extract   = QtWidgets.QRadioButton("Extract sequences")
        self._radio_extract.setToolTip(
            "Extract a subset of sequences: unique ones, duplicated ones, or those\n"
            "whose header matches a pattern.")
        self._radio_unique    = QtWidgets.QRadioButton("Unique sequences")
        self._radio_identical = QtWidgets.QRadioButton("Identical sequences (duplicates)")
        self._radio_identical.setToolTip(
            "Extracts sequences that appear more than once.\n"
            "Generates two output files:\n"
            "  • <name>_identical.fasta — duplicated sequences\n"
            "  • <name>_duplicate_groups.xlsx — IDs grouped by shared sequence\n"
            "    · Sheet 'IDs': list of IDs per group\n"
            "    · Sheet 'Groups': group number, ID count and sequence"
        )
        self._radio_grep         = QtWidgets.QRadioButton("By pattern (grep)")
        self._radio_filter_fields = QtWidgets.QRadioButton("By header fields")
        self._radio_append       = QtWidgets.QRadioButton("Append info to headers from .xlsx file")
        self._radio_reformat     = QtWidgets.QRadioButton("Reformat sequence lines")
        self._radio_sort         = QtWidgets.QRadioButton("Sort sequences")
        self._radio_split        = QtWidgets.QRadioButton("Split FASTA file")
        self._radio_split.setToolTip(
            "Split into several FASTA files: by number of sequences, into a number\n"
            "of files, by total bases per file, or one file per header field value.")
        self._radio_orf          = QtWidgets.QRadioButton(
            "Trim to coding ORF (remove stop codon and 3′ tail)")
        self._radio_orf.setToolTip(
            "Trims each barcode to its longest clean open reading frame (6 frames),\n"
            "removing the gene's stop codon and the non-coding 3′ tail, as coding-marker\n"
            "mode does. Use it to make full-length non-coding barcodes (e.g. CytB with\n"
            "stop) comparable in length with coding-mode barcodes before Compare.\n"
            "Sequences that already translate cleanly are left unchanged; a header\n"
            "length field (>id;LEN;…) is updated. Writes <name>_orf.fasta."
        )

        self._op_group = QtWidgets.QButtonGroup(self)
        for rb in (self._radio_stats, self._radio_extract,
                   self._radio_append, self._radio_reformat, self._radio_orf,
                   self._radio_sort, self._radio_split):
            self._op_group.addButton(rb)

        # Sub-options of "Extract sequences"
        self._extract_group = QtWidgets.QButtonGroup(self)
        for rb in (self._radio_unique, self._radio_identical, self._radio_grep,
                   self._radio_filter_fields):
            self._extract_group.addButton(rb)
        self._radio_unique.setChecked(True)
        self._extract_group.buttonClicked.connect(self._on_operation_changed)

        self._radio_stats.setChecked(True)
        self._op_group.buttonClicked.connect(self._on_operation_changed)

        ops_layout.addWidget(self._radio_stats)
        self._stats_widget = QtWidgets.QWidget()
        stl = QtWidgets.QHBoxLayout(self._stats_widget)
        stl.setContentsMargins(20, 4, 0, 4)
        stl.setSpacing(8)
        self._stats_stops_chk = QtWidgets.QCheckBox("Check stop codons — genetic code:")
        self._stats_stops_chk.setToolTip(
            "Optional. Scans all 6 reading frames (forward and reverse complement),\n"
            "codon by codon including the last one, and reports the frame with fewest stops.\n"
            "Ambiguous codons count as stop only if every IUPAC reading is a stop (e.g. TAR).\n"
            "Stop codons in coding barcodes (COI, rbcL, matK…) suggest pseudogenes (NUMTs)\n"
            "or indel errors. Leave unchecked for non-coding markers or mixed-marker files."
        )
        stl.addWidget(self._stats_stops_chk)
        self._stats_gcode_combo = QtWidgets.QComboBox()
        for name in self._FASTA_GCODES:
            self._stats_gcode_combo.addItem(name)
        self._stats_gcode_combo.setCurrentText("Invertebrate mitochondrial")
        self._stats_gcode_combo.setEnabled(False)
        self._stats_stops_chk.toggled.connect(self._stats_gcode_combo.setEnabled)
        stl.addWidget(self._stats_gcode_combo)
        stl.addStretch(1)
        ops_layout.addWidget(self._stats_widget)
        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_extract)
        self._extract_widget = QtWidgets.QWidget()
        exl = QtWidgets.QVBoxLayout(self._extract_widget)
        exl.setContentsMargins(20, 4, 0, 4)
        exl.setSpacing(4)
        exl.addWidget(self._radio_unique)
        exl.addWidget(self._radio_identical)
        exl.addWidget(self._radio_grep)
        self._extract_widget.hide()
        ops_layout.addWidget(self._extract_widget)
        self._grep_widget = QtWidgets.QWidget()
        gl = QtWidgets.QVBoxLayout(self._grep_widget)
        gl.setContentsMargins(20, 4, 0, 4)
        gl.setSpacing(8)

        self._grep_use_file_chk = QtWidgets.QCheckBox("Load patterns from file")
        self._grep_use_file_chk.stateChanged.connect(self._on_grep_source_changed)
        gl.addWidget(self._grep_use_file_chk)

        self._grep_pattern_row = QtWidgets.QWidget()
        gpr = QtWidgets.QHBoxLayout(self._grep_pattern_row)
        gpr.setContentsMargins(0, 0, 0, 0)
        gpr.setSpacing(8)
        gpr.addWidget(make_label("Pattern:", color=TEXT_SEC))
        self._grep_pattern_edit = QtWidgets.QLineEdit()
        self._grep_pattern_edit.setPlaceholderText("Enter search pattern…")
        gpr.addWidget(self._grep_pattern_edit, 1)
        gl.addWidget(self._grep_pattern_row)

        self._grep_file_row = QtWidgets.QWidget()
        gfr = QtWidgets.QHBoxLayout(self._grep_file_row)
        gfr.setContentsMargins(0, 0, 0, 0)
        gfr.setSpacing(8)
        lbl_pf = make_label("Patterns file:", color=TEXT_SEC)
        lbl_pf.setToolTip(
            "Plain text file — one pattern per line.\n"
            "Empty lines are ignored.\n"
            "Each pattern is matched against the full FASTA header.\n"
            "Enable 'Regex' to treat each line as a regular expression."
        )
        gfr.addWidget(lbl_pf)
        self._grep_file_edit = _DragDropLineEdit(accepted_extensions=[".txt"])
        self._grep_file_edit.setReadOnly(True)
        self._grep_file_edit.setPlaceholderText("No file selected… (or drag & drop)")
        self._grep_file_edit.setToolTip(
            "Plain text file — one pattern per line.\n"
            "Empty lines are ignored.\n"
            "Each pattern is matched against the full FASTA header.\n"
            "Enable 'Regex' to treat each line as a regular expression."
        )
        gfr.addWidget(self._grep_file_edit, 1)
        self._grep_file_clear = self._make_clear_btn()
        self._grep_file_clear.setVisible(False)
        self._grep_file_clear.clicked.connect(lambda: (
            self._grep_file_edit.clear(),
            self._grep_file_clear.setVisible(False),
        ))
        self._grep_file_edit.textChanged.connect(
            lambda t: self._grep_file_clear.setVisible(bool(t)))
        gfr.addWidget(self._grep_file_clear)
        self._grep_file_btn = QtWidgets.QPushButton("Browse…")
        self._grep_file_btn.setObjectName("secondary_btn")
        self._grep_file_btn.setFixedWidth(120)
        self._grep_file_btn.clicked.connect(self._browse_grep_file)
        gfr.addWidget(self._grep_file_btn)
        self._grep_file_row.hide()
        gl.addWidget(self._grep_file_row)

        _grep_file_note = make_label(
            "Plain text file · one pattern per line · empty lines ignored",
            size=14, color=TEXT_HINT,
        )
        gl.addWidget(_grep_file_note)
        self._grep_file_note = _grep_file_note
        _grep_file_note.hide()

        self._grep_regex_chk = QtWidgets.QCheckBox("Use regex (regular expression)")
        gl.addWidget(self._grep_regex_chk)
        self._grep_case_chk = QtWidgets.QCheckBox("Ignore case")
        self._grep_case_chk.setToolTip(
            "Match upper and lower case alike (e.g. 'dns-1' finds DNS-1).\n"
            "Applies to plain text, regex and whole-ID matching.")
        gl.addWidget(self._grep_case_chk)
        self._grep_exact_chk = QtWidgets.QCheckBox(
            "Match whole ID (exact, not substring) — for lists of IDs")
        self._grep_exact_chk.setToolTip(
            "Each pattern must equal the sequence ID, so 'DNS-1' selects DNS-1\n"
            "only, not DNS-10 or DNS-11. Turned on automatically when patterns\n"
            "are loaded from a file.")
        gl.addWidget(self._grep_exact_chk)
        self._grep_id_rule = _IdRuleWidget()
        self._grep_id_rule.hide()
        gl.addWidget(self._grep_id_rule)
        self._grep_exact_chk.toggled.connect(self._on_grep_exact_toggled)
        self._grep_regex_chk.toggled.connect(self._on_grep_exact_toggled)
        self._grep_id_rule.changed.connect(self._update_grep_preview)
        self._grep_case_chk.toggled.connect(self._update_grep_preview)
        self._grep_pattern_edit.textChanged.connect(self._update_grep_preview)
        self._grep_file_edit.textChanged.connect(self._update_grep_preview)
        self._grep_widget.hide()
        exl.addWidget(self._grep_widget)

        exl.addWidget(self._radio_filter_fields)
        self._filter_widget = QtWidgets.QWidget()
        ffw = QtWidgets.QVBoxLayout(self._filter_widget)
        ffw.setContentsMargins(20, 4, 0, 4)
        ffw.setSpacing(10)

        ff_note = make_label(
            "Split each header by a separator and filter by field position or, for "
            "key=value fields, by key name. Combine the criteria with AND or OR. "
            "Numeric operators: >=  <=  >  <  =  !=   ·   Text operators (ignore case): "
            "=  !=  contains  not contains",
            color=TEXT_SEC, size=15,
        )
        ff_note.setWordWrap(True)
        ffw.addWidget(ff_note)

        # ── Separator row ────────────────────────────────────────────────────
        self._ff_sep = _SeparatorPicker()
        ffw.addWidget(self._ff_sep)
        self._ff_sep.changed.connect(self._update_ff_preview)

        # ── Field preview ────────────────────────────────────────────────────
        self._ff_preview_frame = QtWidgets.QFrame()
        self._ff_preview_frame.setObjectName("ff_preview_card")
        self._ff_preview_frame.setStyleSheet(f"""
            QFrame#ff_preview_card {{
                background: {GRAY_BG};
                border: 1px solid {GRAY_LINE};
                border-radius: 8px;
            }}
        """)
        ff_pfl = QtWidgets.QVBoxLayout(self._ff_preview_frame)
        ff_pfl.setContentsMargins(12, 8, 12, 10)
        ff_pfl.setSpacing(6)
        ff_pfl.addWidget(make_label("Field preview  (first loaded sequence):", size=15, color=TEXT_SEC))
        self._ff_preview_lbl = QtWidgets.QLabel("—  no files loaded")
        self._ff_preview_lbl.setStyleSheet(
            f"font-family:'Courier New',Consolas,monospace; font-size:14px; color:{TEXT_PRI};"
            f" background:transparent;"
        )
        self._ff_preview_lbl.setWordWrap(True)
        ff_pfl.addWidget(self._ff_preview_lbl)
        ffw.addWidget(self._ff_preview_frame)

        # ── Criteria rows ────────────────────────────────────────────────────
        ffw.addWidget(make_label("Criteria:", color=TEXT_SEC))

        self._filter_criteria_container = QtWidgets.QWidget()
        self._filter_criteria_layout = QtWidgets.QVBoxLayout(self._filter_criteria_container)
        self._filter_criteria_layout.setContentsMargins(0, 0, 0, 0)
        self._filter_criteria_layout.setSpacing(6)
        ffw.addWidget(self._filter_criteria_container)

        self._filter_criterion_rows: list = []
        self._ff_field_count: int = 99
        self._ff_keys: list = []
        self._add_filter_criterion()

        add_crit_btn = QtWidgets.QPushButton("+ Add criterion")
        add_crit_btn.setObjectName("secondary_btn")
        add_crit_btn.setFixedWidth(180)
        add_crit_btn.clicked.connect(self._add_filter_criterion)
        add_row = QtWidgets.QHBoxLayout()
        add_row.setSpacing(12)
        add_row.addWidget(add_crit_btn)
        add_row.addWidget(make_label("Combine criteria with:", color=TEXT_SEC))
        self._ff_match_combo = QtWidgets.QComboBox()
        self._ff_match_combo.addItem("AND — all must match", "all")
        self._ff_match_combo.addItem("OR — any can match", "any")
        add_row.addWidget(self._ff_match_combo)
        add_row.addStretch()
        ffw.addLayout(add_row)

        self._filter_widget.hide()
        exl.addWidget(self._filter_widget)

        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_append)
        self._append_widget = QtWidgets.QWidget()
        al = QtWidgets.QVBoxLayout(self._append_widget)
        al.setContentsMargins(20, 4, 0, 4)
        al.setSpacing(8)

        note = make_label(
            "Excel format (.xlsx): column 1 = sequence ID, columns 2+ = fields to "
            "append. The ID is the FASTA header text (without '>') before the "
            "chosen separator — by default ';', e.g. ONTbarcoder's "
            "'DNS-1_all.fa;758;807' → 'DNS-1_all.fa', or 'DNS-1' with suffix "
            "'_all.fa' removed; pick 'space' for plain headers. IDs and cell values are normalized: "
            "leading/trailing spaces removed, spaces/dots/commas → '_', accents removed.",
            color=TEXT_SEC, size=15,
        )
        note.setWordWrap(True)
        al.addWidget(note)

        excel_row = QtWidgets.QHBoxLayout()
        excel_row.setSpacing(8)
        excel_row.addWidget(make_label("Excel file:", color=TEXT_SEC))
        self._excel_edit = _DragDropLineEdit(accepted_extensions=[".xlsx"])
        self._excel_edit.setReadOnly(True)
        self._excel_edit.setPlaceholderText("No file selected… (or drag & drop)")
        excel_row.addWidget(self._excel_edit, 1)
        self._excel_clear = self._make_clear_btn()
        self._excel_clear.setVisible(False)
        self._excel_clear.clicked.connect(lambda: (
            self._excel_edit.clear(),
            self._excel_clear.setVisible(False),
        ))
        self._excel_edit.textChanged.connect(
            lambda t: self._excel_clear.setVisible(bool(t)))
        excel_row.addWidget(self._excel_clear)
        self._excel_btn = QtWidgets.QPushButton("Browse…")
        self._excel_btn.setObjectName("secondary_btn")
        self._excel_btn.setFixedWidth(120)
        self._excel_btn.clicked.connect(self._browse_excel)
        excel_row.addWidget(self._excel_btn)
        al.addLayout(excel_row)

        self._append_id_rule = _IdRuleWidget()
        self._append_id_rule.changed.connect(self._update_append_preview)
        self._excel_edit.textChanged.connect(self._update_append_preview)
        al.addWidget(self._append_id_rule)

        self._append_sep = _SeparatorPicker()
        al.addWidget(self._append_sep)

        self._append_widget.hide()
        ops_layout.addWidget(self._append_widget)

        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_reformat)
        self._reformat_widget = QtWidgets.QWidget()
        rfl = QtWidgets.QVBoxLayout(self._reformat_widget)
        rfl.setContentsMargins(20, 4, 0, 4)
        rfl.setSpacing(10)

        self._reformat_mode_group = QtWidgets.QButtonGroup(self)
        self._rf_radio_linearize  = QtWidgets.QRadioButton(
            "Multi-line → Single line  (linearize / unwrap)")
        self._rf_radio_wrap       = QtWidgets.QRadioButton(
            "Single line → Multi-line  (wrap / fold)")
        self._rf_radio_linearize.setChecked(True)
        self._reformat_mode_group.addButton(self._rf_radio_linearize)
        self._reformat_mode_group.addButton(self._rf_radio_wrap)
        self._reformat_mode_group.buttonClicked.connect(self._on_reformat_mode_changed)
        rfl.addWidget(self._rf_radio_linearize)
        rfl.addWidget(self._rf_radio_wrap)

        self._wrap_cols_row = QtWidgets.QWidget()
        wcr = QtWidgets.QHBoxLayout(self._wrap_cols_row)
        wcr.setContentsMargins(0, 0, 0, 0)
        wcr.setSpacing(10)
        wcr.addWidget(make_label("Nucleotides per row:", color=TEXT_SEC))
        self._wrap_cols_spin = QtWidgets.QSpinBox()
        self._wrap_cols_spin.setRange(10, 10000)
        self._wrap_cols_spin.setValue(80)
        self._wrap_cols_spin.setFixedWidth(90)
        wcr.addWidget(self._wrap_cols_spin)
        wcr.addStretch()
        self._wrap_cols_row.hide()
        rfl.addWidget(self._wrap_cols_row)

        self._reformat_widget.hide()
        ops_layout.addWidget(self._reformat_widget)

        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_orf)
        self._orf_widget = QtWidgets.QWidget()
        orl = QtWidgets.QHBoxLayout(self._orf_widget)
        orl.setContentsMargins(20, 4, 0, 4)
        orl.setSpacing(8)
        orl.addWidget(make_label("Genetic code:", color=TEXT_SEC))
        self._orf_gcode_combo = QtWidgets.QComboBox()
        for name in self._FASTA_GCODES:
            self._orf_gcode_combo.addItem(name)
        self._orf_gcode_combo.setCurrentText("Invertebrate mitochondrial")
        orl.addWidget(self._orf_gcode_combo)
        orl.addSpacing(12)
        orl.addWidget(make_label("Min. ORF coverage:", color=TEXT_SEC))
        self._orf_cov_spin = QtWidgets.QSpinBox()
        self._orf_cov_spin.setRange(50, 100)
        self._orf_cov_spin.setValue(95)
        self._orf_cov_spin.setSuffix(" %")
        self._orf_cov_spin.setFixedWidth(90)
        self._orf_cov_spin.setToolTip(
            "A sequence is trimmed only if its clean ORF covers at least this share\n"
            "of its length; otherwise it is left intact and listed in the log\n"
            "(possible NUMT / pseudogene, unusual length or wrong genetic code).")
        orl.addWidget(self._orf_cov_spin)
        orl.addStretch(1)
        self._orf_widget.hide()
        ops_layout.addWidget(self._orf_widget)

        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_sort)

        self._sort_widget = QtWidgets.QWidget()
        svl = QtWidgets.QVBoxLayout(self._sort_widget)
        svl.setContentsMargins(20, 4, 0, 4)
        svl.setSpacing(10)

        sep_note = make_label(
            "If headers contain a separator, each resulting field can be used as a sort key.",
            color=TEXT_SEC, size=15,
        )
        sep_note.setWordWrap(True)
        svl.addWidget(sep_note)

        self._sort_sep = _SeparatorPicker(allow_none=True)
        svl.addWidget(self._sort_sep)
        self._sort_sep.changed.connect(self._update_sort_preview)

        self._sort_preview_frame = QtWidgets.QFrame()
        self._sort_preview_frame.setObjectName("sort_preview_card")
        self._sort_preview_frame.setStyleSheet(f"""
            QFrame#sort_preview_card {{
                background: {GRAY_BG};
                border: 1px solid {GRAY_LINE};
                border-radius: 8px;
            }}
        """)
        pfl = QtWidgets.QVBoxLayout(self._sort_preview_frame)
        pfl.setContentsMargins(12, 8, 12, 10)
        pfl.setSpacing(6)
        pfl.addWidget(make_label("Field preview  (first loaded sequence):", size=15, color=TEXT_SEC))
        self._sort_preview_lbl = QtWidgets.QLabel("—  no files loaded")
        self._sort_preview_lbl.setStyleSheet(
            f"font-family:'Courier New',Consolas,monospace; font-size:14px; color:{TEXT_PRI};"
            f" background:transparent;"
        )
        self._sort_preview_lbl.setWordWrap(True)
        pfl.addWidget(self._sort_preview_lbl)
        svl.addWidget(self._sort_preview_frame)

        svl.addWidget(make_label("Sort levels (applied top to bottom):", color=TEXT_SEC))

        self._sort_levels_container = QtWidgets.QWidget()
        self._sort_levels_layout = QtWidgets.QVBoxLayout(self._sort_levels_container)
        self._sort_levels_layout.setContentsMargins(0, 0, 0, 0)
        self._sort_levels_layout.setSpacing(6)
        svl.addWidget(self._sort_levels_container)

        self._sort_level_rows: list = []
        self._add_sort_level()

        add_level_btn = QtWidgets.QPushButton("+ Add sort level")
        add_level_btn.setObjectName("secondary_btn")
        add_level_btn.setFixedWidth(180)
        add_level_btn.clicked.connect(self._add_sort_level)
        svl.addWidget(add_level_btn)

        self._sort_widget.hide()
        ops_layout.addWidget(self._sort_widget)

        ops_layout.addSpacing(4)
        ops_layout.addWidget(self._radio_split)

        self._split_widget = QtWidgets.QWidget()
        spl = QtWidgets.QVBoxLayout(self._split_widget)
        spl.setContentsMargins(20, 4, 0, 4)
        spl.setSpacing(8)

        split_note = make_label(
            "Every sequence goes to exactly one file; sequences are never cut. "
            "Files are written to a <name>_split folder.",
            color=TEXT_SEC, size=15,
        )
        split_note.setWordWrap(True)
        spl.addWidget(split_note)

        self._split_mode_group = QtWidgets.QButtonGroup(self)

        def _mode_row(label, spin, unit, tip):
            row = QtWidgets.QHBoxLayout()
            row.setSpacing(8)
            rb = QtWidgets.QRadioButton(label)
            rb.setToolTip(tip)
            self._split_mode_group.addButton(rb)
            row.addWidget(rb)
            if spin is not None:
                rb.setMinimumWidth(250)   # line the number boxes up
                spin.setFixedWidth(130)
                spin.setGroupSeparatorShown(True)
                row.addWidget(spin)
                row.addWidget(make_label(unit, color=TEXT_SEC))
            row.addStretch()
            spl.addLayout(row)
            return rb

        self._split_count_spin = QtWidgets.QSpinBox()
        self._split_count_spin.setRange(1, 10_000_000)
        self._split_count_spin.setValue(50)
        self._split_by_count = _mode_row(
            "By number of sequences:", self._split_count_spin, "sequences per file",
            "Consecutive files of this many sequences (the last one may hold fewer).\n"
            "E.g. batches for BLAST on the NCBI website.")

        self._split_files_spin = QtWidgets.QSpinBox()
        self._split_files_spin.setRange(2, 10_000)
        self._split_files_spin.setValue(4)
        self._split_by_files = _mode_row(
            "Into a number of files:", self._split_files_spin, "files",
            "Sequences shared out in order into this many files of near-equal size\n"
            "(sizes differ by one at most).")

        self._split_bases_spin = QtWidgets.QSpinBox()
        self._split_bases_spin.setRange(1_000, 2_000_000_000)
        self._split_bases_spin.setSingleStep(100_000)
        self._split_bases_spin.setValue(1_000_000)
        self._split_by_bases = _mode_row(
            "By total length:", self._split_bases_spin, "bases per file at most",
            "Consecutive files whose sequences add up to at most this many bases.\n"
            "1,000,000 is NCBI BLAST's limit per blastn query. A single sequence\n"
            "longer than the limit gets a file of its own.")

        self._split_by_field = _mode_row(
            "By header field — one file per distinct value", None, "",
            "One file per distinct value of the chosen header field, named after it\n"
            "(e.g. one file per sample or per species). Unlike Extract, you\n"
            "do not need to know or type the values, and no sequence is left out:\n"
            "headers without that field go to <name>_no_field.fasta.")
        self._split_by_count.setChecked(True)

        # Field options, shown only for "By header field"
        self._split_field_box = QtWidgets.QWidget()
        sfl = QtWidgets.QVBoxLayout(self._split_field_box)
        sfl.setContentsMargins(24, 0, 0, 0)
        sfl.setSpacing(6)
        self._split_sep = _SeparatorPicker(stretch=False)
        sf_row = self._split_sep.row
        sf_row.addSpacing(12)
        sf_row.addWidget(make_label("Field:", color=TEXT_SEC))
        self._split_field = _FieldSelector()
        sf_row.addWidget(self._split_field)
        sf_row.addStretch()
        sfl.addWidget(self._split_sep)
        self._split_field_preview = make_label("", size=14, color=TEXT_HINT)
        self._split_field_preview.setWordWrap(True)
        sfl.addWidget(self._split_field_preview)
        self._split_field_box.hide()
        spl.addWidget(self._split_field_box)

        self._split_mode_group.buttonClicked.connect(self._on_split_mode_changed)
        self._split_sep.changed.connect(self._update_split_preview)
        self._split_field.changed.connect(self._update_split_preview)

        self._split_widget.hide()
        ops_layout.addWidget(self._split_widget)

        self._layout.addWidget(ops_box)

        self._status_lbl = make_label("", color=TEXT_SEC)
        self._layout.addWidget(self._status_lbl)

        self._progress_bar = QtWidgets.QProgressBar()
        self._progress_bar.setRange(0, 0)
        self._progress_bar.setFixedHeight(6)
        self._progress_bar.hide()
        self._layout.addWidget(self._progress_bar)

        self._layout.addStretch()

        self._log_edit = QtWidgets.QPlainTextEdit()
        self._log_edit.setReadOnly(True)
        self._log_edit.setFixedHeight(110)
        self._log_edit.setStyleSheet(f"""
            QPlainTextEdit {{
                background: {GRAY_BG};
                border: 1px solid {GRAY_LINE};
                border-top: 1px solid {GRAY_LINE};
                border-radius: 0px;
                font-family: 'Courier New', Consolas, monospace;
                font-size: 16px;
                color: {TEXT_PRI};
                padding: 6px;
            }}
        """)
        self._log_edit.hide()
        outer.addWidget(self._log_edit)

        footer = QtWidgets.QWidget()
        footer.setObjectName("fasta_tools_footer")
        footer.setStyleSheet(f"""
            QWidget#fasta_tools_footer {{
                background: {GRAY_CARD};
                border-top: 1px solid {GRAY_LINE};
            }}
        """)
        fl = QtWidgets.QHBoxLayout(footer)
        fl.setContentsMargins(20, 10, 20, 10)
        fl.setSpacing(8)

        self._clear_btn = QtWidgets.QPushButton("Clear")
        self._clear_btn.setObjectName("danger_btn")
        self._clear_btn.setFixedHeight(44)
        self._clear_btn.setFixedWidth(140)
        self._clear_btn.clicked.connect(self._clear)
        fl.addWidget(self._clear_btn)

        self._open_folder_btn = QtWidgets.QPushButton("Open folder")
        self._open_folder_btn.setObjectName("secondary_btn")
        self._open_folder_btn.setFixedHeight(44)
        self._open_folder_btn.hide()
        self._open_folder_btn.clicked.connect(self._open_output_folder)
        fl.addWidget(self._open_folder_btn)
        fl.addStretch()

        self._run_btn = QtWidgets.QPushButton("Run  →")
        self._run_btn.setObjectName("primary_btn")
        self._run_btn.setFixedHeight(44)
        self._run_btn.setFixedWidth(300)
        self._run_btn.clicked.connect(self._on_run_clicked)
        self._set_run_enabled(False)
        fl.addWidget(self._run_btn)
        outer.addWidget(footer)

        self._worker: Optional[_FastaToolsWorker] = None
        # Workers stopped by Clear that were still busy: referenced here so
        # Python does not destroy a running QThread (which aborts the app).
        self._retired_workers: list = []
        self._hdr_cache: dict = {}     # path -> (stamp, headers) for ID previews
        self._excel_cache: dict = {}   # path -> (stamp, lookup) for ID previews
        self._files: list = []
        self._split_blocked = False
        self._split_distinct = 0
        self._split_truncated = False
        self._running = False
        self._last_outputs: list = []

    @staticmethod
    def _make_section_lbl(text):
        lbl = make_section_label(text)
        lbl.setStyleSheet(
            f"font-size:17px; font-weight:600; color:{BLUE}; letter-spacing:0.4px;"
        )
        return lbl

    @staticmethod
    def _make_clear_btn():
        btn = QtWidgets.QPushButton("✕")
        btn.setFixedSize(26, 26)
        btn.setToolTip("Clear")
        btn.setStyleSheet(
            f"QPushButton {{ background-color: transparent; color: {TEXT_HINT};"
            f" border: none; border-radius: 5px; font-size: 13px; font-weight: bold; }}"
            f"QPushButton:hover {{ background-color: {RED_LT}; color: {RED}; }}"
            f"QPushButton:pressed {{ background-color: {RED}; color: white; }}"
        )
        return btn

    def _set_run_enabled(self, enabled: bool):
        self._run_btn.setEnabled(enabled)
        if enabled:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color:{BLUE}; color:white; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
                f"QPushButton:hover {{ background-color:#0C4A82; }}"
            )
        else:
            self._run_btn.setStyleSheet(
                f"QPushButton {{ background-color:{GRAY_LINE}; color:{TEXT_HINT}; border:none; "
                f"border-radius:8px; padding:9px 20px; font-size:18px; font-weight:500; }}"
            )

    def _refresh_run_state(self):
        """Run is enabled when files are loaded, unless "By header field"
        would not split anything (see _update_split_preview)."""
        self._set_run_enabled(bool(self._files) and not self._split_blocked
                              and not self._running)

    def _on_files(self, files: list):
        self._files = files
        self._split_blocked = False
        self._refresh_run_state()
        self._status_lbl.setText("")
        self._status_lbl.setStyleSheet("")
        self._update_sort_preview()
        self._update_ff_preview()
        self._update_append_preview()
        self._update_grep_preview()
        self._update_split_preview()

    def _on_operation_changed(self, _btn):
        self._stats_widget.setVisible(self._radio_stats.isChecked())
        extract = self._radio_extract.isChecked()
        self._extract_widget.setVisible(extract)
        self._grep_widget.setVisible(extract and self._radio_grep.isChecked())
        self._filter_widget.setVisible(extract and self._radio_filter_fields.isChecked())
        self._append_widget.setVisible(self._radio_append.isChecked())
        self._reformat_widget.setVisible(self._radio_reformat.isChecked())
        self._orf_widget.setVisible(self._radio_orf.isChecked())
        self._sort_widget.setVisible(self._radio_sort.isChecked())
        self._split_widget.setVisible(self._radio_split.isChecked())
        self._update_split_preview()
        if self._radio_sort.isChecked():
            self._update_sort_preview()
        if extract and self._radio_filter_fields.isChecked():
            self._update_ff_preview()
        self._log_edit.clear()
        self._log_edit.hide()
        self._status_lbl.setText("")
        self._status_lbl.setStyleSheet("")
        self._open_folder_btn.hide()

    def _on_split_mode_changed(self, _btn):
        self._split_field_box.setVisible(self._split_by_field.isChecked())
        self._update_split_preview()

    def _update_split_preview(self, *_):
        """Under the field options: how many files "By header field" would
        write, with a few of the values. Uses the cached headers of the other
        previews, so changing the field never re-reads a large file."""
        self._split_blocked = False
        self._split_distinct, self._split_truncated = 0, False
        if not (self._radio_split.isChecked() and self._split_by_field.isChecked()):
            self._refresh_run_state()
            return
        if not self._files:
            self._split_field_preview.setText("—  no files loaded")
            return
        headers, truncated = self._headers()
        if not headers:
            self._split_field_preview.setText("—  no sequences found")
            return
        sep = self._split_sep.separator()
        # Cap the field box at the largest field count found with this separator.
        with_sep = [h for h in headers if not sep or sep in h]
        max_fields = max((len(h.split(sep)) if sep else 1 for h in with_sep),
                         default=1)
        sel = self._split_field
        if not with_sep:
            self._split_blocked = True
            self._refresh_run_state()
            self._split_field_preview.setText(
                f'—  separator "{sep}" not found in the headers → nothing to split by')
            return
        if sel.by_key():
            keys = header_keys(with_sep, sep)
            sel.set_keys(keys)
            if not keys:
                self._split_blocked = True
                self._refresh_run_state()
                self._split_field_preview.setText(
                    "—  no key=value fields found in the headers → use Position")
                return
        else:
            sel.set_max(max_fields, quiet=True)
        field = sel.value()
        values, missing = [], 0
        for hdr in headers:
            v = header_field_value(hdr, sep, field)
            if v is None:
                missing += 1
            else:
                values.append(v)
        # Nothing to split by (no header has this separator / field): block Run.
        self._split_blocked = not values
        self._refresh_run_state()
        distinct = list(dict.fromkeys(values))
        self._split_distinct, self._split_truncated = len(distinct), truncated
        example = ", ".join(distinct[:5]) + (" …" if len(distinct) > 5 else "")
        scope = (f"first {len(headers):,} headers" if truncated
                 else os.path.basename(self._files[0]) if len(self._files) == 1
                 else f"{len(self._files)} files")
        text = (f"{scope}: {field_label(field)} has "
                f"{len(distinct)} distinct value(s) → {len(distinct)} file(s)"
                + (f"  ·  e.g. {example}" if distinct else ""))
        if len(distinct) > _SPLIT_FILE_WARN:
            text += f"  ·  ⚠ more than {_SPLIT_FILE_WARN} files"
        if missing:
            text += f"  ·  {missing} header(s) without this field → _no_field file"
        self._split_field_preview.setText(text)

    def _on_reformat_mode_changed(self, _btn):
        self._wrap_cols_row.setVisible(self._rf_radio_wrap.isChecked())

    def _update_sort_preview(self, *_):
        if not self._files:
            self._sort_preview_lbl.setText("—  no files loaded")
            return
        try:
            header = _first_header(self._files[0])
        except Exception:
            self._sort_preview_lbl.setText("—  could not read file")
            return
        if not header:
            self._sort_preview_lbl.setText("—  no sequences found")
            return
        sep = self._sort_sep.separator()
        # Match the worker's split (unconditional when a separator is set), so the
        # field numbering in the preview always lines up with the actual sort.
        parts = header.split(sep) if sep else [header]
        self._sort_preview_lbl.setText(
            "\n".join(f"Field {i:>2}:  {p.strip()}" for i, p in enumerate(parts, 1))
        )
        self._sort_field_count = len(parts)
        self._update_field_spin_limits()

    def _update_field_spin_limits(self):
        for row in self._sort_level_rows:
            field_spin = row[1]
            field_spin.setMaximum(self._sort_field_count)
            if field_spin.value() > self._sort_field_count:
                field_spin.setValue(self._sort_field_count)

    def _add_sort_level(self):
        n = len(self._sort_level_rows) + 1

        row_w = QtWidgets.QWidget()
        rl = QtWidgets.QHBoxLayout(row_w)
        rl.setContentsMargins(0, 0, 0, 0)
        rl.setSpacing(8)

        lbl = make_label(f"Level {n}:", color=TEXT_SEC)
        lbl.setFixedWidth(60)
        rl.addWidget(lbl)

        rl.addWidget(make_label("Field:", color=TEXT_SEC))
        field_spin = QtWidgets.QSpinBox()
        field_spin.setRange(1, self._sort_field_count)
        field_spin.setValue(min(n, self._sort_field_count))
        field_spin.setFixedWidth(70)
        rl.addWidget(field_spin)

        rl.addSpacing(8)
        rl.addWidget(make_label("Order:", color=TEXT_SEC))
        order_combo = QtWidgets.QComboBox()
        order_combo.addItem("Ascending  ↑", "asc")
        order_combo.addItem("Descending ↓", "desc")
        order_combo.setFixedWidth(160)
        rl.addWidget(order_combo)

        if n > 1:
            remove_btn = QtWidgets.QPushButton("✕")
            remove_btn.setFixedSize(28, 28)
            remove_btn.setStyleSheet(
                "QPushButton { background:transparent; border:1px solid #CCC;"
                " border-radius:4px; color:#888; font-size:11px; }"
                "QPushButton:hover { background:#FEE2E2; border-color:#EF4444; color:#EF4444; }"
            )
            remove_btn.clicked.connect(lambda _, w=row_w: self._remove_sort_level(w))
            rl.addWidget(remove_btn)

        rl.addStretch()
        self._sort_levels_layout.addWidget(row_w)
        self._sort_level_rows.append((row_w, field_spin, order_combo))

    def _remove_sort_level(self, row_widget):
        for i, (w, _fs, _oc) in enumerate(self._sort_level_rows):
            if w is row_widget:
                self._sort_level_rows.pop(i)
                self._sort_levels_layout.removeWidget(w)
                w.deleteLater()
                self._renumber_sort_levels()
                break

    def _renumber_sort_levels(self):
        for i, (w, _fs, _oc) in enumerate(self._sort_level_rows):
            lbl = w.layout().itemAt(0).widget()
            if isinstance(lbl, QtWidgets.QLabel):
                lbl.setText(f"Level {i + 1}:")

    def _update_ff_preview(self, *_):
        if not self._files:
            self._ff_preview_lbl.setText("—  no files loaded")
            return
        try:
            header = _first_header(self._files[0])
        except Exception:
            self._ff_preview_lbl.setText("—  could not read file")
            return
        if not header:
            self._ff_preview_lbl.setText("—  no sequences found")
            return
        sep = self._ff_sep.separator() or "|"
        self._ff_keys = header_keys(self._headers()[0] or [header], sep)
        if sep not in header:
            # The run treats such headers as having no fields at all.
            self._ff_preview_lbl.setText(
                f'—  separator "{sep}" not found in the first header')
            self._ff_field_count = 1
            self._update_ff_spin_limits()
            return
        parts = header.split(sep)
        self._ff_preview_lbl.setText(
            "\n".join(f"Field {i:>2}:  {p.strip()}" for i, p in enumerate(parts, 1))
        )
        self._ff_field_count = len(parts)
        self._update_ff_spin_limits()

    def _update_ff_spin_limits(self):
        for _w, fs, _oc, _ve in self._filter_criterion_rows:
            fs.set_max(self._ff_field_count)
            fs.set_keys(self._ff_keys)

    def _add_filter_criterion(self):
        _REMOVE_STYLE = (
            "QPushButton { background:transparent; border:1px solid #CCC;"
            " border-radius:4px; color:#888; font-size:11px; }"
            "QPushButton:hover { background:#FEE2E2; border-color:#EF4444; color:#EF4444; }"
        )
        n = len(self._filter_criterion_rows) + 1
        row_w = QtWidgets.QWidget()
        rl = QtWidgets.QHBoxLayout(row_w)
        rl.setContentsMargins(0, 0, 0, 0)
        rl.setSpacing(6)

        rl.addWidget(make_label("Field:", color=TEXT_SEC))
        field_spin = _FieldSelector()
        field_spin.set_max(self._ff_field_count)
        field_spin.set_position(min(n, self._ff_field_count))
        field_spin.set_keys(self._ff_keys)
        rl.addWidget(field_spin)

        op_combo = QtWidgets.QComboBox()
        for op in ("=", "!=", ">=", "<=", ">", "<", "contains", "not contains"):
            op_combo.addItem(op)
        op_combo.setFixedWidth(130)
        rl.addWidget(op_combo)

        value_edit = QtWidgets.QLineEdit()
        value_edit.setPlaceholderText("value")
        value_edit.setFixedWidth(140)
        rl.addWidget(value_edit)

        if self._filter_criterion_rows:
            remove_btn = QtWidgets.QPushButton("✕")
            remove_btn.setFixedSize(28, 28)
            remove_btn.setStyleSheet(_REMOVE_STYLE)
            remove_btn.clicked.connect(lambda _, w=row_w: self._remove_filter_criterion(w))
            rl.addWidget(remove_btn)

        rl.addStretch()
        self._filter_criteria_layout.addWidget(row_w)
        self._filter_criterion_rows.append((row_w, field_spin, op_combo, value_edit))

    def _remove_filter_criterion(self, row_widget):
        for i, (w, _fs, _oc, _ve) in enumerate(self._filter_criterion_rows):
            if w is row_widget:
                self._filter_criterion_rows.pop(i)
                self._filter_criteria_layout.removeWidget(w)
                w.deleteLater()
                break

    def _on_grep_source_changed(self, state):
        use_file = (state == QtCore.Qt.Checked)
        self._grep_pattern_row.setVisible(not use_file)
        self._grep_file_row.setVisible(use_file)
        self._grep_file_note.setVisible(use_file)
        # A pattern file is almost always a list of IDs: match them whole.
        if use_file and not self._grep_regex_chk.isChecked():
            self._grep_exact_chk.setChecked(True)
        self._update_grep_preview()

    # ── ID previews (Append info / exact-ID grep) ─────────────────────────
    _PREVIEW_MAX_HEADERS = 20000   # per file; enough for a representative count

    @staticmethod
    def _stamp(path):
        try:
            return (os.path.getmtime(path), os.path.getsize(path))
        except OSError:
            return None

    def _headers(self) -> tuple:
        """(headers of the loaded FASTA files, truncated?) for the previews."""
        import gzip
        out, truncated = [], False
        for path in self._files:
            stamp = self._stamp(path)
            hit = self._hdr_cache.get(path)
            if not hit or hit[0] != stamp:
                hdrs, more = [], False
                try:
                    opener = gzip.open if path.lower().endswith(".gz") else open
                    with opener(path, "rt", encoding="utf-8", errors="replace") as fh:
                        for line in fh:
                            if line.startswith(">"):
                                if len(hdrs) >= self._PREVIEW_MAX_HEADERS:
                                    more = True
                                    break
                                hdrs.append(line[1:].rstrip("\r\n"))
                except OSError:
                    pass
                hit = (stamp, hdrs, more)
                self._hdr_cache[path] = hit
            out.extend(hit[1])
            truncated |= hit[2]
        return out, truncated

    def _count_line(self, headers, truncated, n_found, what) -> str:
        scope = f"{len(headers)}" + ("+ (first headers only)" if truncated else "")
        return f"{n_found}/{scope} sequence(s) {what}"

    def _update_append_preview(self, *_):
        if not hasattr(self, "_append_id_rule"):
            return
        w = self._append_id_rule
        headers, truncated = self._headers()
        if not headers:
            w.set_preview("Load a FASTA file to preview the ID taken from each header.")
            return
        rule = w.rule()
        first = extract_id(headers[0], rule)
        head = (f"First header → ID: <b>{first or '(empty)'}</b>")
        path = self._excel_edit.text().strip()
        if not path or not os.path.isfile(path):
            w.set_preview(head + " · select the Excel file to check matches")
            return
        stamp = self._stamp(path)
        hit = self._excel_cache.get(path)
        if not hit or hit[0] != stamp:
            try:
                hit = (stamp, _FastaToolsWorker.read_excel_lookup(path))
            except Exception as e:
                w.set_preview(head + f" · cannot read the Excel file: {e}", False)
                return
            self._excel_cache[path] = hit
        lookup = hit[1]
        norm = _FastaToolsWorker._normalize_field
        n = sum(1 for h in headers if norm(extract_id(h, rule)) in lookup)
        w.set_preview(head + " · " + self._count_line(headers, truncated, n,
                                                      "found in the Excel"), n > 0)

    def _grep_patterns_preview(self) -> list:
        if self._grep_use_file_chk.isChecked():
            path = self._grep_file_edit.text().strip()
            try:
                with open(path, encoding="utf-8", errors="replace") as fh:
                    return [ln.strip() for ln in fh if ln.strip()]
            except OSError:
                return []
        text = self._grep_pattern_edit.text().strip()
        return [text] if text else []

    def _update_grep_preview(self, *_):
        if not hasattr(self, "_grep_id_rule") or not self._grep_exact_chk.isChecked():
            return
        w = self._grep_id_rule
        headers, truncated = self._headers()
        if not headers:
            w.set_preview("Load a FASTA file to preview the ID taken from each header.")
            return
        rule = w.rule()
        head = f"First header → ID: <b>{extract_id(headers[0], rule) or '(empty)'}</b>"
        ic = self._grep_case_chk.isChecked()
        norm = (lambda t: t.lower()) if ic else (lambda t: t)
        wanted = {norm(w) for w in self._grep_patterns_preview()}
        if not wanted:
            w.set_preview(head + " · enter a pattern or load a pattern file")
            return
        n = sum(1 for h in headers if norm(extract_id(h, rule)) in wanted)
        w.set_preview(head + " · " + self._count_line(headers, truncated, n,
                                                      "match the list"), n > 0)

    def _on_grep_exact_toggled(self, *_):
        regex = self._grep_regex_chk.isChecked()
        # Exact ID and regex are alternative ways of matching.
        self._grep_exact_chk.setEnabled(not regex)
        exact = self._grep_exact_chk.isChecked() and not regex
        self._grep_id_rule.setVisible(exact)
        self._update_grep_preview()

    def _browse_grep_file(self):
        path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, "Select patterns file", "", "Text files (*.txt);;All files (*)"
        )
        if path:
            self._grep_file_edit.setText(path)

    def _browse_excel(self):
        path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, "Select Excel file", "", "Excel workbook (*.xlsx)"
        )
        if path:
            self._excel_edit.setText(path)

    def _clear(self):
        if self._worker and self._worker.isRunning():
            for sig, slot in (
                (self._worker.progress, self._on_progress),
                (self._worker.log_line, self._on_log_line),
                (self._worker.done,     self._on_finished),
                (self._worker.error,    self._on_error),
            ):
                try:
                    sig.disconnect(slot)
                except RuntimeError:
                    pass
            self._worker.stop()
            if not self._worker.wait(500):
                self._retired_workers.append(self._worker)
            self._worker = None
        self._running = False

        self._drop.clear()
        self._files = []
        self._last_outputs = []
        self._status_lbl.setText("")
        self._status_lbl.setStyleSheet("")
        self._progress_bar.hide()
        self._log_edit.clear()
        self._log_edit.hide()
        self._open_folder_btn.hide()
        self._set_run_enabled(False)
        self._merge_chk.setChecked(False)

        self._radio_stats.setChecked(True)
        self._stats_widget.show()
        self._stats_stops_chk.setChecked(False)
        self._stats_gcode_combo.setCurrentText("Invertebrate mitochondrial")
        self._extract_widget.hide()
        self._radio_unique.setChecked(True)
        self._grep_widget.hide()
        self._filter_widget.hide()
        self._append_widget.hide()
        self._reformat_widget.hide()
        self._orf_widget.hide()
        self._orf_gcode_combo.setCurrentText("Invertebrate mitochondrial")
        self._orf_cov_spin.setValue(95)
        self._sort_widget.hide()
        self._split_widget.hide()
        self._split_by_count.setChecked(True)
        self._split_count_spin.setValue(50)
        self._split_files_spin.setValue(4)
        self._split_bases_spin.setValue(1_000_000)
        self._split_sep.reset()
        self._split_field.reset()
        self._split_field_box.hide()

        self._grep_use_file_chk.setChecked(False)
        self._grep_pattern_edit.clear()
        self._grep_file_edit.clear()
        self._grep_regex_chk.setChecked(False)
        self._grep_case_chk.setChecked(False)
        self._grep_exact_chk.setChecked(False)
        self._grep_id_rule.reset()
        self._grep_id_rule.hide()
        self._grep_pattern_row.show()
        self._grep_file_row.hide()
        self._grep_file_note.hide()

        self._excel_edit.clear()
        self._append_sep.reset()
        self._append_id_rule.reset()
        self._hdr_cache.clear()
        self._excel_cache.clear()

        self._rf_radio_linearize.setChecked(True)
        self._wrap_cols_spin.setValue(80)
        self._wrap_cols_row.hide()

        self._sort_sep.reset()
        while len(self._sort_level_rows) > 1:
            w, _fs, _oc = self._sort_level_rows[-1]
            self._sort_level_rows.pop()
            self._sort_levels_layout.removeWidget(w)
            w.deleteLater()
        if self._sort_level_rows:
            _w, fs, oc = self._sort_level_rows[0]
            fs.setValue(1)
            oc.setCurrentIndex(0)
        self._update_sort_preview()

        self._ff_sep.reset()
        self._ff_match_combo.setCurrentIndex(0)
        self._ff_keys = []
        while self._filter_criterion_rows:
            w, _fs, _oc, _ve = self._filter_criterion_rows.pop()
            self._filter_criteria_layout.removeWidget(w)
            w.deleteLater()
        self._add_filter_criterion()

    def _on_run_clicked(self):
        if not self._files:
            return

        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        folder_name = f"ont-barcoder_{ts}_fastools"
        auto_dir = os.path.join(_get_base_dir(), "output", folder_name)

        out_dir = self._ask_output_dir(auto_dir, folder_name)
        if out_dir is None:
            return

        operation = (
            "stats"          if self._radio_stats.isChecked()          else
            ("unique"        if self._radio_unique.isChecked()        else
             "identical"     if self._radio_identical.isChecked()     else
             "filter_fields" if self._radio_filter_fields.isChecked() else
             "grep")         if self._radio_extract.isChecked()       else
            "append"         if self._radio_append.isChecked()         else
            "reformat"       if self._radio_reformat.isChecked()       else
            "orf_trim"       if self._radio_orf.isChecked()            else
            "split"          if self._radio_split.isChecked()          else
            "sort"
        )

        params: dict = {"merge": self._merge_chk.isChecked()}
        if operation == "stats":
            if self._stats_stops_chk.isChecked():
                params["genetic_code"] = self._stats_gcode_combo.currentText()
        elif operation == "grep":
            params["use_regex"] = self._grep_regex_chk.isChecked()
            params["ignore_case"] = self._grep_case_chk.isChecked()
            if self._grep_use_file_chk.isChecked():
                params["pattern_file"] = self._grep_file_edit.text().strip()
            else:
                params["pattern"] = self._grep_pattern_edit.text().strip()
            if self._grep_exact_chk.isChecked() and not params["use_regex"]:
                params["exact_id"] = self._grep_id_rule.rule()
        elif operation == "filter_fields":
            if self._ff_sep.custom_empty():
                QtWidgets.QMessageBox.warning(
                    self, "Filter configuration",
                    "Custom separator is empty.\n"
                    "Please enter a separator character or choose a different option."
                )
                return
            params["separator"] = self._ff_sep.separator()
            params["match"] = self._ff_match_combo.currentData()
            criteria = [
                {"field": fs.value(), "op": oc.currentText(), "value": ve.text().strip()}
                for _w, fs, oc, ve in self._filter_criterion_rows
                if ve.text().strip()
            ]
            if not criteria:
                QtWidgets.QMessageBox.warning(
                    self, "Filter configuration",
                    "No valid criteria.\n"
                    "Please fill in at least one value."
                )
                return
            if any(c["field"] == "" for c in criteria):
                QtWidgets.QMessageBox.warning(
                    self, "Filter configuration",
                    "A criterion selects its field by key name, but the key is empty.\n"
                    "Type a key (e.g. ambs) or switch to Position."
                )
                return
            # Numeric operators need a numeric value; otherwise the criterion
            # would fall into the text branch (which has no case for them) and
            # silently pass every sequence.
            for c in criteria:
                if c["op"] in (">=", "<=", ">", "<"):
                    try:
                        float(c["value"])
                    except ValueError:
                        QtWidgets.QMessageBox.warning(
                            self, "Filter configuration",
                            f"Operator '{c['op']}' ({field_label(c['field'])}) requires a "
                            f"numeric value, but got “{c['value']}”.\n\n"
                            "Enter a number or use a text operator "
                            "(=, !=, contains, not contains)."
                        )
                        return
            params["criteria"] = criteria
        elif operation == "append":
            params["excel_file"] = self._excel_edit.text().strip()
            if self._append_sep.custom_empty():
                QtWidgets.QMessageBox.warning(
                    self, "Append configuration",
                    "Custom separator is empty.\n"
                    "Please enter a separator character or choose a different option."
                )
                return
            params["separator"]  = self._append_sep.separator()
            params["id_rule"]    = self._append_id_rule.rule()
        elif operation == "reformat":
            params["mode"]      = "wrap" if self._rf_radio_wrap.isChecked() else "linearize"
            params["wrap_cols"] = self._wrap_cols_spin.value()
        elif operation == "split":
            if self._split_by_field.isChecked():
                sep = self._split_sep.separator()
                if not sep:
                    QtWidgets.QMessageBox.warning(
                        self, "Split configuration",
                        "Custom separator is empty.\n"
                        "Please enter a separator character or choose a different option."
                    )
                    return
                field = self._split_field.value()
                if field == "":
                    QtWidgets.QMessageBox.warning(
                        self, "Split configuration",
                        "The field is selected by key name, but the key is empty.\n"
                        "Type a key (e.g. ambs) or switch to Position."
                    )
                    return
                n = self._split_distinct
                if n > _SPLIT_FILE_WARN:
                    many = len(self._files) > 1
                    ans = QtWidgets.QMessageBox.question(
                        self, "Split configuration",
                        f"The chosen field has {'at least ' if self._split_truncated else ''}"
                        f"{n:,} distinct values{' across the input files' if many else ''}, "
                        f"so this split will write up to {n:,} files"
                        f"{' per input file' if many else ''}.\n\nContinue?",
                        QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
                        QtWidgets.QMessageBox.No)
                    if ans != QtWidgets.QMessageBox.Yes:
                        return
                params.update(mode="field", separator=sep, field=field)
            elif self._split_by_files.isChecked():
                params.update(mode="files", value=self._split_files_spin.value())
            elif self._split_by_bases.isChecked():
                params.update(mode="bases", value=self._split_bases_spin.value())
            else:
                params.update(mode="count", value=self._split_count_spin.value())
        elif operation == "orf_trim":
            params["genetic_code"] = self._orf_gcode_combo.currentText()
            params["min_coverage"] = self._orf_cov_spin.value() / 100.0
        elif operation == "sort":
            if self._sort_sep.custom_empty():
                QtWidgets.QMessageBox.warning(
                    self, "Sort configuration",
                    "Custom separator is empty.\n"
                    "Please enter a separator character or choose a different option."
                )
                return
            params["separator"] = self._sort_sep.separator()
            params["levels"] = [
                {"field": fs.value(), "order": oc.currentData()}
                for _, fs, oc in self._sort_level_rows
            ]
            if not params["separator"]:
                bad = [lv for lv in params["levels"] if lv["field"] > 1]
                if bad:
                    QtWidgets.QMessageBox.warning(
                        self, "Sort configuration",
                        "Separator is set to None (whole header = 1 field).\n"
                        "All sort levels must use Field 1.\n\n"
                        "Please select a separator or set all levels to Field 1."
                    )
                    return

        self._progress_bar.show()
        self._status_lbl.setStyleSheet("")
        self._status_lbl.setText("Running…")
        self._running = True
        self._set_run_enabled(False)
        self._log_edit.clear()
        self._log_edit.show()

        self._retired_workers = [w for w in self._retired_workers if w.isRunning()]
        self._worker = _FastaToolsWorker(self._files, operation, params, out_dir)
        self._worker.progress.connect(self._on_progress)
        self._worker.log_line.connect(self._on_log_line)
        self._worker.done.connect(self._on_finished)
        self._worker.error.connect(self._on_error)
        self._worker.start()

    def _ask_output_dir(self, auto_dir: str, folder_name: str) -> Optional[str]:
        dlg = QtWidgets.QDialog(self)
        dlg.setWindowTitle("Output folder")
        dlg.setMinimumWidth(480)
        dlg.setStyleSheet(f"""
            QDialog {{ background-color: {GRAY_CARD}; }}
            QLabel {{ color: {TEXT_PRI}; background-color: transparent; }}
            QRadioButton {{
                color: {TEXT_PRI}; background-color: transparent;
                font-size: 15px; padding: 6px 0;
            }}
            QRadioButton::indicator {{ width: 16px; height: 16px; }}
            QPushButton {{
                border-radius: 8px; padding: 8px 20px;
                font-size: 15px; font-weight: 500;
            }}
            #dlg_ok_btn {{ background-color: {BLUE}; color: white; border: none; }}
            #dlg_ok_btn:hover {{ background-color: #0C4A82; }}
            #dlg_cancel_btn {{
                background-color: transparent; color: {BLUE};
                border: 1px solid {BLUE};
            }}
            #dlg_cancel_btn:hover {{ background-color: {BLUE_LIGHT}; }}
        """)

        vlay = QtWidgets.QVBoxLayout(dlg)
        vlay.setSpacing(16)
        vlay.setContentsMargins(24, 24, 24, 20)

        title_lbl = QtWidgets.QLabel("Where to save the results?")
        title_lbl.setStyleSheet(f"font-size:17px; font-weight:700; color:{TEXT_PRI};")
        vlay.addWidget(title_lbl)

        radio_auto = QtWidgets.QRadioButton(
            f"Automatic folder (recommended)\n  …/output/{folder_name}/"
        )
        radio_auto.setChecked(True)
        radio_custom = QtWidgets.QRadioButton("Select folder manually")
        vlay.addWidget(radio_auto)
        vlay.addWidget(radio_custom)
        vlay.addSpacing(8)

        btn_row = QtWidgets.QHBoxLayout()
        btn_row.addStretch()
        btn_cancel = QtWidgets.QPushButton("Cancel")
        btn_cancel.setObjectName("dlg_cancel_btn")
        btn_cancel.setFixedHeight(38)
        btn_ok = QtWidgets.QPushButton("Run")
        btn_ok.setObjectName("dlg_ok_btn")
        btn_ok.setFixedHeight(38)
        btn_ok.setDefault(True)
        btn_cancel.clicked.connect(dlg.reject)
        btn_ok.clicked.connect(dlg.accept)
        btn_row.addWidget(btn_cancel)
        btn_row.addSpacing(8)
        btn_row.addWidget(btn_ok)
        vlay.addLayout(btn_row)

        if dlg.exec_() != QtWidgets.QDialog.Accepted:
            return None

        if radio_auto.isChecked():
            return auto_dir
        parent_dir = QtWidgets.QFileDialog.getExistingDirectory(
            self, "Select output folder"
        )
        if not parent_dir:
            return None
        return os.path.join(parent_dir, folder_name)

    def _on_progress(self, msg: str):
        self._status_lbl.setText(msg)

    def _on_log_line(self, line: str):
        self._log_edit.appendPlainText(line)

    def _on_finished(self, outputs: list):
        self._progress_bar.hide()
        self._running = False
        self._refresh_run_state()
        self._last_outputs = outputs
        self._status_lbl.setStyleSheet(f"color:{GREEN};")
        self._status_lbl.setText(f"Done — {len(outputs)} file(s) saved.")
        self._log_edit.appendPlainText(
            f"─── {len(outputs)} file(s) saved to: "
            f"{os.path.dirname(outputs[0]) if outputs else '—'}"
        )
        if outputs:
            self._open_folder_btn.show()

    def _on_error(self, msg: str):
        self._progress_bar.hide()
        self._running = False
        self._refresh_run_state()
        self._status_lbl.setStyleSheet(f"color:{RED};")
        self._status_lbl.setText(f"Error: {error_summary(msg)}")
        self._log_edit.appendPlainText(f"ERROR: {msg}")
        show_error_dialog(self, "FASTA Tools error", msg)

    def _open_output_folder(self):
        if not self._last_outputs:
            return
        folder = os.path.dirname(self._last_outputs[0])
        if os.path.isdir(folder):
            QtGui.QDesktopServices.openUrl(
                QtCore.QUrl.fromLocalFile(folder)
            )
